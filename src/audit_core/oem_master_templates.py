"""The standard upload templates for the OEM masters (decision 2026-10-06): one Excel template per master
that has an Excel reader, built here beside the readers so a test can prove each template is read by its own
reader. People are asked to upload only in the template; a template still carrying its sample rows is refused
(see `refuse_template_samples`) so a sample can never be loaded as real data.

The consumer scheme and the exchange / scrappage bulletins are the OEM's own PDFs, read as printed: they have
no template."""
from __future__ import annotations

from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from audit_core.oem_master_parsers import (
    STANDARD_PRICE_COLUMNS,
    STANDARD_PRICE_META_KEY,
    STANDARD_PRICE_META_SHEET,
    STANDARD_PRICE_VERSION,
)

# Every sample cell of a template carries this text; the upload is refused while any cell still does.
TEMPLATE_SAMPLE_MARKER = "SAMPLE - DELETE"

TEMPLATE_KINDS = ("PRICE_LIST", "CORPORATE_POLICY", "DISCOUNT_GRID")
TEMPLATE_FILENAMES = {
    "PRICE_LIST": "Verigence price list template.xlsx",
    "CORPORATE_POLICY": "Verigence corporate policy template.xlsx",
    "DISCOUNT_GRID": "Verigence discount grid template.xlsx",
}

_BOLD = Font(bold=True)
_HEAD_FILL = PatternFill("solid", fgColor="E6F4F1")
_SAMPLE_FILL = PatternFill("solid", fgColor="FFF3CD")


def _readme(wb: Workbook, title: str, lines: list[str]) -> None:
    ws = wb.create_sheet("READ ME")
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _style_header(ws: Any, row: int, width: int) -> None:
    for col in range(1, width + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = _BOLD
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _mark_sample(ws: Any, first_row: int, last_row: int, width: int) -> None:
    for row in range(first_row, last_row + 1):
        for col in range(1, width + 1):
            ws.cell(row=row, column=col).fill = _SAMPLE_FILL


def _to_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ── price list ──────────────────────────────────────────────────────────────────
def sample_price_rows() -> list[list[Any]]:
    """Two sample vehicles whose on-road prices add up, so the template is itself a valid file."""
    rows = []
    for variant, trim, ex, reg in (
        (f"{TEMPLATE_SAMPLE_MARKER} variant 1", "S1", 1_000_000, 100_000),
        (f"{TEMPLATE_SAMPLE_MARKER} variant 2", "S2", 1_200_000, 120_000),
    ):
        tcs, ins, ew4, ew45, kit, essential, rsa, fastag = 10_000, 45_000, 18_000, 33_000, 30_000, 0, 2_000, 500
        on_road = ex + tcs + ins + kit + essential + rsa + fastag + reg
        rows.append([
            f"{TEMPLATE_SAMPLE_MARKER} model", variant, trim, "DIESEL", "MT", "2WD", 5, "", ex, tcs, ins, ew4, ew45, kit, essential,
            rsa, fastag, reg, 1_500, on_road, 21_000, "Permit charges", 1_700, None, None,
        ])
    return rows


def price_list_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Price List"
    ws.append(["WEF date (dd-mm-yyyy)", None])
    ws.append(list(STANDARD_PRICE_COLUMNS))
    _style_header(ws, 2, len(STANDARD_PRICE_COLUMNS))
    ws["A1"].font = _BOLD
    first_sample = ws.max_row + 1
    for row in sample_price_rows():
        ws.append(row)
    _mark_sample(ws, first_sample, ws.max_row, len(STANDARD_PRICE_COLUMNS))
    ws.freeze_panes = "C3"
    for col in range(1, len(STANDARD_PRICE_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.row_dimensions[2].height = 60
    _readme(wb, "Verigence price list template", [
        "Upload your OEM price list ONLY in this template. Do not add, rename, move or delete a column; a changed template is refused.",
        "1. One row per vehicle. Keep one sheet; add as many rows as you need. A vehicle may appear once only.",
        "2. Type the WEF date (the date the prices take effect) in the cell next to 'WEF date' at the top. You are asked for it again when you upload, and what you type there wins.",
        "3. Model, Variant, Ex-Showroom Price, Insurance, Registration and the On-Road Price are required for every row. Trim, Fuel (Petrol / Diesel / CNG / Electric), Transmission, Drive and Seater help search; please fill them.",
        "4. Insurance Type: leave blank for one insurance price. If a vehicle has a private AND a commercial insurance price, enter it twice, once with PRIVATE and once with COMMERCIAL.",
        "5. Enter prices as plain numbers: no commas, no rupee sign. A line the vehicle does not have can be left blank, 0 or NA (all read as nil). Any other text in a price cell is an error. Insurance is never blank.",
        "6. The On-Road Price is Ex-Showroom + TCS + Insurance + Accessories Kit + Essential Accessories + RSA + FASTag + Registration. It carries NO extended warranty and NO hypothecation. Both warranty options are listed in their own columns. The file is refused if a row does not add up.",
        "7. Hypothecation Charge (the extra amount when the vehicle is financed), Minimum Booking Amount and up to two Extra Charges (a name and an amount, for example a permit charge) are optional.",
        "8. The yellow rows are samples. Delete them, and every cell that says 'SAMPLE - DELETE', before uploading; a file that still has them is refused.",
        "9. The file keeps the name you give it, so name it so that you can find it later (for example the OEM, the model and the date).",
    ])
    meta = wb.create_sheet(STANDARD_PRICE_META_SHEET)
    meta["A1"], meta["B1"] = "master_key", STANDARD_PRICE_META_KEY
    meta["A2"], meta["B2"] = "template_version", STANDARD_PRICE_VERSION
    meta.sheet_state = "hidden"
    return _to_bytes(wb)


# ── corporate policy ────────────────────────────────────────────────────────────
_CORPORATE_CATEGORIES = ("Cat-B", "Cat-A", "Cat-F", "Cat-Y (Premium)", "Cat-Z (Signature)")
_COMPANY_BLOCKS = ('"Z" Signature', '"Y" Premium', '"F" Focus', '"A"', '"B"')


def corporate_policy_template() -> bytes:
    wb = Workbook()
    policy = wb.active
    policy.title = "Corporate Policy"
    readme_lines = [
        "Upload the corporate privilege policy ONLY in this template. It has two sheets: 'Corporate Policy' and 'Companies List'. Keep both sheet names.",
        "Corporate Policy: type the brand (as the OEM names it) over each M&M / Dealer / Total group. For every category give the OEM's share (M&M), the dealer's share and the total. Total must equal M&M + Dealer.",
        "Type the date the policy takes effect in the first line as 'From dd.mm.yyyy'. You are also asked for it when you upload, and what you type there wins.",
        "Companies List: one block per category. Each company needs its type, its name and its corporate code. A code may appear once only.",
        "The yellow cells are samples. Delete every sample row and every cell that says 'SAMPLE - DELETE' before uploading; a file that still has them is refused.",
    ]
    policy.append(["Corporate Privilege Policy - From DD.MM.YYYY (replace with the date)"])
    policy.append([])
    policy.append([None, "Corporate Category", f"{TEMPLATE_SAMPLE_MARKER} brand", None, None])
    policy.append([None, None, "M&M", "Dealer", "Total"])
    _style_header(policy, 3, 5)
    _style_header(policy, 4, 5)
    for category in _CORPORATE_CATEGORIES:
        policy.append([None, category, 3000, 2000, 5000])
    _mark_sample(policy, 5, policy.max_row, 5)
    policy.column_dimensions["B"].width = 24
    companies = wb.create_sheet("Companies List")
    title_row: list[Any] = []
    header_row: list[Any] = []
    sample_row: list[Any] = []
    for index, block in enumerate(_COMPANY_BLOCKS, start=1):
        title_row += [block, None, None, None]
        header_row += ["S. No.", "Corporate Type", "Description", "Corporate Code"]
        sample_row += [1, f"{TEMPLATE_SAMPLE_MARKER} type", f"{TEMPLATE_SAMPLE_MARKER} company {index}", f"{TEMPLATE_SAMPLE_MARKER}-{index}"]
    companies.append(title_row)
    companies.append(header_row)
    companies.append(sample_row)
    _style_header(companies, 1, len(header_row))
    _style_header(companies, 2, len(header_row))
    _mark_sample(companies, 3, 3, len(header_row))
    for col in range(1, len(header_row) + 1):
        companies.column_dimensions[get_column_letter(col)].width = 22
    _readme(wb, "Verigence corporate privilege policy template", readme_lines)
    return _to_bytes(wb)


# ── dealer discount grid ────────────────────────────────────────────────────────
GRID_HEADER_ROW = ["Model", "Booking Protection", "Agreed Buffer", "Insurance OD %", "Out of Territory"]


def discount_grid_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Discount Grid"
    readme_lines = [
        "Upload the dealer discount grid ONLY in this template (one sheet, 'Discount Grid').",
        "Type the date it takes effect in the first line as 'Effective dd Mon yyyy' (for example Effective 01 Sep 2026). You are also asked for it when you upload, and what you type there wins.",
        "One row per model; a cell may name several models separated by commas. Booking Protection: days (for example 30 days). Agreed Buffer and Out of Territory: rupees (for example 5000 or 5K), or Nil. Insurance OD %: a percentage (for example 60% or 0.6), or Nil. Write 'Out of scope' where the grid does not apply.",
        "Under the table, after a row that reads Parameter | Notes, list the policy wording, one parameter per row.",
        "The yellow cells are samples. Delete every sample row and every cell that says 'SAMPLE - DELETE' before uploading; a file that still has them is refused.",
    ]
    ws.append(["Dealer discount grid - Effective DD Mon YYYY (replace with the date)"])
    ws.append(GRID_HEADER_ROW)
    _style_header(ws, 2, len(GRID_HEADER_ROW))
    ws.append([f"{TEMPLATE_SAMPLE_MARKER} model", "30 days", "Nil", "60%", "Out of scope"])
    _mark_sample(ws, 3, 3, len(GRID_HEADER_ROW))
    ws.append([])
    ws.append(["Parameter", "Notes"])
    _style_header(ws, 5, 2)
    ws.append([f"{TEMPLATE_SAMPLE_MARKER} parameter", f"{TEMPLATE_SAMPLE_MARKER} note"])
    _mark_sample(ws, 6, 6, 2)
    ws.column_dimensions["A"].width = 36
    for col in "BCDE":
        ws.column_dimensions[col].width = 20
    _readme(wb, "Verigence dealer discount grid template", readme_lines)
    return _to_bytes(wb)


_BUILDERS = {
    "PRICE_LIST": price_list_template,
    "CORPORATE_POLICY": corporate_policy_template,
    "DISCOUNT_GRID": discount_grid_template,
}


def build_template(kind: str) -> bytes | None:
    """The template for a master kind, or None where the master is the OEM's own PDF."""
    builder = _BUILDERS.get(kind)
    return builder() if builder else None


def refuse_template_samples(content: bytes) -> str | None:
    """Plain words for the person when a workbook still has a template's sample cells; None when it has none."""
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError):  # not a workbook: the reader says so itself
        return None
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str) and TEMPLATE_SAMPLE_MARKER in cell:
                    return (
                        f"The file still has the template's sample rows (cells that say '{TEMPLATE_SAMPLE_MARKER}', "
                        f"for example on sheet '{sheet.title}'). Delete every sample row and note, then upload again."
                    )
    return None
