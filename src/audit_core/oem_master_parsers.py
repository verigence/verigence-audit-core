"""Parsers for an OEM's *native* price / discount documents.

Four inputs, each in the OEM's own layout (not a Verigence template):

  PRICE_LIST       .xlsx  consolidated price list  -> one row per sellable SKU,
                          or the dealer's per-model price sheets (one sheet per
                          model, PV/CV and EV layouts; decision 2026-09-30)
  CONSUMER_SCHEME  .pdf   monthly consumer scheme bulletin
  EXCHANGE_SCHEME  .pdf   Xmart exchange / scrappage / welcome ready reckoner
  CORPORATE_POLICY .xlsx  Corporate Privilege Policy (matrix + company list)

Every parser returns a :class:`ParseResult`. Money figures are only trusted when
they reconcile against the document's own control totals (on-road price, A+B =
total, cash+other = total consumer offer); a row that does not reconcile is put
in ``errors`` and never ingested. Nothing is inferred.

Pure module: no DB, no network. ``pdfplumber`` is imported lazily so the two
xlsx parsers work without it.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

# ── vocabulary ──────────────────────────────────────────────────────────────────
# price_list_items.component_key — the standard line items a vehicle is priced on.
PRICE_COMPONENT_KEYS: tuple[str, ...] = (
    "EX_SHOWROOM",
    "TCS",
    "INSURANCE",
    "EXT_WARRANTY_4TH_YR",
    "EXT_WARRANTY_4TH_5TH_YR",
    "ACCESSORIES_KIT",
    "RSA_1YR",
    "FASTAG",
    "REGISTRATION_INDIVIDUAL",
    "REGISTRATION_CORPORATE",
)

# discount_scheme_benefits.benefit_key — a discount is "how much is taken off a
# line item" (component-aligned) plus the trade-in / privilege buckets.
DISCOUNT_BENEFIT_KEYS: frozenset[str] = frozenset(
    {
        "CASH_DISCOUNT",
        "ACCESSORIES_KIT",
        "EXT_WARRANTY_4TH_YR",
        "EXT_WARRANTY_4TH_5TH_YR",
        "INSURANCE",
        "OTHER_SCHEME",
        "EXCHANGE_BONUS",
        "SCRAPPAGE_BONUS_DEALER",
        "SCRAPPAGE_BONUS_COD",
        "WELCOME_BONUS",
        "CORPORATE_PRIVILEGE",
    }
)

SCHEME_CATEGORIES: frozenset[str] = frozenset(
    {"CONSUMER", "EXCHANGE", "SCRAPPAGE", "WELCOME", "CORPORATE"}
)

_MONEY_TOLERANCE = Decimal("1.00")  # source figures carry sub-rupee float noise
_TWO_PLACES = Decimal("0.01")


# ── result shapes ───────────────────────────────────────────────────────────────
@dataclass
class PriceRow:
    row_no: int
    category: str  # PV / CV / BEV
    model_name: str
    variant_name: str
    trim: str | None
    fuel: str | None
    transmission: str | None
    drive: str | None
    seater: str | None
    components: dict[str, Decimal]
    onroad_individual: Decimal
    onroad_corporate: Decimal
    registration_basis: str = "STANDARD"  # STANDARD / PRIVATE / COMMERCIAL
    source_sheet: str | None = None


@dataclass
class DiscountRow:
    row_no: int
    scheme_category: str
    model_alias: str
    variant_texts: list[str]  # [] => whole model / brand
    benefits: list[tuple[str, Decimal]]  # (benefit_key, amount)
    total_customer_offer: Decimal
    config: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


@dataclass
class CorporateBenefitRow:
    row_no: int
    privilege_category: str  # Z / Y / F / A / B
    brand_alias: str
    m_and_m: Decimal
    dealer: Decimal
    total: Decimal


@dataclass
class CorporateCompany:
    corporate_code: str
    corporate_name: str
    corporate_type: str | None
    privilege_category: str  # Z / Y / F / A / B


@dataclass
class GridRow:
    """One model line of the dealer's discount grid (decision 2026-09-30):
    booking protection, the agreed buffer, the insurance OD percentage
    (shown as a maximum for now) and the out-of-territory addition."""

    row_no: int
    model_alias: str
    model_aliases: list[str]
    in_scope: bool
    booking_protection_days: int | None
    agreed_buffer_amount: Decimal | None
    insurance_od_percent: Decimal | None
    out_of_territory_amount: Decimal | None
    raw: dict[str, Any] = field(default_factory=dict)
    # The October 2026 grid gives the buffer by tier (2026-10-20): BASE
    # (without in-house insurance or EW), WITH_INSURANCE, WITH_EW_INSURANCE.
    # A single-buffer grid has no tiers; agreed_buffer_amount is then the
    # buffer for every deal, and otherwise it is the BASE tier.
    tiers: dict[str, Decimal] = field(default_factory=dict)


@dataclass
class ParseResult:
    kind: str
    effective_from_hint: date | None = None
    price_rows: list[PriceRow] = field(default_factory=list)
    discount_rows: list[DiscountRow] = field(default_factory=list)
    corporate_benefits: list[CorporateBenefitRow] = field(default_factory=list)
    corporate_companies: list[CorporateCompany] = field(default_factory=list)
    grid_rows: list[GridRow] = field(default_factory=list)
    grid_parameters: list[dict[str, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    meta: dict[str, Any] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.errors


class MasterParseError(ValueError):
    """The file could not be recognised as the declared master kind."""


# ── shared helpers ──────────────────────────────────────────────────────────────
def _money(value: Any) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    text = str(value).strip()
    text = text.replace("₹", "").replace("Rs.", "").replace("Rs", "")
    text = text.replace(",", "").replace("/-", "")
    text = re.sub(r"\s+", "", text)  # OEM PDFs split digits: "2 1,186" -> "21186"
    if text in ("", "-", "NA", "N/A"):
        return Decimal(0)
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _component_money(value: Any) -> Decimal | None:
    """A price-list component cell. An empty cell is 0: the OEM's legend
    says a component not charged on that sheet reads 0, and the October 2026
    book leaves both warranty cells of the THAR ROXX AX7 Star Edition rows
    empty (2026-10-20 product owner rules). The row still has to add up to
    its own on-road price, so a blank that hid a real figure is caught
    there. Text that is not a number is still an error."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return Decimal(0)
    return _money(value)


def _q2(value: Decimal) -> Decimal:
    return value.quantize(_TWO_PLACES, rounding=ROUND_HALF_UP)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


_HEADER_NOISE_RE = re.compile(r"[₹$€£]+$")


def _header_text(value: Any) -> str:
    """Header-cell text, with trailing currency-symbol noise stripped.

    Confirmed live against a real uploaded master: a header cell literally
    reads "Model₹" (the Rupee symbol stuck to the label, not a stray column
    -- Sl. No./Category/Variant/Trim/... on either side are all otherwise
    unaffected), which fails parse_price_list's strict header-layout check
    outright and rejects the whole file -- with a real Trim column the
    picker needs, but nothing downstream ever saw it. A header cell is a
    label, never itself a currency value (the monetary *columns* are
    validated separately via _money on the data rows), so stripping a
    trailing currency glyph here can never mask a genuine layout mismatch.
    """
    return _HEADER_NOISE_RE.sub("", _text(value) or "")


def _slug_model(name: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", name.strip().upper()).strip("_")


_WEF_RE = re.compile(
    r"(?:w\.?e\.?f\.?|valid\s*from|from)\s*:?\s*(?:dt\.?\s*)?(\d{1,2})\s*(?:st|nd|rd|th)?\s*"
    r"[.\-/ ]*\s*(\d{1,2}|[A-Za-z]{3,9})\.?\s*[.\-/ '’;]*\s*(\d{2,4})",
    re.IGNORECASE,
)
_MONTHS = {
    m[:3].lower(): i
    for i, m in enumerate(
        [
            "January",
            "February",
            "March",
            "April",
            "May",
            "June",
            "July",
            "August",
            "September",
            "October",
            "November",
            "December",
        ],
        start=1,
    )
}


def _date_hint(blob: str) -> date | None:
    match = _WEF_RE.search(blob or "")
    if not match:
        return None
    day_raw, month_raw, year_raw = match.groups()
    try:
        day = int(day_raw)
        month = (
            int(month_raw)
            if month_raw.isdigit()
            else _MONTHS.get(month_raw[:3].lower())
        )
        year = int(year_raw)
        if year < 100:
            year += 2000
        if not month:
            return None
        return date(year, month, day)
    except (ValueError, TypeError):
        return None


_LOOSE_DATE_RE = re.compile(
    r"(?<!\d)(\d{1,2})\s*(?:st|nd|rd|th)?[\s_.\-]*([A-Za-z]{3,9})[\s_.\-]*'?(\d{2,4})(?!\d)"
)


def _loose_date_hint(blob: str) -> date | None:
    """A date written the way a file is named: "03Sep2026", "01th_Sept26",
    "1_Aug_2026". Month by name only, so a version number never reads as one."""
    for match in _LOOSE_DATE_RE.finditer(blob or ""):
        day_raw, month_raw, year_raw = match.groups()
        month = _MONTHS.get(month_raw[:3].lower())
        if not month:
            continue
        year = int(year_raw)
        if year < 100:
            year += 2000
        try:
            return date(year, month, int(day_raw))
        except ValueError:
            continue
    return None


# ── 1. price list ───────────────────────────────────────────────────────────────
_PRICE_HEADER = (
    "Sl. No.",
    "Category",
    "Model",
    "Variant",
    "Trim",
    "Fuel",
    "Transmission",
    "Drive",
    "Seater",
    "Ex-Showroom Price",
    "TCS",
    "Insurance",
)
_PRICE_COL = {
    "EX_SHOWROOM": 9,
    "TCS": 10,
    "INSURANCE": 11,
    "EXT_WARRANTY_4TH_YR": 12,
    "EXT_WARRANTY_4TH_5TH_YR": 13,
    "ACCESSORIES_KIT": 14,
    "RSA_1YR": 15,
    "FASTAG": 16,
    "REGISTRATION_INDIVIDUAL": 17,
    "REGISTRATION_CORPORATE": 19,
}
_ONROAD_INDIVIDUAL_COL = 18
_ONROAD_CORPORATE_COL = 20
_SOURCE_SHEET_COL = 21
_VALID_CATEGORIES = {"PV", "CV", "BEV"}
# on-road = Σ(all components) with REGISTRATION_INDIVIDUAL for the individual total
_ONROAD_INDIVIDUAL_PARTS = [
    k for k in PRICE_COMPONENT_KEYS if k != "REGISTRATION_CORPORATE"
]
_ONROAD_CORPORATE_PARTS = [
    k for k in PRICE_COMPONENT_KEYS if k != "REGISTRATION_INDIVIDUAL"
]


def parse_price_list(content: bytes, *, filename: str | None = None) -> ParseResult:
    """The OEM's consolidated list, or the dealer's per-model sheets: the
    layout is recognised from the workbook, never declared."""
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    sheet = None
    for name in workbook.sheetnames:
        if name.strip().lower() in {"price list", "pricelist", "consolidated price list"}:
            sheet = workbook[name]
            break
    if sheet is None:
        sheet = workbook[workbook.sheetnames[0]]

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise MasterParseError("Price list sheet is empty.")
    header = [_header_text(c) for c in rows[0][: len(_PRICE_HEADER)]]
    if header != list(_PRICE_HEADER):
        result = _parse_dealer_sheets(workbook, header_seen=header)
    else:
        result = _parse_consolidated(rows)
    from_name = _loose_date_hint(filename or "")
    if result.effective_from_hint is None and from_name is not None:
        result.effective_from_hint = from_name
        result.meta["effectiveFromSource"] = "FILENAME"
    elif (
        result.meta.get("layout") == "DEALER_PER_MODEL"
        and from_name is not None
        and from_name > result.effective_from_hint
    ):
        # A dealer sheet re-issued with additions keeps the month's
        # "w.e.f. 01.09.2026" header while its name carries the day the
        # additions came ("02 Nos Added EV PRICE -07th Sept'26"). Dating it
        # by the header priced the two added BE 6 SPORTEQ variants from
        # 01 Sep, before they existed, and dropped them from 03 Sep (2026-10-20
        # review). The later date is the file's own: the admin can still
        # enter another.
        result.warnings.append(
            f"The sheet says effective {result.effective_from_hint.isoformat()} but the file name says "
            f"{from_name.isoformat()}; the later date, {from_name.isoformat()}, is applied. Enter a date to override."
        )
        result.effective_from_hint = from_name
        result.meta["effectiveFromSource"] = "FILENAME"
    elif result.effective_from_hint is not None:
        result.meta["effectiveFromSource"] = "SHEET"
        if from_name is not None and from_name != result.effective_from_hint:
            result.warnings.append(
                f"The sheet says effective {result.effective_from_hint.isoformat()} but the file name says "
                f"{from_name.isoformat()}; enter the date to be sure."
            )
    return result


def _parse_consolidated(rows: list[Any]) -> ParseResult:
    result = ParseResult(kind="PRICE_LIST", meta={"layout": "OEM_CONSOLIDATED"})

    hint_blob = " ".join(
        str(c)
        for r in rows[-14:]
        for c in r
        if isinstance(c, str)
    )
    result.effective_from_hint = _date_hint(hint_blob)

    seen: dict[tuple[str, ...], Any] = {}
    for idx, raw in enumerate(rows[1:], start=2):
        category = _text(raw[1]) if len(raw) > 1 else None
        if category not in _VALID_CATEGORIES:
            continue
        ex_showroom = _money(raw[9]) if len(raw) > 9 else None
        if ex_showroom is None or ex_showroom <= 0:
            continue
        model_name = _text(raw[2])
        variant_name = _text(raw[3])
        if not model_name or not variant_name:
            result.errors.append(f"row {idx}: missing model/variant")
            continue

        components: dict[str, Decimal] = {}
        bad = False
        for comp_key, col in _PRICE_COL.items():
            value = _component_money(raw[col] if len(raw) > col else None)
            if value is None:
                result.errors.append(
                    f"row {idx}: non-numeric {comp_key} '{raw[col] if len(raw) > col else None}'"
                )
                bad = True
                break
            components[comp_key] = value
        if bad:
            continue

        onroad_ind = _money(raw[_ONROAD_INDIVIDUAL_COL]) if len(raw) > _ONROAD_INDIVIDUAL_COL else None
        onroad_corp = _money(raw[_ONROAD_CORPORATE_COL]) if len(raw) > _ONROAD_CORPORATE_COL else None
        if onroad_ind is None or onroad_corp is None:
            result.errors.append(f"row {idx}: missing on-road price")
            continue

        calc_ind = sum((components[k] for k in _ONROAD_INDIVIDUAL_PARTS), Decimal(0))
        calc_corp = sum((components[k] for k in _ONROAD_CORPORATE_PARTS), Decimal(0))
        if abs(calc_ind - onroad_ind) > _MONEY_TOLERANCE:
            result.errors.append(
                f"row {idx} ({model_name} / {variant_name}): components sum to "
                f"{_q2(calc_ind)} but on-road (individual) is {_q2(onroad_ind)}"
            )
            continue
        if abs(calc_corp - onroad_corp) > _MONEY_TOLERANCE:
            result.errors.append(
                f"row {idx} ({model_name} / {variant_name}): components sum to "
                f"{_q2(calc_corp)} but on-road (corporate) is {_q2(onroad_corp)}"
            )
            continue

        fuel = _text(raw[5]) if len(raw) > 5 else None
        transmission = _text(raw[6]) if len(raw) > 6 else None
        drive = _text(raw[7]) if len(raw) > 7 else None
        seater = _text(raw[8]) if len(raw) > 8 else None
        source_sheet = _text(raw[_SOURCE_SHEET_COL]) if len(raw) > _SOURCE_SHEET_COL else None
        registration_basis = _registration_basis(source_sheet)
        signature = (
            _q2(onroad_ind),
            _q2(onroad_corp),
            tuple(_q2(components[k]) for k in PRICE_COMPONENT_KEYS),
        )
        key = (
            _slug_model(model_name),
            variant_name.upper(),
            (fuel or "").upper(),
            (transmission or "").upper(),
            (drive or "").upper(),
            (seater or "").upper(),
            registration_basis,
        )
        if key in seen:
            if seen[key] != signature:
                result.errors.append(
                    f"row {idx} ({model_name} / {variant_name}): conflicting prices for the "
                    "same model/variant/fuel/transmission/drive"
                )
            else:
                result.meta["duplicateRowsCollapsed"] = (
                    result.meta.get("duplicateRowsCollapsed", 0) + 1
                )
            continue
        seen[key] = signature

        result.price_rows.append(
            PriceRow(
                row_no=idx,
                category=category,
                model_name=model_name,
                variant_name=variant_name,
                trim=_text(raw[4]) if len(raw) > 4 else None,
                fuel=fuel,
                transmission=transmission,
                drive=drive,
                seater=seater,
                components={k: _q2(v) for k, v in components.items()},
                onroad_individual=_q2(onroad_ind),
                onroad_corporate=_q2(onroad_corp),
                registration_basis=registration_basis,
                source_sheet=source_sheet,
            )
        )

    result.meta["models"] = sorted({r.model_name for r in result.price_rows})
    result.meta["categoryCounts"] = _counts(r.category for r in result.price_rows)
    if not result.price_rows and not result.errors:
        raise MasterParseError("Price list contained no recognisable vehicle rows.")
    return result


def _registration_basis(source_sheet: str | None) -> str:
    blob = (source_sheet or "").upper()
    if re.search(r"\bCOM\b|-COM|COMMERCIAL", blob):
        return "COMMERCIAL"
    if "PVT" in blob or "PRIVATE" in blob:
        return "PRIVATE"
    return "STANDARD"


def _counts(values: Any) -> dict[str, int]:
    out: dict[str, int] = {}
    for value in values:
        out[value] = out.get(value, 0) + 1
    return out


# ── 1b. the dealer's per-model price sheets ─────────────────────────────────────
# One worksheet per model. A "MODEL NAME" cell names it, a "Model & Variant"
# header row starts the table, a row of column letters "(A) (B) ... (I)=(A+B+
# C+D+E+F+G+H)" carries the sheet's own on-road formula, and the data rows
# follow. Component columns are recognised by their label; the registration
# and on-road columns come in pairs, the first pair individual (without
# hypothecation), the second corporate (with). Every row must reconcile
# against the sheet's own formula or it is not ingested.
_LETTER_RE = re.compile(r"^\(([A-Z])\)\s*(?:=\s*\(?\s*([A-Z+\s]+?)\s*\)?\s*)?$")
_TRANSMISSION_TOKENS = {"MT", "AT", "AMT", "CVT", "DCT", "IVT"}
_DRIVE_TOKENS = {"2WD", "4WD", "AWD", "4X4", "4X2"}
_FUEL_TOKENS = {"PETROL": "PETROL", "DIESEL": "DIESEL", "CNG": "CNG", "ELECTRIC": "ELECTRIC", "EV": "ELECTRIC"}
_SEATER_RE = re.compile(r"\b(\d{1,2})\s*(?:STR|SEATER|SEAT)\b", re.IGNORECASE)


def _component_for_label(label: str) -> str | None:
    up = label.upper()
    if "EX-SHOWROOM" in up or "EX SHOWROOM" in up or "EXSHOWROOM" in up:
        return "EX_SHOWROOM"
    if "TAX COLLECTION" in up or re.search(r"\bTCS\b", up):
        return "TCS"
    if "INSURANCE" in up:
        return "INSURANCE"
    if "WARRANTY" in up:
        return "EXT_WARRANTY_4TH_5TH_YR" if re.search(r"5\s*TH", up) else "EXT_WARRANTY_4TH_YR"
    if "ACCESSOR" in up:
        return "ACCESSORIES_KIT"
    if re.search(r"\bRSA\b", up) or "ROAD SIDE" in up or "ROADSIDE" in up:
        return "RSA_1YR"
    if "FASTAG" in up or "FAST TAG" in up:
        return "FASTAG"
    return None


def _variant_attributes(variant: str, *, fuel_hint: str | None, electric: bool) -> dict[str, str | None]:
    tokens = [t for t in re.split(r"[\s\-/,()]+", variant.upper()) if t]
    fuel = "ELECTRIC" if electric else (fuel_hint or None)
    if fuel is None:
        fuel = next((_FUEL_TOKENS[t] for t in tokens if t in _FUEL_TOKENS), None)
    transmission = next((t for t in tokens if t in _TRANSMISSION_TOKENS), None)
    drive = next((t for t in tokens if t in _DRIVE_TOKENS), None)
    seater_match = _SEATER_RE.search(variant)
    return {
        "fuel": fuel,
        "transmission": transmission,
        "drive": drive,
        "seater": seater_match.group(1) if seater_match else None,
    }


def _cell_reader(raw: list[Any]) -> Any:
    def cell(col: int) -> Any:
        return raw[col] if col < len(raw) else None
    return cell


def _labelled_cell_after(rows: list[Any], label: str) -> str | None:
    """The first non-empty cell to the right of a cell that reads ``label``."""
    for raw in rows:
        for col, cell in enumerate(raw):
            if _text(cell) and _text(cell).upper().rstrip(":") == label:
                for value in raw[col + 1:]:
                    if _text(value):
                        return _text(value)
    return None


def _dealer_sheet_layout(rows: list[Any]) -> dict[str, Any] | None:
    letter_row = next(
        (i for i, raw in enumerate(rows)
         if sum(1 for c in raw if _text(c) and _LETTER_RE.match(_text(c).replace(" ", "").upper())) >= 3),
        None,
    )
    if letter_row is None:
        return None
    header_row = next(
        (i for i in range(letter_row - 1, -1, -1)
         if any(_text(c) and "VARIANT" in _text(c).upper() for c in rows[i])),
        None,
    )
    if header_row is None:
        return None
    letter_by_col: dict[int, str] = {}
    formula_by_letter: dict[str, list[str]] = {}
    for col, cell in enumerate(rows[letter_row]):
        match = _LETTER_RE.match((_text(cell) or "").replace(" ", "").upper())
        if not match:
            continue
        letter_by_col[col] = match.group(1)
        if match.group(2):
            formula_by_letter[match.group(1)] = [x for x in match.group(2).replace(" ", "").split("+") if x]
    col_by_letter = {letter: col for col, letter in letter_by_col.items()}

    def label(col: int) -> str:
        return " ".join(
            _text(rows[i][col]) or "" for i in range(header_row, letter_row) if col < len(rows[i]) and _text(rows[i][col])
        )

    variant_col = next(col for col, c in enumerate(rows[header_row]) if _text(c) and "VARIANT" in _text(c).upper())
    components: dict[str, int] = {}
    registration_cols: list[int] = []
    onroad_cols: list[int] = []
    for col in sorted(letter_by_col):
        text_label = label(col).upper()
        if "REGISTRATION" in text_label:
            registration_cols.append(col)
        elif "ON ROAD" in text_label or "ON-ROAD" in text_label or "ONROAD" in text_label:
            onroad_cols.append(col)
        else:
            key = _component_for_label(text_label)
            if key and key not in components:
                components[key] = col
    if "EX_SHOWROOM" not in components or not registration_cols or not onroad_cols:
        return None
    return {
        "letter_row": letter_row,
        "variant_col": variant_col,
        "components": components,
        "registration_cols": registration_cols,
        "onroad_cols": onroad_cols,
        "letter_by_col": letter_by_col,
        "col_by_letter": col_by_letter,
        "formula_by_letter": formula_by_letter,
    }


# Mahindra's electric models (2026-10-20 product owner rules, R4): BE 6 /
# BE 6 SPORTEQ / BE6, XEV 9E, XEV 9S, XUV400 and XUV 3XO EV, or any EV /
# ELECTRIC wording. "XEV" and "BE 6" never contain the word "EV", so the
# dealer's EV sheets for them used to come out as category UNSPECIFIED.
_ELECTRIC_RE = re.compile(
    r"\bEV\b|ELECTRIC|\bBEV\b|\bBE\s*6|\bXEV\s*9|\bXUV\s*400|\b3XO\s*EV|WITH\s+CHARGER"
)


def _is_electric_sheet(model_name: str, title: str | None, fuel_hint: str | None, header_rows: list[Any]) -> bool:
    """An EV price sheet: the model is one of the EV models, the fuel says
    electric, or the ex-showroom column is priced "with Charger"."""
    header_text = " ".join(str(c) for raw in header_rows for c in raw if isinstance(c, str))
    blob = f"{model_name} {title or ''} {fuel_hint or ''} {header_text}".upper()
    return bool(_ELECTRIC_RE.search(blob))


def _parse_dealer_sheets(workbook: Any, *, header_seen: list[str]) -> ParseResult:
    result = ParseResult(kind="PRICE_LIST", meta={"layout": "DEALER_PER_MODEL"})
    seen: dict[tuple[str, ...], Any] = {}
    sheets_read = 0
    for sheet in workbook.worksheets:
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        layout = _dealer_sheet_layout(rows)
        if layout is None:
            continue
        sheets_read += 1
        model_name = _labelled_cell_after(rows[: layout["letter_row"]], "MODEL NAME") or _text(sheet.title)
        if not model_name:
            result.errors.append(f"sheet '{sheet.title}': no model name")
            continue
        fuel_hint = _labelled_cell_after(rows[: layout["letter_row"]], "FUEL TYPE")
        fuel_hint = fuel_hint.upper() if fuel_hint else None
        electric = _is_electric_sheet(model_name, sheet.title, fuel_hint, rows[: layout["letter_row"] + 1])
        if result.effective_from_hint is None:
            blob = " ".join(str(c) for raw in rows[: layout["letter_row"]] for c in raw if isinstance(c, str))
            result.effective_from_hint = _date_hint(blob)
        pairs = list(zip(layout["registration_cols"], layout["onroad_cols"]))
        if len(pairs) == 1:
            result.warnings.append(
                f"sheet '{sheet.title}': one registration / on-road pair; the corporate basis takes the same figures"
            )
        comp_cols: dict[str, int] = layout["components"]

        for idx, raw in enumerate(rows[layout["letter_row"] + 1:], start=layout["letter_row"] + 2):
            cell = _cell_reader(raw)
            variant_name = _text(cell(layout["variant_col"]))
            ex_showroom = _money(cell(comp_cols["EX_SHOWROOM"]))
            if not variant_name or ex_showroom is None or ex_showroom <= 0:
                continue
            components: dict[str, Decimal] = {}
            bad = False
            for key, col in comp_cols.items():
                value = _component_money(cell(col))
                if value is None:
                    result.errors.append(f"sheet '{sheet.title}' row {idx}: non-numeric {key} '{cell(col)}'")
                    bad = True
                    break
                components[key] = value
            if bad:
                continue
            basis_values: list[tuple[Decimal, Decimal]] = []
            for reg_col, onroad_col in pairs:
                registration = _component_money(cell(reg_col))
                onroad = _money(cell(onroad_col))
                if registration is None or onroad is None:
                    result.errors.append(f"sheet '{sheet.title}' row {idx} ({variant_name}): missing registration or on-road price")
                    bad = True
                    break
                letters = layout["formula_by_letter"].get(layout["letter_by_col"].get(onroad_col, ""))
                if letters:
                    calc = sum((_component_money(cell(layout["col_by_letter"][x])) or Decimal(0) for x in letters
                                if x in layout["col_by_letter"]), Decimal(0))
                else:
                    calc = sum(components.values(), Decimal(0)) + registration
                if abs(calc - onroad) > _MONEY_TOLERANCE:
                    result.errors.append(
                        f"sheet '{sheet.title}' row {idx} ({model_name} / {variant_name}): components sum to "
                        f"{_q2(calc)} but on-road is {_q2(onroad)}"
                    )
                    bad = True
                    break
                basis_values.append((registration, onroad))
            if bad:
                continue
            (reg_ind, onroad_ind) = basis_values[0]
            (reg_corp, onroad_corp) = basis_values[1] if len(basis_values) > 1 else basis_values[0]
            components["REGISTRATION_INDIVIDUAL"] = reg_ind
            components["REGISTRATION_CORPORATE"] = reg_corp
            attrs = _variant_attributes(variant_name, fuel_hint=fuel_hint, electric=electric)
            signature = (_q2(onroad_ind), _q2(onroad_corp), tuple(sorted((k, _q2(v)) for k, v in components.items())))
            key = (_slug_model(model_name), variant_name.upper(), attrs["fuel"] or "", attrs["transmission"] or "",
                   attrs["drive"] or "", attrs["seater"] or "", "STANDARD")
            if key in seen:
                if seen[key] != signature:
                    result.errors.append(
                        f"sheet '{sheet.title}' row {idx} ({model_name} / {variant_name}): conflicting prices for the same variant"
                    )
                else:
                    result.meta["duplicateRowsCollapsed"] = result.meta.get("duplicateRowsCollapsed", 0) + 1
                continue
            seen[key] = signature
            result.price_rows.append(
                PriceRow(
                    row_no=idx,
                    category="BEV" if attrs["fuel"] == "ELECTRIC" else "ICE" if attrs["fuel"] else "UNSPECIFIED",
                    model_name=model_name,
                    variant_name=variant_name,
                    trim=None,
                    fuel=attrs["fuel"],
                    transmission=attrs["transmission"],
                    drive=attrs["drive"],
                    seater=attrs["seater"],
                    components={k: _q2(v) for k, v in components.items()},
                    onroad_individual=_q2(onroad_ind),
                    onroad_corporate=_q2(onroad_corp),
                    registration_basis="STANDARD",
                    source_sheet=sheet.title,
                )
            )
    if not sheets_read:
        raise MasterParseError(
            "Price list not recognised: neither the consolidated layout (header got "
            f"{header_seen}) nor a dealer per-model sheet (Model & Variant header with column letters)."
        )
    result.meta["sheets"] = sheets_read
    result.meta["models"] = sorted({r.model_name for r in result.price_rows})
    result.meta["categoryCounts"] = _counts(r.category for r in result.price_rows)
    if not result.price_rows and not result.errors:
        raise MasterParseError("Price sheets contained no recognisable vehicle rows.")
    return result


# ── 2. consumer scheme (PDF) ────────────────────────────────────────────────────
_CASH_RE = re.compile(r"cash\s*discount\s*of\s*rs\.?\s*([\d,\s]+)", re.IGNORECASE)
_ACC_RE = re.compile(r"access(?:ories|ory)?\s*worth\s*of\s*rs\.?\s*([\d,\s]+)", re.IGNORECASE)
_EW_RE = re.compile(
    r"(\d)\s*(?:st|nd|rd|th)?\s*year\s*extended\s*warranty\s*rs\.?\s*([\d,\s]+)", re.IGNORECASE
)
_INS_RE = re.compile(r"insurance\s*(?:worth|of)?\s*(?:rs\.?)?\s*([\d,\s]+)", re.IGNORECASE)


def _pdf_pages(content: bytes) -> tuple[list[list[list[list[Any]]]], str]:
    try:
        import pdfplumber
    except ModuleNotFoundError as exc:  # pragma: no cover - dependency guard
        raise MasterParseError(
            "pdfplumber is required to read this PDF scheme document."
        ) from exc
    with pdfplumber.open(BytesIO(content)) as pdf:
        tables = [page.extract_tables() for page in pdf.pages]
        text = "\n".join((page.extract_text() or "") for page in pdf.pages)
    return tables, text


def parse_consumer_scheme(content: bytes) -> ParseResult:
    page_tables, full_text = _pdf_pages(content)
    return _parse_consumer_tables(page_tables, full_text)


# ── scheme wording: EV models, variant cells, model-cell qualifiers ────────────
# 2026-10-20 product owner rules (R4): EV and ICE are kept apart. The bulletin's
# BNDP factor is the GST gross-up of the M&M contribution: 1.05 is 5% GST, which
# only an electric vehicle carries (ICE rows read 1.18 or 1.40).
_EV_BNDP_FACTOR = Decimal("1.05")
_EV_MODEL_RE = re.compile(
    r"\bBE\s*6\b|\bBE6|\bXEV\s*9\s*[ES]\b|\bXUV\s*400\b|\b3XO\s*EV\b|\bEV\b|\bELECTRIC\b",
    re.IGNORECASE,
)


def is_ev_model_text(text: str | None) -> bool:
    """True when a model name is an electric vehicle by its words (R4)."""
    return bool(_EV_MODEL_RE.search(text or ""))


# A bulletin variant cell wraps long names over several lines. A line is a
# continuation of the one above when that one ends mid-phrase ("... Kwh -",
# "... 6 &") or when it starts with a word no variant name starts with (an
# emission norm, a body or edition word, a charger/battery unit).
_CONTINUATION_WORDS = frozenset({
    "N", "ADAS", "CHARGER", "KWH", "KW", "BLACK", "EDITION", "CBC", "STEERING", "TIPPER", "HOPPER",
    "TPR", "HPR", "MS", "PS", "REFRESH", "REEFER", "AC", "NEW", "INTERIORS", "MOCHA",
})
_TRANSMISSIONS = frozenset({"AT", "MT", "AS", "AMT", "DCT", "IVT", "CVT"})


def _is_continuation(previous: str, line: str) -> bool:
    if previous.rstrip().endswith(("-", "&", ",", "/", "(")):
        return True
    if line.lstrip().startswith(("-", "&", ")", "/", ",")):
        return True
    first = line.split()[0].upper().strip(",")
    return first in _CONTINUATION_WORDS or bool(re.match(r"^BS\d", first))


def _expand_slash(variant: str) -> list[str] | None:
    """One "A/B" written inside a name ("LXT D AT/MT 4WD", "AX7 DS AS DT/MT DT",
    "V4/V6", "BS6.2 - TPR/HPR") as the two names it stands for, or None.

    When the right side starts with a gearbox the slash offers a second
    gearbox: the right side's words replace the left side's gearbox word and
    as many words after it as the right side repeats ("AS DT/MT DT" is
    "AS DT" or "MT DT"). Otherwise one word is swapped for another."""
    match = re.search(r"(\S+)/(\S+)", variant)
    if match is None:
        return None
    left = variant[: match.end(1)].split()
    right = variant[match.start(2):].split()
    if right and right[0].upper() in _TRANSMISSIONS:
        positions = [i for i, token in enumerate(left) if token.upper() in _TRANSMISSIONS]
        if positions:
            at = positions[-1]
            parallel = len(left) - at
            first = " ".join(left + right[parallel:])
            second = " ".join(left[:at] + right)
            return [first, second]
    # a seater pair "6/7STR" -> "6 STR", "7 STR"
    seat = re.fullmatch(r"(\d+)", match.group(1)), re.fullmatch(r"(\d+)\s*(STR)", match.group(2), re.IGNORECASE)
    if seat[0] and seat[1]:
        head, tail = variant[: match.start(1)], variant[match.end(2):]
        return [f"{head}{seat[0].group(1)} STR{tail}", f"{head}{seat[1].group(1)} STR{tail}"]
    head, tail = variant[: match.start(1)], variant[match.end(2):]
    return [f"{head}{match.group(1)}{tail}", f"{head}{match.group(2)}{tail}"]


def _expand_ampersand(variant: str) -> list[str] | None:
    """"AT & MT" (two gearboxes) and "6 & 7STR" (two seaters) as two names.
    Any other "&" ("& Mocha Interiors", "LX & CBC") is part of the name."""
    gear = re.search(r"\b(AT|MT|AS|AMT|DCT|IVT|CVT)\s*&\s*(AT|MT|AS|AMT|DCT|IVT|CVT)\b", variant, re.IGNORECASE)
    if gear is not None:
        head, tail = variant[: gear.start()], variant[gear.end():]
        return [f"{head}{gear.group(1)}{tail}", f"{head}{gear.group(2)}{tail}"]
    seat = re.search(r"\b(\d)\s*&\s*(\d)\s*STR\b", variant, re.IGNORECASE)
    if seat is not None:
        head, tail = variant[: seat.start()], variant[seat.end():]
        return [f"{head}{seat.group(1)} STR{tail}", f"{head}{seat.group(2)} STR{tail}"]
    return None


def split_variant_cell(cell: Any) -> list[str]:
    """A bulletin's variant cell as the list of variant names it names.

    2026-10-20 product owner rules: wrapped lines are joined back together,
    "A / B" lists and "AT & MT", "AT/MT", "MT/AT", "AS/MT" gearbox pairs are
    expanded into one name each, so a name can be matched word for word with
    the price list. A name still not found there is reported for an admin to
    map; it is never applied to the whole model."""
    # A note in brackets ("(Discontinued Model)") is not part of a name; a
    # bracket that is ("M(O)", "EDGE(O)") has no space in it.
    text = re.sub(r"\((?=[^()]*\s)[^()]*\)", " ", str(cell or ""))
    lines = [re.sub(r"[ \t]+", " ", part).strip() for part in text.split("\n")]
    lines = [line for line in lines if line]
    joined: list[str] = []
    for line in lines:
        if joined and _is_continuation(joined[-1], line):
            joined[-1] = re.sub(r"\s+", " ", f"{joined[-1]} {line}").strip()
        else:
            joined.append(line)

    names: list[str] = []
    for line in joined:
        parts = [p.strip() for p in re.split(r"\s+/\s*|\s*/\s+", line) if p.strip()]
        if len(parts) > 1 and all(len(p.split()) == 1 for p in parts[1:]) and len(parts[0].split()) > 1:
            # "... - Tipper / Hopper": the last word has alternatives
            stem = parts[0].rsplit(" ", 1)[0]
            parts = [parts[0], *(f"{stem} {p}" for p in parts[1:])]
        names.extend(parts)

    expanded: list[str] = []
    pending = list(names)
    guard = 0
    while pending and guard < 200:
        guard += 1
        name = re.sub(r"\s+", " ", pending.pop(0)).strip(" -")
        if not name:
            continue
        alternatives = _expand_ampersand(name) or _expand_slash(name)
        if alternatives:
            pending[:0] = alternatives
            continue
        if name not in expanded:
            expanded.append(name)
    return expanded


def _split_outside_parentheses(text: str, separators: str) -> list[str]:
    parts: list[str] = []
    depth = 0
    current = ""
    for char in text:
        if char == "(":
            depth += 1
        elif char == ")":
            depth = max(0, depth - 1)
        if depth == 0 and char in separators:
            parts.append(current)
            current = ""
            continue
        current += char
    parts.append(current)
    return [p.strip() for p in parts if p.strip()]


def _balanced_groups(text: str) -> tuple[str, list[str]]:
    """Text with every balanced (...) group removed, and the groups' contents.
    "XUV 3XO (Excluding EDGE & REVX M(O) PM AT )" -> ("XUV 3XO", ["Excluding
    EDGE & REVX M(O) PM AT"]): the inner "(O)" stays part of the variant."""
    out = ""
    groups: list[str] = []
    depth = 0
    current = ""
    for char in text:
        if char == "(":
            if depth > 0:
                current += char
            depth += 1
            continue
        if char == ")" and depth > 0:
            depth -= 1
            if depth == 0:
                groups.append(current.strip())
                current = ""
            else:
                current += char
            continue
        if depth > 0:
            current += char
        else:
            out += char
    if depth > 0 and current.strip():  # unbalanced: keep the tail as a group
        groups.append(current.strip())
    return re.sub(r"\s+", " ", out).strip(), groups


def _variant_list(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\s*&\s*|\s*,\s*", text) if p.strip()]


def old_vehicle_requirement(text: str | None) -> dict[str, str] | None:
    """An exchange / welcome row's old-vehicle condition, or None when any old
    vehicle qualifies ("Any Personal + Commercial Vehicle", "Not Applicable").
    "Tata Ace Family" -> {"brand": "TATA", "model": "ACE"}; "Old Maruti
    Vehicle" -> {"brand": "MARUTI"}."""
    words = re.sub(r"[^A-Za-z0-9 ]+", " ", text or "").upper().split()
    words = [w for w in words if w not in {"OLD", "VEHICLE", "VEHICLES", "FAMILY", "THE", "A", "AN"}]
    if not words or words[0] in {"ANY", "NOT", "NA", "ALL"}:
        return None
    requirement = {"brand": words[0]}
    if len(words) > 1:
        requirement["model"] = " ".join(words[1:])
    return requirement


def parse_model_qualifiers(part: str) -> tuple[str, dict[str, Any]]:
    """One model named in a scheme's model cell, and what its wording adds.

    2026-10-20 product owner rules: the qualifiers are eligibility, not noise.
      "(Not on 3XO EV)", "(Not on Star Edition)", "(Excluding EDGE & REVX
      M(O) PM AT)"          -> excludedVariantTexts
      "against Old Maruti Vehicle"   -> requiredOldVehicle {"brand": "MARUTI"}
      "for All Variants", "(All Variants)" -> allVariants
      "(Including REVX)", "including Refresh", "(Incl Thar OG Variants)" -> includes
    The parentheses are stripped balanced, so "M(O)" inside one survives."""
    text = re.sub(r"\*+", " ", part or "")
    base, groups = _balanced_groups(text)
    qualifiers: dict[str, Any] = {}

    def note_group(content: str) -> None:
        low = content.lower().strip()
        excl = re.match(r"^(?:not\s+on|excluding|except|excl\.?)\s+(.+)$", low)
        if excl:
            original = content.strip()[len(content.strip()) - len(excl.group(1)):]
            qualifiers.setdefault("excludedVariantTexts", []).extend(_variant_list(original))
            return
        if re.match(r"^(?:for\s+)?all\s+variants?$", low):
            qualifiers["allVariants"] = True
            return
        inc = re.match(r"^(?:including|incl\.?)\s+(.+)$", low)
        if inc:
            qualifiers.setdefault("includes", []).append(content.strip()[len(content.strip()) - len(inc.group(1)):])
            return
        qualifiers.setdefault("notes", []).append(content.strip())

    for group in groups:
        note_group(group)

    against = re.search(r"\bagainst\s+(.+)$", base, re.IGNORECASE)
    if against:
        requirement = old_vehicle_requirement(against.group(1))
        if requirement:
            qualifiers["requiredOldVehicle"] = requirement
        base = base[: against.start()].strip()
    all_variants = re.search(r"\bfor\s+all\s+variants?\b", base, re.IGNORECASE)
    if all_variants:
        qualifiers["allVariants"] = True
        base = (base[: all_variants.start()] + base[all_variants.end():]).strip()
    including = re.search(r"\b(?:including|incl\.?)\s+(.+)$", base, re.IGNORECASE)
    if including:
        qualifiers.setdefault("includes", []).append(including.group(1).strip())
        base = base[: including.start()].strip()
    return re.sub(r"\s+", " ", base).strip(" -"), qualifiers


def variant_key(name: str | None) -> str:
    """A variant name in the form two spellings of the same name share:
    upper case, punctuation as spaces, "7STR" as "7 STR", and without the
    price list's "(Private)" / "(Commercial)" registration suffix (the scheme
    names the variant, not its registration basis)."""
    text = re.sub(r"\((?:PRIVATE|COMMERCIAL)\)", " ", (name or "").upper())
    text = re.sub(r"(\d)\s*STR\b", r"\1 STR", text)
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def variant_excluded(variant_name: str | None, excluded_texts: list[str] | None) -> bool:
    """True when a scheme's exclusion ("EDGE", "REVX M(O) PM AT", "Star
    Edition") names this variant: every word of the exclusion is a word of
    the variant's name."""
    words = set(variant_key(variant_name).split())
    for excluded in excluded_texts or []:
        wanted = set(variant_key(excluded).split())
        if wanted and wanted <= words:
            return True
    return False


def split_model_cell(cell: str | None) -> list[str]:
    """A scheme's model cell ("Pick up/Camper/MAXX City", "XUV700 | Scorpio-N")
    as the models it names; a separator inside parentheses does not split."""
    return _split_outside_parentheses(cell or "", "|/")


# The September BEV "bulletin" is an e-mail that carries the August schemes
# forward for some EVs and declares none for others. It has no scheme table:
# the August figures it refers to are not in it. 2026-10-20: it is refused
# with what it says, never turned into schemes by guesswork.
_CARRY_FORWARD_RE = re.compile(r"continu\w*\s+the\s+same\s+consumer\s+schemes?\s+of\s+([A-Za-z]+[’'`]?\s*\d{0,4})",
                               re.IGNORECASE)


def _carry_forward_declaration(full_text: str) -> str | None:
    match = _CARRY_FORWARD_RE.search(full_text or "")
    if match is None:
        return None
    lines = [line.strip() for line in full_text.splitlines()]

    def block(start_re: str) -> list[str]:
        out: list[str] = []
        for i, line in enumerate(lines):
            if re.search(start_re, line, re.IGNORECASE):
                for following in lines[i + 1:]:
                    if not following or re.match(r"^(pls|please|regards|there are|note)\b", following, re.IGNORECASE):
                        break
                    out.append(following)
                break
        return out

    continued = block(r"no\s+changes\s+in\s+schemes?\s+of")
    none = block(r"no\s+consumer\s+schemes?\s+in")
    return (
        "This file is an e-mail, not a consumer scheme bulletin: it has no scheme table to load. "
        f"It says the consumer schemes of {match.group(1).strip()} continue unchanged for: "
        f"{', '.join(continued) or 'the models it names'}; and that there are no consumer schemes for: "
        f"{', '.join(none) or 'none named'}. Nothing was loaded. To apply it, upload that month's BEV "
        "scheme bulletin with this month's effective date; models with no scheme need no upload."
    )


def _parse_consumer_tables(page_tables: list[list[list[list[Any]]]], full_text: str) -> ParseResult:
    result = ParseResult(kind="CONSUMER_SCHEME")
    result.effective_from_hint = _date_hint(full_text)

    row_no = 0
    found_header = False
    previous_page_product: str | None = None
    previous_page_row: DiscountRow | None = None
    previous_page_signature: tuple[Any, ...] | None = None
    last_row_on_page: DiscountRow | None = None
    last_signature_on_page: tuple[Any, ...] | None = None
    for page_index, tables in enumerate(page_tables, start=1):
        # 2026-10-20 product owner rules: a product cell is never carried over
        # a page break. The September bulletin's XUV400 block starts page 2
        # with its product cell empty; carried, it read as XUV 3XO and a
        # Rs 2,50,000 EV offer reached petrol and diesel XUV 3XOs.
        carry_product: str | None = None
        last_on_page: str | None = None
        for table in tables:
            data_rows = [
                raw for raw in table
                if raw and len(raw) >= 10 and _money(raw[6]) is not None and "₹" in str(raw[6] or "")
            ]
            next_product: dict[int, str | None] = {}
            following: str | None = None
            for index in range(len(data_rows) - 1, -1, -1):
                next_product[index] = following
                own = _text(data_rows[index][0])
                if own:
                    following = own
            data_index = -1
            for raw in table:
                if not raw or len(raw) < 10:
                    continue
                cells = [(_text(c) or "") for c in raw[:10]]
                if cells[0].lower().startswith("product") and "variant" in cells[1].lower():
                    found_header = True
                    continue
                total = _money(raw[6])
                if total is None or "₹" not in str(raw[6] or ""):
                    continue  # not a data row (title / legend)
                data_index += 1
                factor = _money(raw[3])
                is_ev_row = factor is not None and abs(factor - _EV_BNDP_FACTOR) < Decimal("0.001")
                signature = tuple(_money(c) for c in raw[2:9])
                product = cells[0] or carry_product
                if not product and previous_page_row is not None and signature == previous_page_signature:
                    # The page break cut one bulletin row in two (October's
                    # VEERO "V2 / V2 CNG / V4 CNG"): same figures, the rest of
                    # its variants. It is that row, not a new one.
                    for name in split_variant_cell(raw[1]):
                        if name not in previous_page_row.variant_texts:
                            previous_page_row.variant_texts.append(name)
                    carry_product = previous_page_row.model_alias
                    continue
                row_no += 1
                if not product:
                    product, problem = _product_at_page_top(
                        previous=previous_page_product, following=next_product.get(data_index), ev_row=is_ev_row,
                    )
                    if problem:
                        result.errors.append(f"consumer row {row_no} (page {page_index}): {problem}")
                        continue
                carry_product = product
                last_on_page = product
                if is_ev_row and not is_ev_model_text(product):
                    result.errors.append(
                        f"consumer row {row_no} ({product}): the BNDP factor 1.05 is the EV rate (5% GST) "
                        "but the model is not an electric vehicle"
                    )
                    continue
                if factor is not None and factor > 0 and not is_ev_row and is_ev_model_text(product):
                    result.errors.append(
                        f"consumer row {row_no} ({product}): an electric vehicle with BNDP factor {factor}, "
                        "not the EV rate 1.05"
                    )
                    continue
                cash = _money(raw[7]) or Decimal(0)
                other = _money(raw[8]) or Decimal(0)
                if abs((cash + other) - total) > _MONEY_TOLERANCE:
                    result.errors.append(
                        f"consumer row {row_no} ({product}): cash {_q2(cash)} + other "
                        f"{_q2(other)} != total consumer offer {_q2(total)}"
                    )
                    continue

                variant_texts = split_variant_cell(raw[1])
                description = re.sub(r"\s+", " ", str(raw[9] or "")).strip()
                benefits, warnings = _consumer_benefits(cash, other, description)
                result.discount_rows.append(
                    DiscountRow(
                        row_no=row_no,
                        scheme_category="CONSUMER",
                        model_alias=product or "",
                        variant_texts=variant_texts,
                        benefits=benefits,
                        total_customer_offer=_q2(total),
                        config={
                            "description": description,
                            "cashDiscount": str(_q2(cash)),
                            "otherSchemes": str(_q2(other)),
                            "mAndMContribution": str(_q2(_money(raw[4]) or Decimal(0))),
                            "dealerContribution": str(_q2(_money(raw[5]) or Decimal(0))),
                            "bndpFactor": str(factor) if factor is not None else None,
                            "vehicleCategory": "EV" if is_ev_row else "ICE",
                            "variantCell": str(raw[1] or ""),
                        },
                        warnings=warnings,
                    )
                )
                last_row_on_page = result.discount_rows[-1]
                last_signature_on_page = signature
        if last_on_page:
            previous_page_product = last_on_page
        previous_page_row, previous_page_signature = last_row_on_page, last_signature_on_page
        last_row_on_page, last_signature_on_page = None, None
    if not found_header:
        declaration = _carry_forward_declaration(full_text)
        if declaration:
            raise MasterParseError(declaration)
        raise MasterParseError("Not a consumer scheme bulletin (header row not found).")
    result.meta["schemeRows"] = len(result.discount_rows)
    return result


def _product_at_page_top(
    *, previous: str | None, following: str | None, ev_row: bool,
) -> tuple[str | None, str | None]:
    """The product of a row that opens a page with its product cell empty.

    The merged product cell either continues from the previous page or is
    the block the next named row belongs to. The BNDP factor decides between
    them; when it cannot, the row is refused rather than guessed."""
    candidates = [c for c in dict.fromkeys((previous, following)) if c]
    if ev_row:
        fitting = [c for c in candidates if is_ev_model_text(c)]
    else:
        fitting = [c for c in candidates if not is_ev_model_text(c)]
    if len(fitting) == 1:
        return fitting[0], None
    return None, (
        "the product cell is empty at the top of the page and the row cannot be placed "
        f"(previous page: {previous or '-'}, next row: {following or '-'}); fix the bulletin or enter it by hand"
    )


def _consumer_benefits(
    cash: Decimal, other: Decimal, description: str
) -> tuple[list[tuple[str, Decimal]], list[str]]:
    benefits: list[tuple[str, Decimal]] = []
    warnings: list[str] = []
    if cash > 0:
        benefits.append(("CASH_DISCOUNT", _q2(cash)))
        d_cash = _amount(_CASH_RE.search(description))
        if d_cash is not None and abs(d_cash - cash) > _MONEY_TOLERANCE:
            warnings.append(
                f"description cash {d_cash} disagrees with cash column {_q2(cash)}"
            )
    if other > 0:
        parts: list[tuple[str, Decimal]] = []
        acc = _amount(_ACC_RE.search(description))
        if acc is not None:
            parts.append(("ACCESSORIES_KIT", acc))
        ew_match = _EW_RE.search(description)
        if ew_match is not None:
            ew_key = (
                "EXT_WARRANTY_4TH_YR"
                if ew_match.group(1) == "4"
                else "EXT_WARRANTY_4TH_5TH_YR"
            )
            parts.append((ew_key, _amount_str(ew_match.group(2))))
        ins = _amount(_INS_RE.search(description)) if "insurance" in description.lower() else None
        if ins is not None:
            parts.append(("INSURANCE", ins))
        split_total = sum((amount for _, amount in parts), Decimal(0))
        if parts and abs(split_total - other) <= _MONEY_TOLERANCE:
            benefits.extend((key, _q2(amount)) for key, amount in parts)
        else:
            benefits.append(("OTHER_SCHEME", _q2(other)))
            if parts:
                warnings.append(
                    f"could not split 'other schemes' {_q2(other)} from description "
                    f"(parsed {_q2(split_total)}); kept as OTHER_SCHEME"
                )
    return benefits, warnings


def _amount(match: re.Match[str] | None) -> Decimal | None:
    if match is None:
        return None
    return _amount_str(match.group(1))


def _amount_str(raw: str) -> Decimal:
    return Decimal(re.sub(r"[,\s]", "", raw))


# ── 3. exchange / scrappage / welcome (PDF) ─────────────────────────────────────
def parse_exchange_scheme(content: bytes) -> ParseResult:
    result = ParseResult(kind="EXCHANGE_SCHEME")
    page_tables, full_text = _pdf_pages(content)
    result.effective_from_hint = _date_hint(
        full_text.replace("valid from", "w.e.f.")
    )

    row_no = 0
    recognised = 0
    for page_index, tables in enumerate(page_tables, start=1):
        section = _exchange_section(full_text, page_index)
        for table in tables:
            if not table or len(table) < 2:
                continue
            header = [(_text(c) or "").lower() for c in table[0]]
            if "brand" not in " ".join(header):
                continue
            width = len(table[0])
            if width == 5:
                brand_i, a_i, b_i, total_i, note_i = 0, 1, 2, 3, 4
                old_i = scheme_i = None
            elif width >= 7:
                brand_i, old_i, scheme_i, a_i, b_i, total_i, note_i = 0, 1, 2, 3, 4, 5, 6
            else:
                continue
            recognised += 1
            for raw in table[1:]:
                brand = _text(raw[brand_i]) if len(raw) > brand_i else None
                total = _money(raw[total_i]) if len(raw) > total_i else None
                if not brand or total is None:
                    continue
                a_val = _money(raw[a_i]) or Decimal(0)
                b_val = _money(raw[b_i]) or Decimal(0)
                if abs((a_val + b_val) - total) > _MONEY_TOLERANCE:
                    result.errors.append(
                        f"{section} row '{brand}': A {_q2(a_val)} + B {_q2(b_val)} "
                        f"!= total {_q2(total)}"
                    )
                    continue
                row_no += 1
                benefit_key, scheme_category = _EXCHANGE_BENEFIT[section]
                config: dict[str, Any] = {
                    "section": section,
                    "mAndMContribution": str(_q2(a_val)),
                    "dealerContribution": str(_q2(b_val)),
                    "creditNoteWithoutGst": str(
                        _q2(_money(raw[note_i]) or Decimal(0))
                    )
                    if len(raw) > note_i
                    else None,
                }
                if old_i is not None and len(raw) > old_i and _text(raw[old_i]):
                    config["oldVehicleModel"] = _text(raw[old_i])
                    # 2026-10-20 product owner rules: the old-vehicle column is
                    # a condition ("Tata Ace Family"), kept structured so the
                    # benefit applies only against that old vehicle.
                    requirement = old_vehicle_requirement(config["oldVehicleModel"])
                    if requirement:
                        config["requiredOldVehicle"] = requirement
                if scheme_i is not None and len(raw) > scheme_i and _text(raw[scheme_i]):
                    config["schemeType"] = _text(raw[scheme_i])
                result.discount_rows.append(
                    DiscountRow(
                        row_no=row_no,
                        scheme_category=scheme_category,
                        model_alias=brand,
                        variant_texts=[],
                        benefits=[(benefit_key, _q2(total))],
                        total_customer_offer=_q2(total),
                        config=config,
                    )
                )
    if not recognised:
        raise MasterParseError("No exchange/scrappage ready-reckoner tables found.")
    result.meta["schemeRows"] = len(result.discount_rows)
    result.meta["sections"] = sorted({r.config.get("section") for r in result.discount_rows})
    return result


_EXCHANGE_BENEFIT = {
    "EXCHANGE_PERSONAL": ("EXCHANGE_BONUS", "EXCHANGE"),
    "EXCHANGE_COMMERCIAL": ("EXCHANGE_BONUS", "EXCHANGE"),
    "WELCOME_BONUS": ("WELCOME_BONUS", "WELCOME"),
    "SCRAPPAGE_DEALER": ("SCRAPPAGE_BONUS_DEALER", "SCRAPPAGE"),
    "SCRAPPAGE_COD": ("SCRAPPAGE_BONUS_COD", "SCRAPPAGE"),
}


def _exchange_section(full_text: str, page_index: int) -> str:
    lines = full_text.splitlines()
    # crude page->section map keyed off the section headings present in the deck
    joined = full_text.lower()
    markers = [
        ("exchange scheme – personal brands", "EXCHANGE_PERSONAL"),
        ("exchange scheme – commercial brands", "EXCHANGE_COMMERCIAL"),
        ("welcome bonus", "WELCOME_BONUS"),
        ("customer scrapped the vehicle through m&m", "SCRAPPAGE_DEALER"),
        ("customer approached m&m dealer with cod", "SCRAPPAGE_COD"),
    ]
    order: list[str] = []
    for line in lines:
        low = line.strip().lower()
        for needle, tag in markers:
            if needle in low and (not order or order[-1] != tag):
                order.append(tag)
    del joined
    if page_index - 2 < len(order) and page_index >= 2:
        return order[page_index - 2]
    return order[-1] if order else "EXCHANGE_PERSONAL"


# ── 4. corporate privilege policy (xlsx) ────────────────────────────────────────
_CAT_ROW_LABELS = {
    "cat-b": "B",
    "cat-a": "A",
    "cat-f": "F",
    "cat-y (premium)": "Y",
    "cat-y(premium)": "Y",
    "cat-z (signature)": "Z",
    "cat-z(signature)": "Z",
}
_COMPANY_BLOCK_CATEGORY = {
    '"z" signature': "Z",
    '"y" premium': "Y",
    '"f" focus': "F",
    '"a"': "A",
    '"b"': "B",
}


def parse_corporate_policy(content: bytes) -> ParseResult:
    result = ParseResult(kind="CORPORATE_POLICY")
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    names = {n.strip().lower(): n for n in workbook.sheetnames}
    policy_name = names.get("corporate policy")
    companies_name = names.get("companies list")
    if not policy_name or not companies_name:
        raise MasterParseError(
            "Corporate policy workbook must have 'Corporate Policy' and 'Companies List' sheets."
        )

    _parse_corporate_matrix(workbook[policy_name], result)
    _parse_company_list(workbook[companies_name], result)

    hint = " ".join(
        str(c)
        for r in list(workbook[policy_name].iter_rows(values_only=True))[:2]
        for c in r
        if isinstance(c, str)
    )
    result.effective_from_hint = _date_hint(hint.replace("From", "w.e.f."))
    result.meta["benefitRows"] = len(result.corporate_benefits)
    result.meta["companies"] = len(result.corporate_companies)
    return result


def _parse_corporate_matrix(sheet: Any, result: ParseResult) -> None:
    rows = list(sheet.iter_rows(values_only=True))
    block: list[tuple[str, int]] = []  # (brand_alias, m&m column index)
    for raw in rows:
        cells = [_text(c) for c in raw]
        label = (cells[1] or "").lower() if len(cells) > 1 else ""
        # a brand-header row: many 3-col groups of M&M/Dealer/Total after col B
        if label == "corporate category":
            block = []
            for col in range(2, len(cells)):
                name = cells[col]
                if name and name.lower() not in {"m&m", "dealer", "total"}:
                    block.append((name, col))
            continue
        category = _CAT_ROW_LABELS.get(label.replace(" ", "")) or _CAT_ROW_LABELS.get(label)
        if not category or not block:
            continue
        for brand_alias, col in block:
            m_and_m = _money(raw[col]) if len(raw) > col else None
            dealer = _money(raw[col + 1]) if len(raw) > col + 1 else None
            total = _money(raw[col + 2]) if len(raw) > col + 2 else None
            if None in (m_and_m, dealer, total):
                continue
            if total == 0:
                continue
            if abs((m_and_m + dealer) - total) > _MONEY_TOLERANCE:
                result.errors.append(
                    f"corporate {category} / {brand_alias}: M&M {m_and_m} + dealer "
                    f"{dealer} != total {total}"
                )
                continue
            result.corporate_benefits.append(
                CorporateBenefitRow(
                    row_no=len(result.corporate_benefits) + 1,
                    privilege_category=category,
                    brand_alias=brand_alias,
                    m_and_m=_q2(m_and_m),
                    dealer=_q2(dealer),
                    total=_q2(total),
                )
            )


def _parse_company_list(sheet: Any, result: ParseResult) -> None:
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 3:
        return
    # row 1 holds block titles; row 2 holds "S. No./Corporate Type/Description/Code"
    header = [(_text(c) or "").lower() for c in rows[1]]
    blocks: list[tuple[str, int]] = []  # (category, code_col_index)
    for col, value in enumerate(header):
        if value == "corporate code":
            # walk left to the block title on row 0
            title = None
            for back in range(col, -1, -1):
                cand = _text(rows[0][back]) if back < len(rows[0]) else None
                if cand:
                    title = cand.lower()
                    break
            category = _COMPANY_BLOCK_CATEGORY.get((title or "").strip())
            if category:
                blocks.append((category, col))
    seen: set[str] = set()
    for raw in rows[2:]:
        for category, code_col in blocks:
            code = raw[code_col] if len(raw) > code_col else None
            name = _text(raw[code_col - 1]) if code_col >= 1 and len(raw) > code_col - 1 else None
            ctype = _text(raw[code_col - 2]) if code_col >= 2 and len(raw) > code_col - 2 else None
            if code is None or not name:
                continue
            code_str = str(int(code)) if isinstance(code, float) and code.is_integer() else str(code).strip()
            if not code_str or code_str in seen:
                continue
            seen.add(code_str)
            result.corporate_companies.append(
                CorporateCompany(
                    corporate_code=code_str,
                    corporate_name=name,
                    corporate_type=ctype,
                    privilege_category=category,
                )
            )


# ── 5. the dealer's discount grid (xlsx) ────────────────────────────────────────
# One sheet: an "Effective ..." line, a header (Model | Booking Protection |
# Agreed Buffer | Insurance OD % | Out of Territory), one row per model (a cell
# may name several models: "PICKUP, MAXX, MAXX HD"), then a Parameter | Notes
# block with the policy wording. Values as written: "30 days", "Nil", "Out of
# scope", 0.6, "Additional 3K".
_GRID_HEADER = {
    "model": ("MODEL",),
    "booking_protection": ("BOOKING PROTECTION",),
    "agreed_buffer": ("AGREED BUFFER", "BUFFER"),
    "insurance_od": ("INSURANCE OD", "OD %"),
    "out_of_territory": ("OUT OF TERRITORY",),
}
_OUT_OF_SCOPE = {"OUT OF SCOPE", "NOT APPLICABLE", "N/A", "NA"}
_NIL = {"NIL", "NONE", "-", "0", "ZERO"}


def _grid_days(value: Any) -> tuple[int | None, str | None]:
    text = (_text(value) or "").upper()
    if not text or text in _OUT_OF_SCOPE:
        return None, None
    if text in _NIL:
        return 0, None
    match = re.search(r"(\d+)", text)
    if match:
        return int(match.group(1)), None
    return None, f"unreadable days '{_text(value)}'"


def _grid_amount(value: Any) -> tuple[Decimal | None, str | None]:
    if isinstance(value, (int, float, Decimal)):
        return _q2(Decimal(str(value))), None
    text = (_text(value) or "").upper()
    if not text or text in _OUT_OF_SCOPE:
        return None, None
    if text in _NIL:
        return Decimal("0.00"), None
    match = re.search(r"(\d+(?:\.\d+)?)\s*(K|L|LAKH|LAC)?", text.replace(",", ""))
    if not match:
        return None, f"unreadable amount '{_text(value)}'"
    amount = Decimal(match.group(1))
    unit = match.group(2) or ""
    if unit == "K":
        amount *= 1000
    elif unit:
        amount *= 100000
    return _q2(amount), None


def _grid_percent(value: Any) -> tuple[Decimal | None, str | None]:
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
        return _q2(number * 100 if number <= 1 else number), None
    text = (_text(value) or "").upper()
    if not text or text in _OUT_OF_SCOPE:
        return None, None
    if text in _NIL:
        return Decimal("0.00"), None
    match = re.search(r"(\d+(?:\.\d+)?)", text)
    if not match:
        return None, f"unreadable percentage '{_text(value)}'"
    number = Decimal(match.group(1))
    return _q2(number * 100 if number <= 1 and "%" not in text else number), None


def parse_discount_grid(content: bytes, *, filename: str | None = None) -> ParseResult:
    result = ParseResult(kind="DISCOUNT_GRID", meta={"layout": "DEALER_DISCOUNT_GRID"})
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    header_row = next(
        (i for i, raw in enumerate(rows)
         if any(_text(c) and _text(c).upper() == "MODEL" for c in raw)
         and any(_text(c) and "BOOKING PROTECTION" in _text(c).upper() for c in raw)),
        None,
    )
    if header_row is None:
        raise MasterParseError(
            "Discount grid not recognised: no 'Model | Booking Protection | ...' header row."
        )
    columns: dict[str, int] = {}
    tier_columns: dict[str, int] = {}
    for col, cell in enumerate(rows[header_row]):
        label = (_text(cell) or "").upper()
        tier = _grid_tier(label)
        if tier and tier not in tier_columns:
            tier_columns[tier] = col
            continue
        for key, needles in _GRID_HEADER.items():
            if key not in columns and label and any(n in label for n in needles):
                columns[key] = col
    # 2026-10-20: the October grid replaces the single agreed buffer with
    # three tiers and moves the insurance OD caps into its notes.
    if "agreed_buffer" not in columns and "BASE" in tier_columns:
        columns["agreed_buffer"] = tier_columns["BASE"]
        result.meta["layout"] = "DEALER_DISCOUNT_GRID_TIERED"
    required = ["model", "booking_protection", "agreed_buffer", "out_of_territory"]
    if not tier_columns:
        required.append("insurance_od")
    for key in required:
        if key not in columns:
            raise MasterParseError(f"Discount grid header lacks the {key.replace('_', ' ')} column.")

    blob = " ".join(str(c) for raw in rows[:header_row] for c in raw if isinstance(c, str))
    result.effective_from_hint = _loose_date_hint(blob)
    if result.effective_from_hint is not None:
        result.meta["effectiveFromSource"] = "SHEET"
    elif filename and _loose_date_hint(filename):
        result.effective_from_hint = _loose_date_hint(filename)
        result.meta["effectiveFromSource"] = "FILENAME"

    in_parameters = False
    od_caps: dict[str, Decimal] = {}
    for idx, raw in enumerate(rows[header_row + 1:], start=header_row + 2):
        cell = _cell_reader(raw)
        first = _text(cell(columns["model"]))
        if not first:
            continue
        if first.upper() == "PARAMETER":
            in_parameters = True
            continue
        if in_parameters:
            raw_note = next((v for v in raw[columns["model"] + 1:] if _text(v)), None)
            note = _text(raw_note)
            result.grid_parameters.append({"parameter": first, "note": note or ""})
            if re.search(r"\bOD\b", first, re.IGNORECASE):
                od_caps.update(_od_caps(str(raw_note or "")))
            continue
        value_cols = [columns["booking_protection"], columns["agreed_buffer"], *tier_columns.values(),
                      columns["out_of_territory"]]
        if "insurance_od" in columns:
            value_cols.append(columns["insurance_od"])
        if all(not _text(cell(c)) for c in value_cols):
            continue  # a heading inside the sheet ("Notes"), not a model row
        values = [cell(columns["booking_protection"]), cell(columns["agreed_buffer"]),
                  cell(columns["insurance_od"]) if "insurance_od" in columns else None,
                  cell(columns["out_of_territory"])]
        in_scope = not all((_text(v) or "").upper() in _OUT_OF_SCOPE for v in values if v is not None)
        days, e1 = _grid_days(values[0])
        buffer, e2 = _grid_amount(values[1])
        od, e3 = _grid_percent(values[2])
        territory, e4 = _grid_amount(values[3])
        tiers: dict[str, Decimal] = {}
        tier_problems: list[str] = []
        for tier, col in tier_columns.items():
            amount, problem = _grid_amount(cell(col))
            if problem:
                tier_problems.append(problem)
            elif amount is not None:
                tiers[tier] = amount
        problems = [e for e in (e1, e2, e3, e4, *tier_problems) if e]
        if problems:
            result.errors.append(f"grid row {idx} ({first}): " + "; ".join(problems))
            continue
        aliases = [a.strip() for a in re.split(r"[,&/|]|\band\b", first, flags=re.IGNORECASE) if a.strip()]
        raw_values = {
            "bookingProtection": _text(values[0]),
            "agreedBuffer": _text(values[1]),
            "insuranceOd": _text(values[2]),
            "outOfTerritory": _text(values[3]),
        }
        if tiers:
            raw_values["tiers"] = {tier: str(amount) for tier, amount in tiers.items()}
        result.grid_rows.append(
            GridRow(
                row_no=idx,
                model_alias=first,
                model_aliases=aliases or [first],
                in_scope=in_scope,
                booking_protection_days=days,
                agreed_buffer_amount=buffer,
                insurance_od_percent=od,
                out_of_territory_amount=territory,
                raw=raw_values,
                tiers=tiers,
            )
        )
    # The insurance OD cap the notes give per vehicle class (October 2026:
    # private car 65%, goods carrier 65%, passenger carrier 50%, 3+3 85%):
    # a row without its own OD column takes its class's cap.
    if od_caps:
        result.meta["odCapsPercent"] = {k: str(v) for k, v in od_caps.items()}
        for grow in result.grid_rows:
            if grow.insurance_od_percent is None and grow.in_scope:
                vehicle_class = "GCV" if _CV_GRID_RE.search(grow.model_alias) else "PRIVATE_CAR"
                grow.insurance_od_percent = od_caps.get(vehicle_class)
                grow.raw["odCapsPercent"] = result.meta["odCapsPercent"]
                grow.raw["odCapClass"] = vehicle_class
    if not result.grid_rows and not result.errors:
        raise MasterParseError("Discount grid contained no model rows.")
    result.meta["gridRows"] = len(result.grid_rows)
    result.meta["parameters"] = len(result.grid_parameters)
    result.meta["models"] = [r.model_alias for r in result.grid_rows]
    if tier_columns:
        result.meta["tiers"] = list(tier_columns)
    return result


_CV_GRID_RE = re.compile(r"PICK\s*-?\s*UP|MAXX|SUPRO|CAMPER|VEERO", re.IGNORECASE)


def _grid_tier(label: str) -> str | None:
    """The October grid's buffer tier a header names, or None."""
    if not label.startswith("MY") and "INSURANCE" not in label:
        return None
    if "WITHOUT" in label:
        return "BASE"
    if "EXTENDED WARRANTY" in label or re.search(r"\bEW\b", label):
        return "WITH_EW_INSURANCE"
    if "WITH INSURANCE" in label:
        return "WITH_INSURANCE"
    return None


def _od_caps(note: str) -> dict[str, Decimal]:
    caps: dict[str, Decimal] = {}
    for line in re.split(r"\n|(?<=%)\s+", note):
        match = re.search(r"^(.*?):.*?(\d+(?:\.\d+)?)\s*%", line.strip())
        if not match:
            continue
        label = match.group(1).lower()
        if "3+3" in label:
            key = "THREE_PLUS_THREE"
        elif "gcv" in label or "goods" in label:
            key = "GCV"
        elif "pcv" in label or "passenger" in label:
            key = "PCV"
        elif "private" in label or "car" in label:
            key = "PRIVATE_CAR"
        else:
            continue
        caps[key] = _q2(Decimal(match.group(2)))
    return caps


# ── dispatch ────────────────────────────────────────────────────────────────────
_PARSERS = {
    "PRICE_LIST": parse_price_list,
    "CONSUMER_SCHEME": parse_consumer_scheme,
    "EXCHANGE_SCHEME": parse_exchange_scheme,
    "CORPORATE_POLICY": parse_corporate_policy,
    "DISCOUNT_GRID": parse_discount_grid,
}


def parse_master(kind: str, content: bytes, *, filename: str | None = None) -> ParseResult:
    parser = _PARSERS.get(kind)
    if parser is None:
        raise MasterParseError(f"Unknown master kind '{kind}'.")
    if kind in ("PRICE_LIST", "DISCOUNT_GRID"):
        return parser(content, filename=filename)
    return parser(content)
