"""Integration tests for OEM native master ingestion into the tenant-scoped
versioned masters. Requires DATABASE_URL (a migrated Audit Core database).
"""
from __future__ import annotations

import os
from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, text

from audit_core.oem_master_parsers import (
    CorporateBenefitRow,
    CorporateCompany,
    DiscountRow,
    ParseResult,
    parse_price_list,
)
from audit_core.oem_price_masters import (
    _project_oem,
    ingest_corporate_policy,
    ingest_discount_document,
    ingest_price_list,
)
from audit_core.price_lists import find_effective_price_plan
from audit_core.uc03_model_resolution import _sku_rows_for_version

_HEADER = [
    "Sl. No.", "Category", "Model", "Variant", "Trim", "Fuel", "Transmission",
    "Drive", "Seater", "Ex-Showroom Price", "TCS", "Insurance",
    "Ext. Warranty (4th Yr)", "Ext. Warranty (4th & 5th Yr)", "Accessories Kit",
    "RSA (1 Yr)", "FASTag", "Registration\n(Individual / w/o Hyp.)",
    "On-Road Price\n(Individual / w/o Hyp.)", "Registration\n(Corporate / with Hyp.)",
    "On-Road Price\n(Corporate / with Hyp.)", "Source Sheet",
]


def _price_bytes(*price_rows: tuple[str, str, int]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Price List"
    ws.append(_HEADER)
    for sl, (model, variant, ex) in enumerate(price_rows, start=1):
        tcs, ins, ew4, ew45, acc, rsa, fastag, reg = 100, 200, 300, 400, 500, 60, 40, 5000
        onroad = ex + tcs + ins + ew4 + ew45 + acc + rsa + fastag + reg
        ws.append([
            sl, "PV", model, variant, variant.split()[0], "PETROL", "MT", "2WD", "5",
            ex, tcs, ins, ew4, ew45, acc, rsa, fastag, reg, onroad, reg, onroad, "S",
        ])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _price_bytes_full(*rows: dict) -> bytes:
    """Same header/layout as _price_bytes, but every structured attribute is
    explicit per row instead of derived from the variant string -- needed to
    reproduce a real master sheet where several variants of the same model
    share some attributes (fuel/transmission/drive, even the same trim) but
    differ in others (seater), which _price_bytes' one-row-per-call helper
    can never exercise."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Price List"
    ws.append(_HEADER)
    for sl, row in enumerate(rows, start=1):
        tcs, ins, ew4, ew45, acc, rsa, fastag, reg = 100, 200, 300, 400, 500, 60, 40, 5000
        ex = row["ex"]
        onroad = ex + tcs + ins + ew4 + ew45 + acc + rsa + fastag + reg
        ws.append([
            sl, "PV", row["model"], row["variant"], row["trim"], row["fuel"],
            row["transmission"], row["drive"], row["seater"],
            ex, tcs, ins, ew4, ew45, acc, rsa, fastag, reg, onroad, reg, onroad, "S",
        ])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def connection():
    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        pytest.skip("DATABASE_URL is required for OEM master integration tests")
    engine = create_engine(database_url)
    suffix = uuid4().hex[:10]
    tenant_id = f"tenant-oemm-{suffix}"
    with engine.begin() as conn:
        oem_id = conn.execute(
            text("SELECT oem_id FROM auditcore.oems WHERE oem_code = 'MAHINDRA'")
        ).scalar_one_or_none()
        if oem_id is None:
            oem_id = conn.execute(
                text("INSERT INTO auditcore.oems (oem_code, oem_name) "
                     "VALUES ('MAHINDRA','Mahindra') RETURNING oem_id")
            ).scalar_one()
        category_id = conn.execute(
            text("INSERT INTO auditcore.product_categories (category_code, category_name) "
                 "VALUES (:c, 'Vehicle') RETURNING product_category_id"),
            {"c": f"CAT-{suffix}"},
        ).scalar_one()
        conn.execute(
            text("""
                INSERT INTO auditcore.projects
                    (tenant_id, project_code, project_name, oem_id, product_category_id, effective_start_date)
                VALUES (:t, :pc, 'OEM Master Project', :o, :cat, CURRENT_DATE)
            """),
            {"t": tenant_id, "pc": f"P-{suffix}", "o": oem_id, "cat": category_id},
        )
    engine.dispose()
    engine = create_engine(database_url)
    with engine.begin() as conn:
        conn.execute(text("SELECT set_config('app.platform_super_admin','true',true)"))
        conn.execute(text("SELECT set_config('app.tenant_id',:t,true)"), {"t": tenant_id})
        conn.tenant_id = tenant_id  # type: ignore[attr-defined]
        yield conn
    engine.dispose()


def test_price_list_ingest_creates_catalogue_and_resolves_plan(connection) -> None:
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    parsed = parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_000_000)))
    version_id, skus = ingest_price_list(
        connection,
        tenant_id=tenant_id,
        oem_id=project["oem_id"],
        effective_from=date(2026, 9, 3),
        parsed=parsed,
        actor_id="admin",
    )
    assert len(skus) == 1
    items = connection.execute(
        text("SELECT component_key, standard_amount FROM auditcore.price_list_items "
             "WHERE tenant_id = :t AND price_list_version_id = :v"),
        {"t": tenant_id, "v": version_id},
    ).mappings().all()
    keyed = {r["component_key"]: r["standard_amount"] for r in items}
    assert keyed["EX_SHOWROOM"] == Decimal("1000000.00")
    assert "REGISTRATION_CORPORATE" in keyed

    plan = find_effective_price_plan(connection, tenant_id=tenant_id, effective_on=date(2026, 9, 20))
    assert plan is not None
    assert plan["price_list_code"] == "OEM_NATIVE_PRICE_LIST"


def test_price_list_ingest_persists_trim_on_the_variant(connection) -> None:
    """oem_master_parsers.py has parsed Trim (a distinct masters column, not
    part of the Variant string) from day one, but oem_price_masters.py never
    included it in the attrs written to product_variants.attributes --
    confirmed live: several distinct trims sharing identical fuel/
    transmission/drive/seater made "Variant" the only disambiguator in the
    Modify Model picker, with no way to narrow by trim first."""
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"],
        effective_from=date(2026, 9, 3),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_000_000))),
        actor_id="admin",
    )
    trim = connection.execute(
        text("""
            SELECT pv.attributes ->> 'trim' FROM auditcore.product_variants pv
            JOIN auditcore.product_models pm ON pm.model_id = pv.model_id
            WHERE pm.oem_id = :o AND pv.variant_name = 'MX1 PMT 2WD'
        """),
        {"o": project["oem_id"]},
    ).scalar_one()
    assert trim == "MX1"


def test_reupload_backfills_trim_onto_an_already_ingested_variant(connection) -> None:
    """A re-upload of the same masters file is a routine admin action, not a
    one-off -- when the parser starts extracting a field it didn't capture
    before, an already-ingested variant must pick it up too, not keep
    serving stale attributes forever just because its variant_code already
    existed."""
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"],
        effective_from=date(2026, 9, 3),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_000_000))),
        actor_id="admin",
    )
    # Simulate the pre-fix state: strip trim back out, as if this variant had
    # been ingested before oem_price_masters.py carried it through.
    connection.execute(
        text("""
            UPDATE auditcore.product_variants SET attributes = attributes - 'trim'
            WHERE variant_id IN (
                SELECT pv.variant_id FROM auditcore.product_variants pv
                JOIN auditcore.product_models pm ON pm.model_id = pv.model_id
                WHERE pm.oem_id = :o AND pv.variant_name = 'MX1 PMT 2WD'
            )
        """),
        {"o": project["oem_id"]},
    )
    assert connection.execute(
        text("""
            SELECT pv.attributes ->> 'trim' FROM auditcore.product_variants pv
            JOIN auditcore.product_models pm ON pm.model_id = pv.model_id
            WHERE pm.oem_id = :o AND pv.variant_name = 'MX1 PMT 2WD'
        """),
        {"o": project["oem_id"]},
    ).scalar_one_or_none() is None

    # Re-upload the exact same file -- same variant_code, so the old
    # insert-only path would have left it untouched forever.
    ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"],
        effective_from=date(2026, 10, 1),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_100_000))),
        actor_id="admin",
    )
    trim = connection.execute(
        text("""
            SELECT pv.attributes ->> 'trim' FROM auditcore.product_variants pv
            JOIN auditcore.product_models pm ON pm.model_id = pv.model_id
            WHERE pm.oem_id = :o AND pv.variant_name = 'MX1 PMT 2WD'
        """),
        {"o": project["oem_id"]},
    ).scalar_one()
    assert trim == "MX1"


def test_price_list_ingest_persists_trim_across_sibling_variants_of_one_model(connection) -> None:
    """Reproduces the live master sheet verbatim: SCORPIO CLASSIC carries
    three variants that share fuel/transmission/drive (DIESEL/MT/2WD) and,
    for two of them, the identical trim "S" -- differing only in seater (7
    vs 9) for those two, and in trim ("S11") for the third. The single-row
    tests above (test_price_list_ingest_persists_trim_on_the_variant,
    test_reupload_backfills_trim_onto_an_already_ingested_variant) only
    ever ingest one variant per model per call, so they cannot catch a bug
    that only shows up when several sibling variants of the same model are
    ingested together in one upload -- which is the real-world shape being
    reported live as "trim is missing" for exactly these three rows."""
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    parsed = parse_price_list(_price_bytes_full(
        {"model": "SCORPIO CLASSIC", "variant": "Classic S BS6.2 - E", "trim": "S",
         "fuel": "DIESEL", "transmission": "MT", "drive": "2WD", "seater": "7", "ex": 1_336_701},
        {"model": "SCORPIO CLASSIC", "variant": "Classic S - 9 STR BS6.2 - E", "trim": "S",
         "fuel": "DIESEL", "transmission": "MT", "drive": "2WD", "seater": "9", "ex": 1_383_901},
        {"model": "SCORPIO CLASSIC", "variant": "Classic S11 BS6.2 - E", "trim": "S11",
         "fuel": "DIESEL", "transmission": "MT", "drive": "2WD", "seater": "7", "ex": 1_739_901},
    ))
    assert not parsed.errors, parsed.errors
    version_id, sku_id_by_code = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"],
        effective_from=date(2026, 9, 3), parsed=parsed, actor_id="admin",
    )
    assert len(sku_id_by_code) == 3

    rows = _sku_rows_for_version(connection, tenant_id=tenant_id, price_list_version_id=version_id)
    by_variant = {r["variant_name"]: r for r in rows}
    assert by_variant["Classic S BS6.2 - E"]["trim"] == "S"
    assert by_variant["Classic S - 9 STR BS6.2 - E"]["trim"] == "S"
    assert by_variant["Classic S11 BS6.2 - E"]["trim"] == "S11"
    for variant_name in by_variant:
        assert by_variant[variant_name]["fuel_powertrain"] == "DIESEL"
        assert by_variant[variant_name]["transmission"] == "MT"
        assert by_variant[variant_name]["drive"] == "2WD"


def test_reupload_supersedes_by_effective_date(connection) -> None:
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"],
        effective_from=date(2026, 9, 3),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_000_000))),
        actor_id="admin",
    )
    v2, _ = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"],
        effective_from=date(2026, 10, 1),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_100_000))),
        actor_id="admin",
    )
    sept = find_effective_price_plan(connection, tenant_id=tenant_id, effective_on=date(2026, 9, 15))
    octo = find_effective_price_plan(connection, tenant_id=tenant_id, effective_on=date(2026, 10, 15))
    assert sept["version_no"] == 1
    assert octo["version_no"] == 2
    assert octo["price_list_version_id"] == v2


def test_discount_ingest_and_tombstone(connection) -> None:
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"],
        effective_from=date(2026, 9, 3),
        parsed=parse_price_list(
            _price_bytes(
                ("THAR ROXX", "MX1 PMT 2WD", 1_000_000),
                ("SCORPIO N", "Z8 L", 2_000_000),
            )
        ),
        actor_id="admin",
    )

    sept = ParseResult(kind="EXCHANGE_SCHEME")
    sept.discount_rows = [
        DiscountRow(1, "EXCHANGE", "Thar Roxx", [], [("EXCHANGE_BONUS", Decimal(25000))],
                    Decimal(25000), {"section": "EXCHANGE_PERSONAL"}),
        DiscountRow(2, "EXCHANGE", "Scorpio-N", [], [("EXCHANGE_BONUS", Decimal(25000))],
                    Decimal(25000), {"section": "EXCHANGE_PERSONAL"}),
    ]
    summary = ingest_discount_document(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
        master_kind="EXCHANGE_SCHEME", effective_from=date(2026, 9, 3), parsed=sept,
        actor_id="admin",
    )
    assert summary["published"] == 2
    assert summary["tombstoned"] == 0

    # October drops Scorpio-N -> its scheme is tombstoned with an empty version
    octo = ParseResult(kind="EXCHANGE_SCHEME")
    octo.discount_rows = [
        DiscountRow(1, "EXCHANGE", "Thar Roxx", [], [("EXCHANGE_BONUS", Decimal(30000))],
                    Decimal(30000), {"section": "EXCHANGE_PERSONAL"}),
    ]
    summary2 = ingest_discount_document(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
        master_kind="EXCHANGE_SCHEME", effective_from=date(2026, 10, 1), parsed=octo,
        actor_id="admin",
    )
    assert summary2["published"] == 1
    assert summary2["tombstoned"] == 1

    live_benefits = connection.execute(
        text("""
            SELECT s.scheme_code, count(b.benefit_id) AS n
            FROM auditcore.discount_schemes s
            JOIN auditcore.discount_scheme_versions v
              ON v.tenant_id = s.tenant_id AND v.discount_scheme_id = s.discount_scheme_id
             AND v.version_no = (SELECT max(v2.version_no) FROM auditcore.discount_scheme_versions v2
                                 WHERE v2.tenant_id = v.tenant_id
                                   AND v2.discount_scheme_id = v.discount_scheme_id)
            LEFT JOIN auditcore.discount_scheme_benefits b
              ON b.tenant_id = v.tenant_id AND b.discount_scheme_version_id = v.discount_scheme_version_id
            WHERE s.tenant_id = :t
            GROUP BY s.scheme_code
        """),
        {"t": tenant_id},
    ).mappings().all()
    by_code = {r["scheme_code"]: r["n"] for r in live_benefits}
    scorpio = next(c for c in by_code if "SCORPIO" in c)
    assert by_code[scorpio] == 0  # tombstoned: latest version carries no benefit


def test_corporate_policy_batches_large_company_registry(connection) -> None:
    # A company list can run to hundreds/thousands of rows — this exercises the
    # batched upsert (chunk_size=500 default) across a chunk boundary, plus a
    # duplicate corporate_code within one file (last one must win, matching the
    # old sequential upsert's semantics) and full-replace on re-upload.
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"],
        effective_from=date(2026, 9, 3),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_000_000))),
        actor_id="admin",
    )

    parsed = ParseResult(kind="CORPORATE_POLICY")
    parsed.corporate_benefits = [
        CorporateBenefitRow(1, "Z", "Thar Roxx", Decimal(10000), Decimal(5000), Decimal(15000)),
    ]
    parsed.corporate_companies = [
        CorporateCompany(f"CORP-{i:04d}", f"Company {i}", "PSU", "Z") for i in range(600)
    ]
    parsed.corporate_companies.append(CorporateCompany("CORP-0000", "Company 0 Renamed", "PSU", "Z"))

    summary = ingest_corporate_policy(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
        effective_from=date(2026, 9, 3), parsed=parsed, upload_id=uuid4(), actor_id="admin",
    )
    assert summary["published"] == 1
    assert summary["companies"] == 601

    rows = connection.execute(
        text("SELECT corporate_code, corporate_name FROM auditcore.corporate_privilege_registry "
             "WHERE oem_code = 'MAHINDRA' ORDER BY corporate_code"),
    ).mappings().all()
    assert len(rows) == 600  # the duplicate code collapsed to one row
    by_corp_code = {r["corporate_code"]: r["corporate_name"] for r in rows}
    assert by_corp_code["CORP-0000"] == "Company 0 Renamed"  # last one wins
    assert by_corp_code["CORP-0599"] == "Company 599"

    # re-upload fully replaces the registry
    parsed2 = ParseResult(kind="CORPORATE_POLICY")
    parsed2.corporate_benefits = parsed.corporate_benefits
    parsed2.corporate_companies = [CorporateCompany("CORP-0000", "Only Company Left", "PSU", "Z")]
    ingest_corporate_policy(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
        effective_from=date(2026, 10, 1), parsed=parsed2, upload_id=uuid4(), actor_id="admin",
    )
    remaining = connection.execute(
        text("SELECT count(*) FROM auditcore.corporate_privilege_registry WHERE oem_code='MAHINDRA'")
    ).scalar_one()
    assert remaining == 1


def _items(connection, tenant_id, version_id) -> dict[tuple[str, str], tuple[Decimal, date]]:
    rows = connection.execute(
        text("""SELECT s.sku_code, i.component_key, i.standard_amount, i.price_since
                FROM auditcore.price_list_items i JOIN auditcore.product_skus s ON s.product_sku_id = i.product_sku_id
                WHERE i.tenant_id = :t AND i.price_list_version_id = :v"""),
        {"t": tenant_id, "v": version_id},
    ).all()
    return {(r[0], r[1]): (Decimal(r[2]), r[3]) for r in rows}


def _skus(connection, tenant_id, version_id) -> set[str]:
    return {code for code, _ in _items(connection, tenant_id, version_id)}


def _code_of(skus: dict, model_prefix_free_variant: str) -> str:
    """The SKU code a returned {code: id} map holds for a variant (the model
    part of the code is whatever the shared catalogue first named it)."""
    (code,) = [c for c in skus if c.split("::", 1)[1] == model_prefix_free_variant]
    return code


def test_a_complete_book_withdraws_what_it_does_not_list_and_dates_each_price(connection) -> None:
    """2026-10-20 product owner rules: the consolidated list is the whole book
    for its date. A SKU it does not list is withdrawn from that date, never
    carried forward at an older month's price (it used to be: the October
    book held 144 September SKUs at September prices). Every row keeps its
    own price-since date. (Replaces the 2026-09-30 test that carried a
    consolidated file's missing SKUs forward.)"""
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    v1, first_skus = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 1),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_000_000), ("XUV 3XO", "MX1", 800_000))),
        actor_id="admin",
    )
    stats: dict = {}
    v2, _ = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 19),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_050_000))),
        actor_id="admin", stats=stats,
    )
    assert stats == {"uploaded": 1, "changed": 1, "unchanged": 0, "carriedForward": 0, "withdrawn": 1,
                     "completeBook": True, "standingVersionId": str(v1), "laterVersionsRederived": 0,
                     "journeysRechecked": 0}
    thar, xuv = _code_of(first_skus, "MX1 PMT 2WD"), _code_of(first_skus, "MX1")
    first, second = _items(connection, tenant_id, v1), _items(connection, tenant_id, v2)
    assert len(first) == 20 and len(second) == 10  # the XUV is not in the second book
    assert second[(thar, "EX_SHOWROOM")] == (Decimal(1_050_000), date(2026, 9, 19))
    assert second[(thar, "TCS")] == (Decimal(100), date(2026, 9, 1))  # same figure: kept its date
    assert (xuv, "EX_SHOWROOM") not in second
    assert find_effective_price_plan(connection, tenant_id=tenant_id, effective_on=date(2026, 9, 20))["price_list_version_id"] == v2
    assert find_effective_price_plan(connection, tenant_id=tenant_id, effective_on=date(2026, 9, 10))["price_list_version_id"] == v1

    # The same file again: nothing changes, every date stands.
    again: dict = {}
    v3, _ = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 10, 1),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_050_000))),
        actor_id="admin", stats=again,
    )
    assert again["changed"] == 0 and again["unchanged"] == 1 and again["carriedForward"] == 0
    assert _items(connection, tenant_id, v3)[(thar, "EX_SHOWROOM")] == (Decimal(1_050_000), date(2026, 9, 19))


def _dealer_ev_bytes(sheets: dict[str, list[tuple[str, int]]]) -> bytes:
    """The dealer's EV price sheets: one sheet per model (the real layout)."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, variants in sheets.items():
        ws = wb.create_sheet(title)
        ws.append([None, "ADITYA MOTORS"])
        ws.append([None, "MODEL NAME", title, "ODISHA", "STATE", "CUTTACK"])
        ws.append([None, "MODEL & VARIANT", "Ex-showroom Price with Charger", "Tax Collection at Source (TCS)",
                   "Insurance", "Essential Accessories", "Fastag", "On Road Price - Individual"])
        ws.append([None] * 7 + ["Registration without Hypothecation", "On Road Price without Hypothecation",
                                "Registration with Hypothecation", "On Road Price with Hypothecation"])
        ws.append([None, None, "(A)", "(B)", "(C)", "(D)", "(E)", "(F)", "(G)=(A+B+C+D+E+F)", "(H)",
                   "(I)=(A+B+C+D+E+H)"])
        for variant, ex in variants:
            tcs, ins, acc, fastag, reg_without, reg_with = round(ex * 0.01), 80000, 25000, 500, 140, 1640
            base = ex + tcs + ins + acc + fastag
            ws.append([None, variant, ex, tcs, ins, acc, fastag, reg_without, base + reg_without, reg_with,
                       base + reg_with])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_a_dealer_sheet_carries_forward_only_the_models_it_does_not_cover(connection) -> None:
    """The dealer's EV sheets price only their own models: SKUs of other
    models are carried forward; a SKU of a covered model the sheet no longer
    lists is dropped."""
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    _, first = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 1),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_000_000), ("BE 6", "BE 6 One B59 R18 NCH", 1_890_000),
                                             ("BE 6", "BE 6 Two B59 R19 C7", 2_240_000))),
        actor_id="admin",
    )
    stats: dict = {}
    v2, second = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 7),
        parsed=parse_price_list(_dealer_ev_bytes({"BE 6": [("BE 6 Two B59 R19 C7", 2_250_000)]}),
                                filename="EV PRICE -07th Sept'26.xlsx"),
        actor_id="admin", stats=stats,
    )
    assert stats["completeBook"] is False and stats["carriedForward"] == 1 and stats["withdrawn"] == 1
    items = _items(connection, tenant_id, v2)
    assert items[(_code_of(first, "MX1 PMT 2WD"), "EX_SHOWROOM")] == (Decimal(1_000_000), date(2026, 9, 1))
    assert items[(_code_of(first, "BE 6 Two B59 R19 C7"), "EX_SHOWROOM")] == (Decimal(2_250_000), date(2026, 9, 7))
    assert (_code_of(first, "BE 6 One B59 R18 NCH"), "EX_SHOWROOM") not in items
    assert set(second) == {_code_of(first, "BE 6 Two B59 R19 C7")}  # the same SKU as the consolidated row


def test_an_earlier_dated_dealer_sheet_reaches_only_later_dealer_sheet_versions(connection) -> None:
    """The EV sheet dated the 1st arrives after books dated later. A later
    complete book is left alone (what it does not list is not sold from its
    date); a later dealer-sheet version, which priced only its own models,
    is re-derived on its own date with the earlier sheet's other models."""
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    consolidated, _ = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 3),
        parsed=parse_price_list(_price_bytes(("XUV 3XO", "MX1", 800_000))), actor_id="admin",
    )
    stats: dict = {}
    ev, _ = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 1),
        parsed=parse_price_list(_dealer_ev_bytes({"BE 6": [("BE 6 PACK ONE", 1_900_000)]}),
                                filename="EV PRICE -01th Sept'26.xlsx"),
        actor_id="admin", stats=stats,
    )
    assert stats["laterVersionsRederived"] == 0
    assert find_effective_price_plan(connection, tenant_id=tenant_id, effective_on=date(2026, 9, 2))["price_list_version_id"] == ev
    assert find_effective_price_plan(connection, tenant_id=tenant_id, effective_on=date(2026, 9, 5))["price_list_version_id"] == consolidated

    # A dealer sheet for another model, dated the 7th, then an earlier one
    # for a third model: the 7th's version gains it, the 3rd's does not.
    later, _ = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 7),
        parsed=parse_price_list(_dealer_ev_bytes({"XEV 9E": [("XEV 9e One B59 R19 NCH", 2_190_000)]}),
                                filename="EV PRICE -07th Sept'26.xlsx"),
        actor_id="admin",
    )
    assert len(_skus(connection, tenant_id, later)) == 2  # XEV + the XUV carried from the 3rd
    again: dict = {}
    ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 5),
        parsed=parse_price_list(_dealer_ev_bytes({"XUV400": [("XUV400 EL PRO", 1_600_000)]}),
                                filename="EV PRICE -05th Sept'26.xlsx"),
        actor_id="admin", stats=again,
    )
    assert again["laterVersionsRederived"] == 1  # the 7th only
    on_the_9th = find_effective_price_plan(connection, tenant_id=tenant_id, effective_on=date(2026, 9, 9))
    assert len(_skus(connection, tenant_id, on_the_9th["price_list_version_id"])) == 3
    assert len(_skus(connection, tenant_id, consolidated)) == 1  # published versions are untouched


def test_discount_grid_ingests_per_model_and_reports_unknown_models(connection) -> None:
    from audit_core.oem_master_parsers import GridRow
    from audit_core.oem_price_masters import grid_row_for_model, ingest_discount_grid

    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    _, skus = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 1),
        parsed=parse_price_list(_price_bytes(("THAR ROXX", "MX1 PMT 2WD", 1_000_000), ("PICK UP", "S6", 800_000))),
        actor_id="admin",
    )
    parsed = ParseResult(kind="DISCOUNT_GRID")
    parsed.grid_rows = [
        GridRow(3, "Thar Roxx", ["Thar Roxx"], True, 60, Decimal(5000), Decimal(60), Decimal(3000)),
        GridRow(4, "PICKUP, MAXX", ["PICKUP", "MAXX"], True, 30, Decimal(7000), Decimal(50), Decimal(3000)),
        # A model no book or catalogue carries (until 2026-11-01 this was the
        # September grid's "SCORPIO N Old", which now names the SCORPIO N model).
        GridRow(5, "QUANTO", ["QUANTO"], False, None, None, None, None),
    ]
    parsed.grid_parameters = [{"parameter": "Penalty Amount", "note": "₹30,000 per case"}]
    summary = ingest_discount_grid(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
        effective_from=date(2026, 9, 1), parsed=parsed, upload_id=None, actor_id="admin",
    )
    assert summary["published"] == 4 and summary["resolved"] == 2 and summary["parameters"] == 1
    assert [u.split("'")[1] for u in summary["unresolved"]] == ["MAXX", "QUANTO"]
    thar_model = connection.execute(
        text("SELECT model_id FROM auditcore.product_skus WHERE product_sku_id = :s"),
        {"s": skus["THAR_ROXX::MX1 PMT 2WD"]},
    ).scalar_one()
    row = grid_row_for_model(connection, tenant_id=tenant_id, model_id=thar_model, effective_on=date(2026, 9, 15))
    assert row["booking_protection_days"] == 60 and Decimal(row["insurance_od_percent"]) == 60
    assert row["parameters"][0]["parameter"] == "Penalty Amount"
    assert grid_row_for_model(connection, tenant_id=tenant_id, model_id=thar_model, effective_on=date(2026, 8, 15)) is None


def test_the_master_object_key_is_filed_by_effective_month() -> None:
    """Where a master document is kept.

    Filed under the month the document is EFFECTIVE for, not the day it
    was uploaded: the August-named grid whose sheet says "Effective 1st
    Sep" belongs to September. The upload id keeps a re-upload of the same
    month from overwriting the file it supersedes, because uploads are not
    idempotent and two files for one month is ordinary.
    """
    from datetime import date as _date
    from uuid import UUID as _UUID

    from audit_core.oem_price_masters import _master_content_type, _master_object_key

    key = _master_object_key(
        tenant_id="tenant-mahindraorissa", oem_code="mahindra", master_kind="PRICE_LIST",
        effective_from=_date(2026, 10, 3),
        upload_id=_UUID("11111111-2222-3333-4444-555555555555"),
        filename="Mahindra Consolidated Price List Oct'26 (1).xlsx",
    )
    assert key == (
        "tenant-mahindraorissa/oem-masters/MAHINDRA/2026-10/PRICE_LIST/"
        "11111111-2222-3333-4444-555555555555_Mahindra_Consolidated_Price_List_Oct_26_1_.xlsx"
    )

    # A grid named for August but effective in September files under September.
    september = _master_object_key(
        tenant_id="t", oem_code="MAHINDRA", master_kind="DISCOUNT_GRID",
        effective_from=_date(2026, 9, 1),
        upload_id=_UUID("66666666-7777-8888-9999-000000000000"),
        filename="Odisha PV & CV Discount Grid-1 Aug 2026.xlsx",
    )
    assert "/2026-09/DISCOUNT_GRID/" in september

    assert _master_content_type("book.xlsx").endswith("spreadsheetml.sheet")
    assert _master_content_type("scheme.PDF") == "application/pdf"
    assert _master_content_type(None) == "application/octet-stream"


# ── scheme resolution (2026-10-20 product owner rules) ──────────────────────────
def _eligibility(connection, tenant_id, code_like: str) -> list[dict]:
    return [dict(r) for r in connection.execute(
        text("""
            SELECT s.scheme_code, v.version_no, e.variant_id, pv.variant_name, pm.model_name, e.criteria,
                   e.customer_type_code
            FROM auditcore.discount_schemes s
            JOIN auditcore.discount_scheme_versions v
              ON v.tenant_id = s.tenant_id AND v.discount_scheme_id = s.discount_scheme_id
            JOIN auditcore.discount_scheme_eligibility e
              ON e.tenant_id = v.tenant_id AND e.discount_scheme_version_id = v.discount_scheme_version_id
            JOIN auditcore.product_models pm ON pm.model_id = e.model_id
            LEFT JOIN auditcore.product_variants pv ON pv.variant_id = e.variant_id
            WHERE s.tenant_id = :t AND s.scheme_code LIKE :c
            ORDER BY s.scheme_code, pv.variant_name
        """),
        {"t": tenant_id, "c": code_like},
    ).mappings().all()]


def test_scheme_variants_that_do_not_resolve_are_stored_unresolved_never_model_wide(connection) -> None:
    from audit_core.oem_master_parsers import split_variant_cell

    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 1),
        parsed=parse_price_list(_price_bytes(
            ("XUV 3XO", "AX5 PM MT", 1_029_900), ("XUV 3XO", "MX2 PRO PM AT", 900_000),
            ("XUV 3XO", "MX2 PRO PM MT", 880_000),
            ("XUV400", "EL PRO FC - Battery 34.5 Kwh - Charger 7.2 Kw - E", 1_700_000),
        )),
        actor_id="admin",
    )
    sept = ParseResult(kind="CONSUMER_SCHEME")
    ice = {"vehicleCategory": "ICE"}
    sept.discount_rows = [
        DiscountRow(1, "CONSUMER", "XUV 3XO", ["AX5 PM MT"], [("CASH_DISCOUNT", Decimal(30000))], Decimal(30000), ice),
        DiscountRow(2, "CONSUMER", "XUV 3XO", split_variant_cell("MX2 PRO PM AT & MT"),
                    [("CASH_DISCOUNT", Decimal(60000))], Decimal(60000), ice),
        DiscountRow(3, "CONSUMER", "XUV 3XO", ["MX9 IMAGINARY"], [("CASH_DISCOUNT", Decimal(10000))],
                    Decimal(10000), ice),
        DiscountRow(4, "CONSUMER", "XUV400", split_variant_cell(
            "EC PRO - Battery 34.5 Kwh-E-NCH\nEL PRO FC - Battery 34.5 Kwh -\nCharger 7.2 Kw - E"),
            [("CASH_DISCOUNT", Decimal(250000))], Decimal(250000), {"vehicleCategory": "EV"}),
    ]
    summary = ingest_discount_document(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
        master_kind="CONSUMER_SCHEME", effective_from=date(2026, 9, 3), parsed=sept, actor_id="admin",
    )
    rows = _eligibility(connection, tenant_id, "MAH_CONSUMER_%")
    assert not [r for r in rows if r["criteria"].get("scope") == "MODEL"]  # nothing named became model-wide
    by_variant = {(r["model_name"], r["variant_name"]): r for r in rows if r["variant_id"]}
    assert ("XUV 3XO", "MX2 PRO PM AT") in by_variant and ("XUV 3XO", "MX2 PRO PM MT") in by_variant
    # The catalogue is shared by every test project of the OEM, so another
    # model spelt "XUV 3XO" may exist without these variants: unresolved
    # rows are checked by what no target model carries (the summary).
    unresolved = [r for r in rows if r["criteria"].get("scope") == "UNRESOLVED"]
    assert unresolved and all(r["variant_id"] is None for r in unresolved)
    texts = {t for r in unresolved for t in r["criteria"]["unresolvedVariantTexts"]}
    assert {"EC PRO - Battery 34.5 Kwh-E-NCH", "MX9 IMAGINARY"} <= texts
    bev = [r for r in rows if "EC PRO" in str(r["criteria"]) or "EL PRO" in str(r["criteria"])]
    assert bev and {r["model_name"] for r in bev} == {"XUV400"}  # the EV offer never reaches XUV 3XO
    assert {u["variantTexts"][0] for u in summary["unresolvedVariants"]} == {
        "MX9 IMAGINARY", "EC PRO - Battery 34.5 Kwh-E-NCH"}

    # An EV offer named against an ICE model is refused, not dropped quietly.
    from audit_core.errors import ValidationError

    wrong = ParseResult(kind="CONSUMER_SCHEME")
    wrong.discount_rows = [DiscountRow(1, "CONSUMER", "XUV 3XO", ["EC PRO"], [("CASH_DISCOUNT", Decimal(250000))],
                                       Decimal(250000), {"vehicleCategory": "EV"})]
    with pytest.raises(ValidationError) as refused:
        ingest_discount_document(
            connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
            master_kind="CONSUMER_SCHEME", effective_from=date(2026, 9, 3), parsed=wrong, actor_id="admin",
        )
    assert "an EV offer named against 'XUV 3XO'" in refused.value.detail


def test_aliases_never_merge_models_and_conditions_make_separate_schemes(connection) -> None:
    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 1),
        parsed=parse_price_list(_price_bytes(
            ("THE BOSS- BOLERO+NEO", "B4 BS6.2 - Refresh", 900_000),
            ("THE BOSS- BOLERO+NEO", "N8 BS6.2 - Refresh", 1_000_000),
            ("THE BOSS- BOLERO+NEO", "P10 BS6.2", 1_100_000),
            ("PICK UP", "BOL CAMPER 2WD BS6.2 - Refresh", 950_000),
            ("PICK UP", "BOL PIK UP FB 1.3T - BS6.2 - PS", 850_000),
            ("XUV 3XO", "AX5 PM MT", 1_029_900), ("XUV 3XO", "REVX M PM MT", 1_100_000),
            ("XUV 3XO", "REVX EDGE PM MT", 1_200_000),
        )),
        actor_id="admin",
    )
    section = {"section": "EXCHANGE_PERSONAL"}
    exchange = ParseResult(kind="EXCHANGE_SCHEME")
    exchange.discount_rows = [
        DiscountRow(1, "EXCHANGE", "Camper", [], [("EXCHANGE_BONUS", Decimal(20000))], Decimal(20000),
                    {"section": "EXCHANGE_COMMERCIAL"}),
        DiscountRow(2, "EXCHANGE", "Pickup", [], [("EXCHANGE_BONUS", Decimal(35000))], Decimal(35000),
                    {"section": "EXCHANGE_COMMERCIAL"}),
        DiscountRow(3, "EXCHANGE", "XUV 3XO (Including REVX)", [], [("EXCHANGE_BONUS", Decimal(50000))],
                    Decimal(50000), section),
        DiscountRow(4, "EXCHANGE", "XUV 3XO (Including REVX) against Old Maruti Vehicle**", [],
                    [("EXCHANGE_BONUS", Decimal(70000))], Decimal(70000), section),
        DiscountRow(5, "EXCHANGE", "XUV300", [], [("EXCHANGE_BONUS", Decimal(20000))], Decimal(20000), section),
    ]
    summary = ingest_discount_document(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
        master_kind="EXCHANGE_SCHEME", effective_from=date(2026, 9, 3), parsed=exchange, actor_id="admin",
    )
    assert summary["published"] == 4  # XUV300 is not sold: reported, not applied
    assert any("'XUV300' is not in the loaded price list" in u for u in summary["unresolved"])
    rows = _eligibility(connection, tenant_id, "MAH_EXCHANGE_%")
    codes = {r["scheme_code"] for r in rows}
    camper = [r for r in rows if r["scheme_code"].endswith("PICK_UP_CAMPER")]
    pickup = [r for r in rows if r["scheme_code"].endswith("PICK_UP_PICKUP")]
    assert [r["variant_name"] for r in camper] == ["BOL CAMPER 2WD BS6.2 - Refresh"]
    assert "BOL PIK UP FB 1.3T - BS6.2 - PS" in [r["variant_name"] for r in pickup]
    assert not [r for r in pickup if "CAMPER" in r["variant_name"]]
    maruti = [r for r in rows if r["scheme_code"].endswith("OLD_MARUTI")]
    assert maruti and all(r["criteria"]["requiredOldVehicle"] == {"brand": "MARUTI"} for r in maruti)
    general = [r for r in rows if "XUV" in r["scheme_code"] and not r["scheme_code"].endswith("OLD_MARUTI")]
    assert general and all(r["version_no"] == 1 for r in general + maruti)  # two schemes, not two versions
    assert len(codes) == 4

    corporate = ParseResult(kind="CORPORATE_POLICY")
    corporate.corporate_benefits = [
        CorporateBenefitRow(1, "Z", "BOLERO NEO+", Decimal(24000), Decimal(16000), Decimal(40000)),
        CorporateBenefitRow(2, "Z", "BOLERO NEO", Decimal(18000), Decimal(12000), Decimal(30000)),
        CorporateBenefitRow(3, "Z", "BOLERO", Decimal(18000), Decimal(12000), Decimal(30000)),
        CorporateBenefitRow(4, "F", "XUV 3XO (Not on 3XO EV)", Decimal(2500), Decimal(2500), Decimal(5000)),
        CorporateBenefitRow(5, "F", "3XO REVX (Not on REVX Edge)", Decimal(5000), Decimal(5000), Decimal(10000)),
    ]
    ingest_corporate_policy(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
        effective_from=date(2026, 9, 3), parsed=corporate, upload_id=uuid4(), actor_id="admin",
    )
    rows = _eligibility(connection, tenant_id, "MAH_CORPORATE_%")
    by_code: dict[str, list] = {}
    for r in rows:
        if r["criteria"].get("scope") != "UNRESOLVED":  # other test models of the same name
            by_code.setdefault(r["scheme_code"], []).append(r["variant_name"])
    # #468 files the September "THE BOSS- BOLERO+NEO" rows under three
    # models; "BOLERO NEO+" is the Neo Plus (the "+" read as PLUS before the
    # identity), never the Neo next to its own row (review of PR 472, B4).
    print(sorted(by_code))
    models = {r["scheme_code"]: r["model_name"] for r in rows}
    assert models["MAH_CORPORATE_Z_BOLERO_NEO_PLUS"] == "BOLERO NEO PLUS"
    assert models["MAH_CORPORATE_Z_BOLERO_NEO"] == "BOLERO NEO"
    assert models["MAH_CORPORATE_Z_BOLERO"] == "BOLERO"
    assert by_code["MAH_CORPORATE_F_XUV_3XO_REVX"] == ["REVX M PM MT"]  # REVX only, Edge excluded
    assert set(by_code["MAH_CORPORATE_F_XUV_3XO"]) == {None}  # the model-wide XUV 3XO line
    assert all(r["customer_type_code"].startswith("CORPORATE:") for r in rows)


def test_a_scheme_row_naming_an_unreadable_model_blocks_the_upload(connection) -> None:
    from audit_core.errors import ValidationError

    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 1),
        parsed=parse_price_list(_price_bytes(("XUV 3XO", "AX5 PM MT", 1_029_900))), actor_id="admin",
    )
    parsed = ParseResult(kind="EXCHANGE_SCHEME")
    parsed.discount_rows = [
        DiscountRow(1, "EXCHANGE", "XUV 3XO EDGE & REVX Variants", [], [("EXCHANGE_BONUS", Decimal(10000))],
                    Decimal(10000), {"section": "EXCHANGE_PERSONAL"}),
    ]
    with pytest.raises(ValidationError) as refused:
        ingest_discount_document(
            connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
            master_kind="EXCHANGE_SCHEME", effective_from=date(2026, 9, 3), parsed=parsed, actor_id="admin",
        )
    assert "name a model that is not known" in refused.value.detail
    assert "XUV 3XO EDGE & REVX Variants" in refused.value.detail


def test_variant_match_ignores_the_registration_suffix() -> None:
    from audit_core.oem_price_masters import _match_variants

    private, commercial = uuid4(), uuid4()
    matches = _match_variants(["P10 BS6.2", "Z8 G MT 2WD 7STR"], [
        ("P10 BS6.2 (Private)", private), ("P10 BS6.2 (Commercial)", commercial), ("Z8 G MT 2WD 7 STR", uuid4()),
    ])
    assert matches[0] == ("P10 BS6.2", [private, commercial])
    assert len(matches[1][1]) == 1


def test_the_october_grid_keeps_tiers_and_variant_groups_and_falls_back_a_month(connection) -> None:
    from audit_core.oem_master_parsers import GridRow
    from audit_core.oem_price_masters import grid_row_for_model, ingest_discount_grid

    tenant_id = connection.tenant_id
    project = _project_oem(connection, tenant_id)
    _, skus = ingest_price_list(
        connection, tenant_id=tenant_id, oem_id=project["oem_id"], effective_from=date(2026, 9, 1),
        parsed=parse_price_list(_price_bytes(
            ("PICK UP", "BOL CAMPER 2WD BS6.2 - Refresh", 950_000), ("PICK UP", "BOL PIK UP FB 1.3T - BS6.2 - PS", 850_000),
            ("THAR ROXX", "MX1 PMT 2WD", 1_000_000),
        )),
        actor_id="admin",
    )
    camper_sku = next(sku for code, sku in skus.items() if "CAMPER" in code.upper())
    pick_up = connection.execute(text("SELECT model_id FROM auditcore.product_skus WHERE product_sku_id = :s"),
                                 {"s": camper_sku}).scalar_one()
    thar = connection.execute(text("SELECT model_id FROM auditcore.product_skus WHERE product_sku_id = :s"),
                              {"s": skus["THAR_ROXX::MX1 PMT 2WD"]}).scalar_one()
    try:
        august = ParseResult(kind="DISCOUNT_GRID")
        august.grid_rows = [GridRow(3, "Thar Roxx", ["Thar Roxx"], True, 60, Decimal(5000), Decimal(60), Decimal(3000))]
        ingest_discount_grid(connection, tenant_id=tenant_id, oem_id=project["oem_id"], oem_code="MAHINDRA",
                             effective_from=date(2026, 9, 1), parsed=august, upload_id=None, actor_id="admin")
        october = ParseResult(kind="DISCOUNT_GRID")
        tiers = {"BASE": Decimal(5000), "WITH_INSURANCE": Decimal(7000), "WITH_EW_INSURANCE": Decimal(7000)}
        october.grid_rows = [
            GridRow(14, "PICKUP, MAXX City", ["PICKUP", "MAXX City"], True, 30, Decimal(5000), Decimal(65),
                    Decimal(5000), tiers=tiers),
            GridRow(17, "Camper", ["Camper"], True, 60, Decimal(0), Decimal(65), Decimal(5000),
                    tiers={k: Decimal(0) for k in tiers}),
        ]
        summary = ingest_discount_grid(connection, tenant_id=tenant_id, oem_id=project["oem_id"],
                                       oem_code="MAHINDRA", effective_from=date(2026, 10, 1), parsed=october,
                                       upload_id=None, actor_id="admin")
        assert summary["resolved"] == 2  # MAXX City is not in this catalogue
        camper = grid_row_for_model(connection, tenant_id=tenant_id, model_id=pick_up, effective_on=date(2026, 10, 9),
                                    variant_name="BOL CAMPER 2WD BS6.2 - Refresh")
        pickup = grid_row_for_model(connection, tenant_id=tenant_id, model_id=pick_up, effective_on=date(2026, 10, 9),
                                    variant_name="BOL PIK UP FB 1.3T - BS6.2 - PS")
        assert camper["model_alias"] == "Camper" and Decimal(camper["agreed_buffer_amount"]) == 0
        assert pickup["model_alias"] == "PICKUP" and pickup["raw"]["tiers"]["WITH_INSURANCE"] == "7000"
        # Thar Roxx is in the August grid only: in October it takes the
        # previous month's line (product owner, 2026-11-02; R5 had none).
        assert grid_row_for_model(connection, tenant_id=tenant_id, model_id=thar, effective_on=date(2026, 9, 9))
        carried = grid_row_for_model(connection, tenant_id=tenant_id, model_id=thar, effective_on=date(2026, 10, 9))
        assert carried["model_alias"] == "Thar Roxx" and carried["carriedFromPreviousGrid"]
        assert carried["effective_from"] == date(2026, 9, 1)
    finally:
        connection.execute(text("DELETE FROM auditcore.dealer_discount_grid_rows WHERE tenant_id=:t"), {"t": tenant_id})
        connection.execute(text("DELETE FROM auditcore.dealer_discount_grid_versions WHERE tenant_id=:t"),
                           {"t": tenant_id})


# ── 2026-10-20 product owner rules: one vehicle, one SKU, across months ────────
@pytest.fixture
def fresh_oem():
    """A throwaway OEM and project in one transaction that is rolled back, so
    model identity is tested against an empty catalogue."""
    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        pytest.skip("DATABASE_URL is required for OEM master integration tests")
    engine = create_engine(database_url)
    suffix = uuid4().hex[:10]
    conn = engine.connect()
    tx = conn.begin()
    try:
        oem_id = conn.execute(text("INSERT INTO auditcore.oems (oem_code, oem_name) VALUES (:c, 'Fresh') RETURNING oem_id"),
                              {"c": f"FRESH{suffix}"}).scalar_one()
        category_id = conn.execute(
            text("INSERT INTO auditcore.product_categories (category_code, category_name) VALUES (:c, 'V') "
                 "RETURNING product_category_id"), {"c": f"FR-CAT-{suffix}"},
        ).scalar_one()
        tenant_id = f"tenant-fresh-{suffix}"
        conn.execute(
            text("INSERT INTO auditcore.projects (tenant_id, project_code, project_name, oem_id, product_category_id, "
                 "effective_start_date) VALUES (:t, :pc, 'Fresh', :o, :cat, CURRENT_DATE)"),
            {"t": tenant_id, "pc": f"FR-{suffix}", "o": oem_id, "cat": category_id},
        )
        conn.execute(text("SELECT set_config('app.platform_super_admin','true',true)"))
        conn.execute(text("SELECT set_config('app.tenant_id',:t,true)"), {"t": tenant_id})
        conn.tenant_id = tenant_id  # type: ignore[attr-defined]
        conn.oem_id = oem_id  # type: ignore[attr-defined]
        yield conn
    finally:
        tx.rollback()
        conn.close()
        engine.dispose()


def _book(*rows: tuple) -> bytes:
    """Consolidated rows (category, model, variant, ex-showroom[, source sheet])."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Price List"
    ws.append(_HEADER)
    for sl, row in enumerate(rows, start=1):
        category, model, variant, ex = row[:4]
        source = row[4] if len(row) > 4 else "S"
        tcs, ins, ew4, ew45, acc, rsa, fastag, reg = 100, 200, 300, 400, 500, 60, 40, 5000
        onroad = ex + tcs + ins + ew4 + ew45 + acc + rsa + fastag + reg
        ws.append([sl, category, model, variant, variant.split()[0], "DIESEL", "MT", "2WD", "5",
                   ex, tcs, ins, ew4, ew45, acc, rsa, fastag, reg, onroad, reg, onroad, source])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _load(c, on: date, content: bytes, **kw) -> tuple:
    stats: dict = {}
    version, skus = ingest_price_list(c, tenant_id=c.tenant_id, oem_id=c.oem_id, effective_from=on,
                                      parsed=parse_price_list(content, **kw), actor_id="admin", stats=stats)
    return version, skus, stats


def _sku_model_names(c, version) -> list[tuple[str, str]]:
    return sorted(tuple(r) for r in c.execute(text("""
        SELECT DISTINCT s.sku_code, pm.model_name FROM auditcore.price_list_items i
        JOIN auditcore.product_skus s ON s.product_sku_id = i.product_sku_id
        JOIN auditcore.product_models pm ON pm.model_id = s.model_id
        WHERE i.price_list_version_id = :v"""), {"v": version}).all())


def test_a_respelt_model_keeps_its_model_and_sku_code_and_each_month_its_price(fresh_oem) -> None:
    """September's "XUV3XO" is October's "XUV 3XO": one model, one SKU code,
    priced by the book in force on the date, whichever spelling the
    document uses. Real figures (AX5 PM MT, 03 Sep and 03 Oct 2026 books)."""
    from audit_core.oem_price_masters import find_book_sku

    c = fresh_oem
    _, sept_skus, _ = _load(c, date(2026, 9, 3), _book(("PV", "XUV3XO", "AX5 PM MT", 1_029_900.54)))
    octo, oct_skus, stats = _load(c, date(2026, 10, 3), _book(("PV", "XUV 3XO", "AX5 PM MT", 999_901.32)))
    assert sept_skus == oct_skus == {"XUV3XO::AX5 PM MT": sept_skus["XUV3XO::AX5 PM MT"]}
    assert stats["changed"] == 1 and stats["withdrawn"] == 0
    assert _sku_model_names(c, octo) == [("XUV3XO::AX5 PM MT", "XUV3XO")]  # never two spellings in one book
    for spelling in ("XUV3XO", "XUV 3XO", "xuv-3xo"):
        on_sept = find_book_sku(c, tenant_id=c.tenant_id, model_text=spelling, variant_text="AX5 PM MT",
                                effective_on=date(2026, 9, 15))
        on_oct = find_book_sku(c, tenant_id=c.tenant_id, model_text=spelling, variant_text="AX5 PM MT",
                               effective_on=date(2026, 10, 15))
        assert on_sept["sku"]["sku_code"] == on_oct["sku"]["sku_code"] == "XUV3XO::AX5 PM MT"
        assert on_sept["sku"]["ex_showroom"] == Decimal("1029900.54")
        assert on_oct["sku"]["ex_showroom"] == Decimal("999901.32")
    before = find_book_sku(c, tenant_id=c.tenant_id, model_text="XUV 3XO", variant_text="AX5 PM MT",
                           effective_on=date(2026, 8, 15))
    assert before["plan"] is None and before["sku"] is None
    assert before["reason"] == "No price master in force on 2026-08-15"


def test_october_names_reach_the_september_models(fresh_oem) -> None:
    """SCORPIO-N is the new Scorpio N (its rows come from the "SCORPIO N ...
    New" sheets), not September's old SCORPIO N; SCORPIO is SCORPIO
    CLASSIC; NEW THAR is NEW THAR 2WD & 4WD. A release marker that changed
    ("- Refresh" to "- N") keeps the SKU; two October rows that differ only
    by "- Refresh" stay two vehicles."""
    c = fresh_oem
    _, sept, _ = _load(c, date(2026, 9, 3), _book(
        ("PV", "SCORPIO N", "Z8 S D AT 2WD 7 STR BS6.2 - N", 1_917_499.58, "SCORPIO N D old"),
        ("PV", "New SCORPIO N", "Z8 S D AT 2WD 7 STR BS6.2 - Refresh", 1_988_996.00, "SCORPIO N D New"),
        ("PV", "SCORPIO CLASSIC", "Classic S BS6.2 - E", 1_336_700.64),
        ("PV", "NEW THAR 2WD & 4WD", "LXT D MT 4WD 4S HT BS6.2", 1_642_500.64),
    ))
    octo, october, stats = _load(c, date(2026, 10, 3), _book(
        ("PV", "SCORPIO-N", "Z8 S D AT 2WD 7 STR BS6.2 - N", 1_998_995.45, "SCORPIO N D New"),
        ("PV", "SCORPIO", "Classic S BS6.2 - E", 1_343_200.67),
        ("PV", "NEW THAR", "LXT D MT 4WD 4S HT BS6.2", 1_642_500.94),
        ("PV", "NEW THAR", "LXT D MT 4WD 4S HT BS6.2 - Refresh", 1_642_501.00),
    ))
    assert set(october) == {
        "NEW_SCORPIO_N::Z8 S D AT 2WD 7 STR BS6.2 - Refresh",  # the new Scorpio N's SKU, not the old one's
        "SCORPIO_CLASSIC::Classic S BS6.2 - E",
        "NEW_THAR_2WD_4WD::LXT D MT 4WD 4S HT BS6.2",
        "NEW_THAR_2WD_4WD::LXT D MT 4WD 4S HT BS6.2 - Refresh",  # genuinely new
    }
    assert october["NEW_SCORPIO_N::Z8 S D AT 2WD 7 STR BS6.2 - Refresh"] == sept["NEW_SCORPIO_N::Z8 S D AT 2WD 7 STR BS6.2 - Refresh"]
    assert stats["withdrawn"] == 1  # the old Scorpio N
    printed = c.execute(text("""
        SELECT i.metadata->'printedAs'->>'model', i.metadata->'printedAs'->>'variant'
        FROM auditcore.price_list_items i WHERE i.price_list_version_id = :v AND i.product_sku_id = :s
          AND i.component_key = 'EX_SHOWROOM'"""),
        {"v": octo, "s": october["NEW_SCORPIO_N::Z8 S D AT 2WD 7 STR BS6.2 - Refresh"]}).one()
    assert tuple(printed) == ("SCORPIO-N", "Z8 S D AT 2WD 7 STR BS6.2 - N")


def test_the_bolero_family_is_split_by_series_in_both_months(fresh_oem) -> None:
    """September prints one model "THE BOSS- BOLERO+NEO"; October prints the
    Bolero, the Bolero Neo and the Bolero Neo Plus. Both months file the rows
    under the three models (B / N / P series), so a SKU keeps its code."""
    c = fresh_oem
    _, sept, _ = _load(c, date(2026, 9, 3), _book(
        ("PV", "THE BOSS- BOLERO+NEO", "B4 BS6.2 - Refresh", 848_999.65, "NEW BOLERO NEO (PVT)"),
        ("PV", "THE BOSS- BOLERO+NEO", "B4 BS6.2 - Refresh", 848_999.65, "NEW BOLERO NEO-COM"),
        ("PV", "THE BOSS- BOLERO+NEO", "N10 BS6.2 - R - Refresh", 1_008_999.65, "NEW BOLERO NEO (PVT)"),
        ("PV", "THE BOSS- BOLERO+NEO", "P4 BS6.2", 1_138_999.58, "NEW BOLERO NEO (PVT)"),
    ))
    octo, october, _ = _load(c, date(2026, 10, 3), _book(
        ("PV", "Bolero", "B4 BS6.2 - Refresh", 856_998.57, "NEW BOLERO NEO (PVT)"),
        ("PV", "Bolero", "B4 BS6.2 - Refresh", 856_998.57, "NEW BOLERO NEO-COM"),
        ("PV", "Bolero neo", "N10 BS6.2 - R - Refresh", 999_898.93, "NEW BOLERO NEO (PVT)"),
        ("PV", "Bolero Neo Plus", "P4 BS6.2", 1_145_998.99, "NEW BOLERO NEO (PVT)"),
    ))
    assert october == sept
    assert _sku_model_names(c, octo) == [
        ("BOLERO::B4 BS6.2 - Refresh::COM", "BOLERO"),
        ("BOLERO::B4 BS6.2 - Refresh::PRI", "BOLERO"),
        ("BOLERO_NEO::N10 BS6.2 - R - Refresh::PRI", "BOLERO NEO"),
        ("BOLERO_NEO_PLUS::P4 BS6.2::PRI", "BOLERO NEO PLUS"),
    ]


def test_three_and_three_plus_are_two_variants_and_a_legacy_sku_is_repaired(fresh_oem) -> None:
    """"+" used to be stripped from the variant code, so the BE 6 SPORTEQ
    THREE+ (Rs 1,00,000 dearer) sat on the THREE variant."""
    c = fresh_oem
    content = _book(("BEV", "BE6 SPORTEQ", "BE 6 SPORTEQ THREE B70 R19 NCH", 2_294_999.49),
                    ("BEV", "BE6 SPORTEQ", "BE 6 SPORTEQ THREE+ B70 R19 NCH", 2_394_999.49))
    _, skus, _ = _load(c, date(2026, 9, 1), content)

    def variants() -> dict[str, tuple]:
        return {r[0]: (r[1], r[2]) for r in c.execute(text("""
            SELECT s.sku_code, v.variant_name, v.variant_code FROM auditcore.product_skus s
            JOIN auditcore.product_variants v ON v.variant_id = s.variant_id WHERE s.oem_id = :o"""),
            {"o": c.oem_id}).all()}

    three, plus = "BE6_SPORTEQ::BE 6 SPORTEQ THREE B70 R19 NCH", "BE6_SPORTEQ::BE 6 SPORTEQ THREE+ B70 R19 NCH"
    assert variants()[three] == ("BE 6 SPORTEQ THREE B70 R19 NCH", "BE_6_SPORTEQ_THREE_B70_R19_NCH")
    assert variants()[plus] == ("BE 6 SPORTEQ THREE+ B70 R19 NCH", "BE_6_SPORTEQ_THREE_PLUS_B70_R19_NCH")
    # The state an earlier load left behind: THREE+ on the THREE variant.
    c.execute(text("UPDATE auditcore.product_skus SET variant_id = (SELECT variant_id FROM auditcore.product_skus "
                   "WHERE product_sku_id = :three) WHERE product_sku_id = :plus"),
              {"three": skus[three], "plus": skus[plus]})
    assert variants()[plus][0] == "BE 6 SPORTEQ THREE B70 R19 NCH"
    _load(c, date(2026, 10, 3), content)
    assert variants()[plus] == ("BE 6 SPORTEQ THREE+ B70 R19 NCH", "BE_6_SPORTEQ_THREE_PLUS_B70_R19_NCH")


def test_a_dealer_sheet_never_downgrades_a_known_category_or_fuel(fresh_oem) -> None:
    """The EV sheets used to turn every BE 6 / XEV variant from BEV to
    UNSPECIFIED. A dealer sheet now says BEV for the EV models; ICE or
    UNSPECIFIED, or a blank, never replaces what the consolidated book said."""
    c = fresh_oem
    _load(c, date(2026, 9, 1), _book(("BEV", "BE 6", "BE 6 Pack One B59", 1_890_000),
                                     ("CV", "VEERO", "1.5XXL SD V2", 839_500)))
    _load(c, date(2026, 9, 2), _dealer_ev_bytes({"BE6": [("BE 6 Pack One B59", 1_890_000)]}))

    wb = Workbook()
    ws = wb.active
    ws.title = "VEERO"
    ws.append(["ADITYA MOTORS"])
    ws.append(["Price list w.e.f. Dt.03.09.2026"])
    ws.append(["MODEL NAME", None, "VEERO", None, "FUEL TYPE", None])
    ws.append(["Model & Variant", None, "Ex-showroom Price", "Tax Collection at Source (TCS)", "Insurance", None,
               "On Road Price - Individual"])
    ws.append([None] * 6 + ["Registration\n without Hypoth", "On Road Price\nwithout\nHypoth"])
    ws.append([None, None, "(A)", "(B)", "(C)", None, "(D)", "(E)=(A+B+C+D)"])
    ws.append(["1.5XXL SD V2", None, 839_500, 0, 34_000, None, 51_110, 839_500 + 34_000 + 51_110])
    buf = BytesIO()
    wb.save(buf)
    _load(c, date(2026, 9, 3), buf.getvalue())

    found = c.execute(text("""
        SELECT pm.model_name, pv.attributes->>'category', pv.fuel_powertrain FROM auditcore.product_variants pv
        JOIN auditcore.product_models pm ON pm.model_id = pv.model_id WHERE pm.oem_id = :o"""), {"o": c.oem_id}).all()
    assert len(found) == 2  # the dealer sheets reached the same two variants
    rows = {r[0]: (r[1], r[2]) for r in found}
    assert rows["BE 6"] == ("BEV", "ELECTRIC")
    assert rows["VEERO"] == ("CV", "DIESEL")


def test_the_book_in_force_ends_the_day_before_the_next_and_ignores_other_lists(fresh_oem) -> None:
    """A version is in force from its date until the next version's date.
    Only the OEM's own book counts: a segment list published later never
    shadows it (defect 14)."""
    from audit_core.price_lists import (
        create_price_list,
        create_price_list_version,
        publish_price_list_version,
    )

    c = fresh_oem
    sept, _, _ = _load(c, date(2026, 9, 3), _book(("CV", "MAXX CITY", "MAXX CITY 1.5 LX BS6.2", 839_501.15)))
    octo, _, _ = _load(c, date(2026, 10, 3), _book(("CV", "MAXX CITY", "MAXX CITY 1.5 LX BS6.2", 852_000.62)))
    segment = create_price_list(c, tenant_id=c.tenant_id, code="MAHINDRA_PASSENGER_VEHICLE", name="Template")
    shadow = create_price_list_version(c, tenant_id=c.tenant_id, price_list_id=segment, version_no=1,
                                       effective_from=date(2026, 10, 10))
    publish_price_list_version(c, tenant_id=c.tenant_id, price_list_version_id=shadow, actor_id="admin")

    def on(d):
        return find_effective_price_plan(c, tenant_id=c.tenant_id, effective_on=d)

    assert on(date(2026, 8, 15)) is None
    assert on(date(2026, 10, 1))["price_list_version_id"] == sept
    assert on(date(2026, 10, 2))["in_force_to"] == date(2026, 10, 2)
    assert on(date(2026, 10, 3))["price_list_version_id"] == octo
    assert on(date(2026, 10, 15))["price_list_version_id"] == octo  # not the later segment list
    assert on(date(2026, 10, 15))["in_force_to"] is None


# ── review of PR 468: live data loaded by the old code ─────────────────────────
def _migration_0151():
    import importlib.util
    from pathlib import Path

    path = Path(__file__).resolve().parents[1] / "migrations" / "versions" / "0151_bolero_split_model_aliases.py"
    spec = importlib.util.spec_from_file_location("migration_0151", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    return module


def _legacy_bolero(c) -> dict[str, object]:
    """The catalogue and September book as the dev code left them: one model
    "THE BOSS- BOLERO+NEO" holding the Bolero, Neo and Neo Plus, a model-wide
    consumer scheme for the Neo, a variant-level one for the N10, a corporate
    one on the combined model, and a grid row for the Bolero."""
    from audit_core.discount_schemes import (
        create_discount_scheme,
        create_discount_scheme_version,
        publish_discount_scheme_version,
    )
    from audit_core.oem_price_masters import _insert_eligibility
    from audit_core.price_lists import (
        add_price_list_item,
        create_price_list,
        create_price_list_version,
        publish_price_list_version,
    )

    model = c.execute(text("INSERT INTO auditcore.product_models (oem_id, model_code, model_name) "
                           "VALUES (:o, 'THE_BOSS_BOLERO_NEO', 'THE BOSS- BOLERO+NEO') RETURNING model_id"),
                      {"o": c.oem_id}).scalar_one()
    price_list = create_price_list(c, tenant_id=c.tenant_id, code="OEM_NATIVE_PRICE_LIST", name="OEM")
    version = create_price_list_version(c, tenant_id=c.tenant_id, price_list_id=price_list, version_no=1,
                                        effective_from=date(2026, 9, 3))
    skus, variants = {}, {}
    for name, ex in (("B4 BS6.2 - Refresh", "848999.65"), ("N10 BS6.2 - R - Refresh", "1008999.65"),
                     ("P4 BS6.2", "1138999.58")):
        code = "_".join("".join(ch if ch.isalnum() else " " for ch in name.upper()).split()) + "__PRI"
        variants[name] = c.execute(text(
            "INSERT INTO auditcore.product_variants (model_id, variant_code, variant_name, attributes) "
            "VALUES (:m, :c, :n, CAST('{\"category\": \"PV\"}' AS jsonb)) RETURNING variant_id"),
            {"m": model, "c": code, "n": f"{name} (Private)"}).scalar_one()
        skus[name] = c.execute(text(
            "INSERT INTO auditcore.product_skus (oem_id, model_id, variant_id, sku_code, attributes) "
            "VALUES (:o, :m, :v, :c, CAST('{\"source\": \"OEM_NATIVE_PRICE_LIST\"}' AS jsonb)) RETURNING product_sku_id"),
            {"o": c.oem_id, "m": model, "v": variants[name], "c": f"THE_BOSS_BOLERO_NEO::{name}::PRI"}).scalar_one()
        add_price_list_item(c, tenant_id=c.tenant_id, price_list_version_id=version, product_sku_id=skus[name],
                            component_key="EX_SHOWROOM", standard_amount=Decimal(ex))
    publish_price_list_version(c, tenant_id=c.tenant_id, price_list_version_id=version, actor_id="dev")
    for scheme_name, variant in (("Mahindra Consumer — Bolero Neo", None),
                                 ("Mahindra Consumer — THE BOSS- BOLERO+NEO", "N10 BS6.2 - R - Refresh"),
                                 ("Mahindra Corporate Privilege Z — THE BOSS- BOLERO+NEO", None)):
        scheme = create_discount_scheme(c, tenant_id=c.tenant_id, code=f"S{uuid4().hex[:8]}", name=scheme_name)
        sv = create_discount_scheme_version(c, tenant_id=c.tenant_id, discount_scheme_id=scheme, version_no=1,
                                            effective_from=date(2026, 9, 3))
        _insert_eligibility(c, tenant_id=c.tenant_id, version_id=sv, model_id=model,
                            variant_id=variants[variant] if variant else None, customer_type=None,
                            criteria={"scope": "MODEL"} if variant is None else {"variantText": variant})
        publish_discount_scheme_version(c, tenant_id=c.tenant_id, discount_scheme_version_id=sv, actor_id="dev")
    grid = c.execute(text("INSERT INTO auditcore.dealer_discount_grid_versions (tenant_id, version_no, effective_from, "
                          "parameters) VALUES (:t, 1, :e, CAST('[]' AS jsonb)) RETURNING grid_version_id"),
                     {"t": c.tenant_id, "e": date(2026, 9, 1)}).scalar_one()
    c.execute(text("INSERT INTO auditcore.dealer_discount_grid_rows (tenant_id, grid_version_id, model_alias, model_id, "
                   "in_scope, booking_protection_days) VALUES (:t, :g, 'BOLERO', :m, true, 30)"),
              {"t": c.tenant_id, "g": grid, "m": model})
    return {"model": model, "skus": skus}


def _model_of(c, sku_id) -> tuple:
    return tuple(c.execute(text("SELECT s.sku_code, pm.model_name FROM auditcore.product_skus s JOIN auditcore.product_models pm "
                                "ON pm.model_id = s.model_id WHERE s.product_sku_id = :s"), {"s": sku_id}).one())


def test_the_bolero_split_keeps_every_live_sku_pinned_and_priced(fresh_oem) -> None:
    """Review of PR 468, B1: re-loading the books after the split must not
    orphan the September Bolero SKUs Journeys are pinned to. Migration 0151
    moves them in place (same id, same code) with their scheme and grid rows;
    the re-load then prices the same SKU ids in both months."""
    c = fresh_oem
    legacy = _legacy_bolero(c)
    counts = _migration_0151().split_bolero_family(c)
    assert counts["skus"] == 3 and counts["models"] == 3
    assert {_model_of(c, s) for s in legacy["skus"].values()} == {
        ("THE_BOSS_BOLERO_NEO::B4 BS6.2 - Refresh::PRI", "BOLERO"),
        ("THE_BOSS_BOLERO_NEO::N10 BS6.2 - R - Refresh::PRI", "BOLERO NEO"),
        ("THE_BOSS_BOLERO_NEO::P4 BS6.2::PRI", "BOLERO NEO PLUS"),
    }
    eligibility = sorted((r[0], r[1] is not None) for r in c.execute(text(
        "SELECT pm.model_name, e.variant_id FROM auditcore.discount_scheme_eligibility e "
        "JOIN auditcore.product_models pm ON pm.model_id = e.model_id WHERE pm.oem_id = :o"), {"o": c.oem_id}))
    # A model-wide row covered all three models before and still does,
    # whatever its scheme is named (re-review B3); the N10 row follows its
    # variant.
    assert eligibility == [
        ("BOLERO", False), ("BOLERO", False),
        ("BOLERO NEO", False), ("BOLERO NEO", False), ("BOLERO NEO", True),
        ("BOLERO NEO PLUS", False), ("BOLERO NEO PLUS", False),
    ]
    grid = sorted(c.execute(text("SELECT pm.model_name FROM auditcore.dealer_discount_grid_rows g JOIN auditcore.product_models pm "
                                 "ON pm.model_id = g.model_id WHERE pm.oem_id = :o"), {"o": c.oem_id}).scalars().all())
    assert grid == ["BOLERO", "BOLERO NEO", "BOLERO NEO PLUS"]  # the one grid row covered all three
    assert _migration_0151().split_bolero_family(c) == {
        "models": 0, "variants": 0, "merged": 0, "skus": 0, "eligibility": 0, "grid": 0}  # idempotent

    _load(c, date(2026, 9, 3), _book(
        ("PV", "THE BOSS- BOLERO+NEO", "B4 BS6.2 - Refresh", 848_999.65, "NEW BOLERO NEO (PVT)"),
        ("PV", "THE BOSS- BOLERO+NEO", "N10 BS6.2 - R - Refresh", 1_008_999.65, "NEW BOLERO NEO (PVT)"),
        ("PV", "THE BOSS- BOLERO+NEO", "P4 BS6.2", 1_138_999.58, "NEW BOLERO NEO (PVT)"),
    ))
    _load(c, date(2026, 10, 3), _book(
        ("PV", "Bolero", "B4 BS6.2 - Refresh", 856_998.57, "NEW BOLERO NEO (PVT)"),
        ("PV", "Bolero neo", "N10 BS6.2 - R - Refresh", 999_898.93, "NEW BOLERO NEO (PVT)"),
        ("PV", "Bolero Neo Plus", "P4 BS6.2", 1_145_998.99, "NEW BOLERO NEO (PVT)"),
    ))
    for on in (date(2026, 9, 15), date(2026, 10, 15)):
        plan = find_effective_price_plan(c, tenant_id=c.tenant_id, effective_on=on)
        priced = set(c.execute(text("SELECT DISTINCT product_sku_id FROM auditcore.price_list_items "
                                    "WHERE price_list_version_id = :v"), {"v": plan["price_list_version_id"]}).scalars())
        assert priced == set(legacy["skus"].values()), on


def test_the_old_three_variant_named_three_plus_is_renamed_on_reload(fresh_oem) -> None:
    """Review of PR 468, S2: the dev code filed the THREE B79 SKU (added on
    07 Sep) under the variant row created for THREE+, still named "THREE+".
    The re-load moves THREE+ to its own variant and gives the old row the
    THREE name, so each name has its own variant."""
    c = fresh_oem
    model = c.execute(text("INSERT INTO auditcore.product_models (oem_id, model_code, model_name) "
                           "VALUES (:o, 'BE6_SPORTEQ', 'BE6 SPORTEQ') RETURNING model_id"), {"o": c.oem_id}).scalar_one()
    legacy = c.execute(text("INSERT INTO auditcore.product_variants (model_id, variant_code, variant_name) "
                            "VALUES (:m, 'BE_6_SPORTEQ_THREE_B79_R19_NCH', 'BE 6 SPORTEQ THREE+ B79 R19 NCH') "
                            "RETURNING variant_id"), {"m": model}).scalar_one()
    ids = {}
    for name in ("BE 6 SPORTEQ THREE+ B79 R19 NCH", "BE 6 SPORTEQ THREE B79 R19 NCH"):
        ids[name] = c.execute(text("INSERT INTO auditcore.product_skus (oem_id, model_id, variant_id, sku_code, attributes) "
                                   "VALUES (:o, :m, :v, :c, CAST('{\"source\": \"OEM_NATIVE_PRICE_LIST\"}' AS jsonb)) "
                                   "RETURNING product_sku_id"),
                              {"o": c.oem_id, "m": model, "v": legacy, "c": f"BE6_SPORTEQ::{name}"}).scalar_one()
    _, skus, _ = _load(c, date(2026, 10, 3), _book(("BEV", "BE6 SPORTEQ", "BE 6 SPORTEQ THREE B79 R19 NCH", 2_394_999.50),
                                                  ("BEV", "BE6 SPORTEQ", "BE 6 SPORTEQ THREE+ B79 R19 NCH", 2_494_999.50)))
    assert set(skus.values()) == set(ids.values())
    named = dict(c.execute(text("SELECT s.product_sku_id, v.variant_name FROM auditcore.product_skus s "
                                "JOIN auditcore.product_variants v ON v.variant_id = s.variant_id WHERE s.model_id = :m"),
                           {"m": model}).all())
    assert named == {ids[n]: n for n in ids}


def test_an_unrelated_older_model_never_captures_a_vehicle(fresh_oem) -> None:
    """Review of PR 468, S3: a "Scorpio" created earlier by some other import
    (owning no price-list SKU) shares SCORPIO CLASSIC's identity only through
    an alias; an older "XUV 3XO" with no SKUs sits beside the "XUV3XO" that
    owns them. Neither captures the vehicle."""
    c = fresh_oem
    for code, name in (("SCORPIO", "Scorpio"), ("XUV_3XO_DECOY", "XUV 3XO")):
        c.execute(text("INSERT INTO auditcore.product_models (oem_id, model_code, model_name, created_at_utc) "
                       "VALUES (:o, :c, :n, now() - interval '1 year')"), {"o": c.oem_id, "c": code, "n": name})
    owner = c.execute(text("INSERT INTO auditcore.product_models (oem_id, model_code, model_name) "
                           "VALUES (:o, 'XUV3XO', 'XUV3XO') RETURNING model_id"), {"o": c.oem_id}).scalar_one()
    variant = c.execute(text("INSERT INTO auditcore.product_variants (model_id, variant_code, variant_name) "
                             "VALUES (:m, 'AX5_PM_MT', 'AX5 PM MT') RETURNING variant_id"), {"m": owner}).scalar_one()
    pinned = c.execute(text("INSERT INTO auditcore.product_skus (oem_id, model_id, variant_id, sku_code, attributes) "
                            "VALUES (:o, :m, :v, 'XUV3XO::AX5 PM MT', CAST('{\"source\": \"OEM_NATIVE_PRICE_LIST\"}' AS jsonb)) "
                            "RETURNING product_sku_id"), {"o": c.oem_id, "m": owner, "v": variant}).scalar_one()
    _, sept, _ = _load(c, date(2026, 9, 3), _book(("PV", "SCORPIO CLASSIC", "Classic S BS6.2 - E", 1_336_700.64),
                                                  ("PV", "XUV 3XO", "AX5 PM MT", 999_901.32)))
    _, octo, stats = _load(c, date(2026, 10, 3), _book(("PV", "SCORPIO", "Classic S BS6.2 - E", 1_343_200.67)))
    assert sept["XUV3XO::AX5 PM MT"] == pinned
    assert octo == {"SCORPIO_CLASSIC::Classic S BS6.2 - E": sept["SCORPIO_CLASSIC::Classic S BS6.2 - E"]}
    assert _model_of(c, octo["SCORPIO_CLASSIC::Classic S BS6.2 - E"])[1] == "SCORPIO CLASSIC"
    assert "modelIdentityConflicts" not in stats


def test_the_07_sep_added_ev_sheet_prices_its_additions_from_07_sep(fresh_oem) -> None:
    """Review of PR 468, S1: the "02 Nos Added ... 07th Sept'26" sheet keeps
    the month's "w.e.f. 01.09.2026" header. Dated by the header, its two
    added variants were priced 01-02 Sep only (the complete 03 Sep book does
    not list them). Dated by its name, they are priced from 07 Sep until
    the October book."""
    from openpyxl import load_workbook

    c = fresh_oem
    _load(c, date(2026, 9, 3), _book(("BEV", "BE6 SPORTEQ", "BE 6 SPORTEQ THREE+ B79 R19 NCH", 2_494_999.50)))
    wb = load_workbook(BytesIO(_dealer_ev_bytes({"BE6 SPORTEQ": [("BE 6 SPORTEQ THREE+ B79 R19 NCH", 2_494_999),
                                                              ("BE 6 SPORTEQ THREE B79 R19 NCH", 2_394_999)]})))
    wb["BE6 SPORTEQ"]["B1"] = "EV PRICE LIST w.e.f. 01.09.2026"
    buf = BytesIO()
    wb.save(buf)
    name = "02 Nos Added EV PRICE -07th  Sept'26.xlsx"
    parsed = parse_price_list(buf.getvalue(), filename=name)
    assert parsed.effective_from_hint == date(2026, 9, 7) and parsed.meta["effectiveFromSource"] == "FILENAME"
    assert "the later date, 2026-09-07, is applied" in parsed.warnings[0]
    _, added, _ = _load(c, parsed.effective_from_hint, buf.getvalue(), filename=name)
    three = added["BE6_SPORTEQ::BE 6 SPORTEQ THREE B79 R19 NCH"]
    _load(c, date(2026, 10, 3), _book(("BEV", "BE6 SPORTEQ", "BE 6 SPORTEQ THREE+ B79 R19 NCH", 2_494_999.50)))

    def priced(on):
        plan = find_effective_price_plan(c, tenant_id=c.tenant_id, effective_on=on)
        return three in set(c.execute(text("SELECT product_sku_id FROM auditcore.price_list_items WHERE "
                                           "price_list_version_id = :v"), {"v": plan["price_list_version_id"]}).scalars())

    assert [priced(d) for d in (date(2026, 9, 6), date(2026, 9, 7), date(2026, 10, 2), date(2026, 10, 3))] == \
        [False, True, True, False]


def test_a_refused_upload_is_recorded_and_a_preview_is_labelled(connection) -> None:
    """Review of PR 468, S8: a preview with errors is kept as FAILED with
    "PREVIEW:" in front of its detail; a refused publish is kept as FAILED
    (written in its own transaction, so it survives the request's rollback)."""
    import asyncio
    from types import SimpleNamespace

    from starlette.datastructures import UploadFile

    from audit_core.errors import ValidationError as AuditValidationError
    from audit_core.oem_price_masters import upload_oem_master

    tenant_id = connection.tenant_id
    wb = Workbook()
    ws = wb.active
    ws.title = "Price List"
    ws.append(_HEADER)
    ws.append([1, "PV", "THAR ROXX", "AX7 DMT 2WD", "AX7", "DIESEL", "MT", "2WD", "5", 1_000_000, 100, 200,
               "see note", 0, 0, 0, 0, 5000, 1_005_300, 5000, 1_005_300, "S"])
    buf = BytesIO()
    wb.save(buf)
    name = f"bad-{uuid4().hex[:6]}.xlsx"

    def call(dry_run: bool):
        return asyncio.run(upload_oem_master(
            admin_request=SimpleNamespace(user_id="admin"), connection=connection, tenant_id=tenant_id,
            master_kind="PRICE_LIST", file=UploadFile(file=BytesIO(buf.getvalue()), filename=name),
            effective_from=date(2026, 10, 3), dry_run=dry_run,
        ))

    preview = call(True)
    assert preview.status == "PREVIEW" and preview.errors
    status, detail = connection.execute(text("SELECT status, error_detail FROM auditcore.oem_master_uploads "
                                             "WHERE tenant_id = :t AND source_filename = :n"),
                                        {"t": tenant_id, "n": name}).one()
    assert status == "FAILED" and detail.startswith("PREVIEW: 1 unreconciled row(s): row 2: non-numeric")

    with pytest.raises(AuditValidationError):
        call(False)
    engine = create_engine(os.environ["DATABASE_URL"])
    with engine.begin() as other:  # another transaction sees the refused publish
        other.execute(text("SELECT set_config('app.tenant_id', :t, true)"), {"t": tenant_id})
        rows = other.execute(text("SELECT error_detail FROM auditcore.oem_master_uploads WHERE tenant_id = :t "
                                  "AND source_filename = :n AND error_detail NOT LIKE 'PREVIEW:%'"),
                             {"t": tenant_id, "n": name}).scalars().all()
    engine.dispose()
    assert len(rows) == 1 and rows[0].startswith("1 unreconciled row(s)")


_SEPT_EXCHANGE = "/mnt/user-data/uploads/mahindra/PriceMastersMahindra/Exchange scheme Sept'26 - Bhubaneswar (2) (2).pdf"


@pytest.mark.skipif(not os.path.exists(_SEPT_EXCHANGE), reason="the real September exchange bulletin is not on this machine")
def test_the_bolero_split_keeps_the_september_exchange_and_scrappage_for_all_three(fresh_oem) -> None:
    """Re-review of PR 468, B3: the September exchange bulletin prints the
    Bolero, Bolero Neo and Bolero Neo Plus rows (exchange 10,000, scrappage
    dealer 20,000, scrappage COD 10,000). Loaded the way the dev code did,
    against the combined model, each set became one scheme named after its
    first row ("— Bolero") with model-wide rows covering all three. After
    0151 the Neo and the Neo Plus must keep every one of them."""
    from audit_core.oem_master_parsers import parse_master
    from audit_core.oem_price_masters import ingest_discount_document

    c = fresh_oem
    _legacy_bolero(c)
    oem_code = c.execute(text("SELECT oem_code FROM auditcore.oems WHERE oem_id = :o"), {"o": c.oem_id}).scalar_one()
    # The alias table as it stood before 0151: every Bolero spelling named
    # the combined model.
    c.execute(text("""
        INSERT INTO auditcore.oem_model_aliases (oem_code, alias_text, canonical_model_name)
        SELECT :code, alias_text,
               CASE WHEN upper(alias_text) LIKE 'BOL%' AND upper(alias_text) NOT LIKE '%MAXX%'
                         AND upper(alias_text) NOT LIKE '%PIK%' AND upper(alias_text) NOT LIKE '%CAMPER%'
                         AND upper(alias_text) NOT LIKE '%PICK%'
                    THEN 'THE BOSS- BOLERO+NEO' ELSE canonical_model_name END
        FROM auditcore.oem_model_aliases WHERE oem_code = 'MAHINDRA'"""), {"code": oem_code})
    from pathlib import Path

    parsed = parse_master("EXCHANGE_SCHEME", Path(_SEPT_EXCHANGE).read_bytes(), filename=os.path.basename(_SEPT_EXCHANGE))
    ingest_discount_document(c, tenant_id=c.tenant_id, oem_id=c.oem_id, oem_code=oem_code, master_kind="EXCHANGE_SCHEME",
                             effective_from=date(2026, 9, 3), parsed=parsed, actor_id="dev")

    def sections() -> dict[str, set[tuple[str, str]]]:
        out: dict[str, set[tuple[str, str]]] = {}
        for model, section, benefit, amount in c.execute(text("""
            SELECT pm.model_name, e.criteria->>'section', b.benefit_key, b.amount_value::text
            FROM auditcore.discount_scheme_eligibility e
            JOIN auditcore.product_models pm ON pm.model_id = e.model_id
            JOIN auditcore.discount_scheme_benefits b
              ON b.tenant_id = e.tenant_id AND b.discount_scheme_version_id = e.discount_scheme_version_id
            WHERE pm.oem_id = :o"""), {"o": c.oem_id}).all():
            out.setdefault(model, set()).add((section, f"{benefit} {Decimal(amount):.0f}"))
        return out

    before = sections()["THE BOSS- BOLERO+NEO"]
    assert {("EXCHANGE_PERSONAL", "EXCHANGE_BONUS 10000"), ("SCRAPPAGE_DEALER", "SCRAPPAGE_BONUS_DEALER 20000"),
            ("SCRAPPAGE_COD", "SCRAPPAGE_BONUS_COD 10000")} <= before
    _migration_0151().split_bolero_family(c)
    after = sections()
    assert "THE BOSS- BOLERO+NEO" not in after
    for model in ("BOLERO", "BOLERO NEO", "BOLERO NEO PLUS"):
        assert after[model] == before, model


def _scheme_rows(c, code_like: str) -> list[tuple]:
    return sorted(tuple(r) for r in c.execute(text("""
        SELECT s.scheme_code, pm.model_name, COALESCE(pv.variant_name, '-'), e.criteria ->> 'scope',
               v.effective_to
        FROM auditcore.discount_schemes s
        JOIN auditcore.discount_scheme_versions v ON v.discount_scheme_id = s.discount_scheme_id
        JOIN auditcore.discount_scheme_eligibility e ON e.discount_scheme_version_id = v.discount_scheme_version_id
        JOIN auditcore.product_models pm ON pm.model_id = e.model_id
        LEFT JOIN auditcore.product_variants pv ON pv.variant_id = e.variant_id
        WHERE s.tenant_id = :t AND s.scheme_code LIKE :c"""), {"t": c.tenant_id, "c": code_like}).all())


def test_scheme_rows_reach_the_variants_they_name_through_model_identity(fresh_oem) -> None:
    """Review of PR 472 (B1, B4, S1, S4), on #468's model identity: a variant
    named without the emission norm, the "- E" suffix or with "& Mocha
    Interiors" is the price list's variant; "Scorpio-N including Refresh"
    reaches the old and the new Scorpio N, a plain "Scorpio-N" in September
    the Scorpio N on sale (the old one); the bulletin's validity end is the
    versions' effective_to."""
    c = fresh_oem
    oem_code = c.execute(text("SELECT oem_code FROM auditcore.oems WHERE oem_id = :o"), {"o": c.oem_id}).scalar_one()
    _load(c, date(2026, 9, 3), _book(
        ("PV", "SCORPIO N", "Z8 S D AT 2WD 7 STR BS6.2 - N", 1_917_499.58, "SCORPIO N D old"),
        ("PV", "New SCORPIO N", "Z8 S D AT 2WD 7 STR BS6.2 - Refresh", 1_988_996.00, "SCORPIO N D New"),
        ("PV", "SCORPIO CLASSIC", "Classic S BS6.2 - E", 1_336_700.64),
        ("CV", "MAXX HD", "MAXX HD 1.3 LX BS6.2", 900_000),
        ("PV", "THAR ROXX", "AX7L DAT 2WD", 2_000_000),
    ))
    consumer = ParseResult(kind="CONSUMER_SCHEME")
    consumer.effective_to_hint = date(2026, 9, 30)
    consumer.discount_rows = [
        DiscountRow(1, "CONSUMER", "Scorpio Classic", ["Classic S BS6.2"], [("CASH_DISCOUNT", Decimal(7001))],
                    Decimal(7001), {}),
        DiscountRow(2, "CONSUMER", "MAXX HD", ["MAXX HD 1.3 LX"], [("CASH_DISCOUNT", Decimal(15000))],
                    Decimal(15000), {}),
        DiscountRow(3, "CONSUMER", "Thar Roxx", ["AX7L DAT 2WD & Mocha Interiors"],
                    [("ACCESSORIES_KIT", Decimal(50001))], Decimal(50001), {}),
    ]
    summary = ingest_discount_document(c, tenant_id=c.tenant_id, oem_id=c.oem_id, oem_code=oem_code,
                                       master_kind="CONSUMER_SCHEME", effective_from=date(2026, 9, 3),
                                       parsed=consumer, actor_id="admin")
    assert summary["unresolvedVariants"] == []
    rows = _scheme_rows(c, "MAH_CONSUMER_%")
    assert {(r[1], r[2], r[3]) for r in rows} == {
        ("SCORPIO CLASSIC", "Classic S BS6.2 - E", "VARIANT"),
        ("MAXX HD", "MAXX HD 1.3 LX BS6.2", "VARIANT"),
        ("THAR ROXX", "AX7L DAT 2WD", "VARIANT"),
    }
    assert {r[4] for r in rows} == {date(2026, 9, 30)}

    exchange = ParseResult(kind="EXCHANGE_SCHEME")
    exchange.discount_rows = [
        DiscountRow(1, "EXCHANGE", "Scorpio-N including Refresh", [], [("EXCHANGE_BONUS", Decimal(50000))],
                    Decimal(50000), {"section": "EXCHANGE_PERSONAL"}),
        DiscountRow(2, "SCRAPPAGE", "Scorpio-N", [], [("SCRAPPAGE_BONUS_DEALER", Decimal(25000))],
                    Decimal(25000), {"section": "SCRAPPAGE"}),
    ]
    ingest_discount_document(c, tenant_id=c.tenant_id, oem_id=c.oem_id, oem_code=oem_code,
                             master_kind="EXCHANGE_SCHEME", effective_from=date(2026, 9, 3), parsed=exchange,
                             actor_id="admin")
    exchange_models = {r[1] for r in _scheme_rows(c, "MAH_EXCHANGE_%")}
    assert exchange_models == {"SCORPIO N", "NEW SCORPIO N"}
    assert {r[1] for r in _scheme_rows(c, "MAH_SCRAPPAGE_%")} == {"SCORPIO N"}


def test_the_grid_reads_its_own_spellings_and_a_pick_up_never_takes_the_camper_line(fresh_oem) -> None:
    from audit_core.oem_master_parsers import GridRow
    from audit_core.oem_price_masters import grid_row_for_model, ingest_discount_grid

    c = fresh_oem
    oem_code = c.execute(text("SELECT oem_code FROM auditcore.oems WHERE oem_id = :o"), {"o": c.oem_id}).scalar_one()
    _load(c, date(2026, 9, 1), _book(
        ("PV", "XUV3XO", "AX5 PM MT", 1_029_900.54), ("CV", "MAXX CITY", "MAXX CITY 1.5 LX BS6.2", 839_501.15),
        ("CV", "PICK UP", "BOL CAMPER 2WD BS6.2 - Refresh", 950_000),
        ("CV", "PICK UP", "BOL PIK UP FB 1.3T - BS6.2 - PS", 850_000),
    ))
    grid = ParseResult(kind="DISCOUNT_GRID")
    grid.grid_rows = [
        GridRow(row_no=1, model_alias="3XO", model_aliases=["3XO"], in_scope=True, booking_protection_days=None,
                agreed_buffer_amount=Decimal(7000), insurance_od_percent=None, out_of_territory_amount=None, raw={}),
        GridRow(row_no=2, model_alias="MAXX", model_aliases=["MAXX"], in_scope=True, booking_protection_days=None,
                agreed_buffer_amount=Decimal(5000), insurance_od_percent=None, out_of_territory_amount=None, raw={}),
        GridRow(row_no=3, model_alias="Camper", model_aliases=["Camper"], in_scope=True, booking_protection_days=None,
                agreed_buffer_amount=Decimal(0), insurance_od_percent=None, out_of_territory_amount=None, raw={}),
        GridRow(row_no=4, model_alias="Pick-Up", model_aliases=["Pick-Up"], in_scope=True,
                booking_protection_days=None, agreed_buffer_amount=Decimal(7000), insurance_od_percent=None,
                out_of_territory_amount=None, raw={}),
    ]
    summary = ingest_discount_grid(c, tenant_id=c.tenant_id, oem_id=c.oem_id, oem_code=oem_code,
                                   effective_from=date(2026, 9, 1), parsed=grid, upload_id=uuid4(), actor_id="admin")
    assert summary["unresolved"] == []
    models = dict(c.execute(text("SELECT model_name, model_id FROM auditcore.product_models WHERE oem_id = :o"),
                            {"o": c.oem_id}).all())
    on = date(2026, 9, 15)
    assert grid_row_for_model(c, tenant_id=c.tenant_id, model_id=models["XUV3XO"], effective_on=on,
                              variant_name="AX5 PM MT")["agreed_buffer_amount"] == Decimal(7000)
    assert grid_row_for_model(c, tenant_id=c.tenant_id, model_id=models["MAXX CITY"], effective_on=on,
                              variant_name="MAXX CITY 1.5 LX BS6.2")["agreed_buffer_amount"] == Decimal(5000)
    pick_up = grid_row_for_model(c, tenant_id=c.tenant_id, model_id=models["PICK UP"], effective_on=on,
                                 variant_name="BOL PIK UP FB 1.3T - BS6.2 - PS")
    assert pick_up["model_alias"] == "Pick-Up" and pick_up["agreed_buffer_amount"] == Decimal(7000)


def test_a_grid_dated_before_its_book_reads_its_models_against_the_book_it_is_used_with(fresh_oem) -> None:
    """Review of PR 472: the October grid is dated 1 Oct but the October book
    starts 3 Oct, and the book of 1 Oct still prices the old Scorpio N. The
    grid's "SCORPIO-N" is the new Scorpio N, which only the October book
    carries: its line must reach NEW SCORPIO N (it was bound to the retired
    SCORPIO N, so a deal on 15 Oct had no grid line), and a check must say
    so when a line is bound to a model its book does not price."""
    from audit_core.oem_master_parsers import GridRow
    from audit_core.oem_price_masters import (
        grid_lines_unbound_to_books,
        grid_row_for_model,
        ingest_discount_grid,
    )

    c = fresh_oem
    oem_code = c.execute(text("SELECT oem_code FROM auditcore.oems WHERE oem_id = :o"), {"o": c.oem_id}).scalar_one()
    _load(c, date(2026, 9, 3), _book(("PV", "XUV3XO", "AX5 PM MT", 1_029_900.54),
                                     ("PV", "SCORPIO N", "Z8L D AT 2WD 7 STR BS6.2", 1_917_499.58)))
    _load(c, date(2026, 10, 3), _book(("PV", "XUV 3XO", "AX5 PM MT", 999_901.32),
                                      ("PV", "SCORPIO-N", "Z8L D AT 2WD 7 STR BS6.2 - Refresh", 1_988_996.00)))
    models = dict(c.execute(text("SELECT model_name, model_id FROM auditcore.product_models WHERE oem_id = :o"),
                            {"o": c.oem_id}).all())
    assert {"SCORPIO N", "NEW SCORPIO N", "XUV3XO"} <= set(models)

    def line(no: int, alias: str, buffer: int) -> GridRow:
        return GridRow(row_no=no, model_alias=alias, model_aliases=[alias], in_scope=True, booking_protection_days=None,
                       agreed_buffer_amount=Decimal(buffer), insurance_od_percent=None, out_of_territory_amount=None,
                       raw={})

    grid = ParseResult(kind="DISCOUNT_GRID")
    grid.grid_rows = [line(1, "3XO", 7000), line(2, "SCORPIO-N", 12000)]
    summary = ingest_discount_grid(c, tenant_id=c.tenant_id, oem_id=c.oem_id, oem_code=oem_code,
                                   effective_from=date(2026, 10, 1), parsed=grid, upload_id=uuid4(), actor_id="admin")
    assert summary["unresolved"] == []

    def buffer_for(model: str, on: date):
        row = grid_row_for_model(c, tenant_id=c.tenant_id, model_id=models[model], effective_on=on)
        return row["agreed_buffer_amount"] if row else None

    assert buffer_for("NEW SCORPIO N", date(2026, 10, 15)) == Decimal(12000)
    assert buffer_for("XUV3XO", date(2026, 10, 15)) == Decimal(7000)
    # The book of 1 - 2 Oct still prices the old Scorpio N, but the October
    # grid's "SCORPIO-N" is the new one only: the old series takes the
    # previous month's line (product owner, 2026-11-02), none in this test.
    assert buffer_for("SCORPIO N", date(2026, 10, 2)) is None
    version = c.execute(text("SELECT grid_version_id FROM auditcore.dealer_discount_grid_versions "
                             "WHERE tenant_id = :t"), {"t": c.tenant_id}).scalar_one()
    assert grid_lines_unbound_to_books(c, tenant_id=c.tenant_id, oem_id=c.oem_id, oem_code=oem_code,
                                       grid_version_id=version) == []

    # A line bound only to the retired model (how it was stored) is reported.
    c.execute(text("UPDATE auditcore.dealer_discount_grid_rows SET model_id = :old WHERE grid_version_id = :v "
                   "AND model_alias = 'SCORPIO-N' AND model_id = :m"),
              {"v": version, "m": models["NEW SCORPIO N"], "old": models["SCORPIO N"]})
    gaps = grid_lines_unbound_to_books(c, tenant_id=c.tenant_id, oem_id=c.oem_id, oem_code=oem_code,
                                       grid_version_id=version)
    assert len(gaps) == 1 and "'SCORPIO-N'" in gaps[0] and "NEW SCORPIO N" in gaps[0] and "SCORPIO N'" in gaps[0]
    assert buffer_for("NEW SCORPIO N", date(2026, 10, 15)) is None
