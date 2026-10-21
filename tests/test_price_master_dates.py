"""A Journey is priced from the price master in force on its ONE pricing date
(2026-10-20 product owner rules, R5). Requires DATABASE_URL.

Real figures: MAXX CITY 1.5 LX BS6.2 ex-showroom 8,39,501.15 in the
September book (from 03 Sep 2026) and 8,52,000.62 in the October book (from
03 Oct 2026).
"""
from __future__ import annotations

import os
from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, text

from audit_core.errors import AuditCoreError
from audit_core.oem_master_parsers import parse_price_list
from audit_core.oem_price_masters import ingest_price_list
from audit_core.uc03_sku_candidates import (
    PRICE_MASTER_NOT_IN_FORCE_TASK,
    SKU_NOT_IN_PRICE_MASTER_TASK,
    _price_plan_for_journey,
    journey_price_plan,
    journey_pricing_date,
    sync_price_master_in_force_task,
)

_HEADER = [
    "Sl. No.", "Category", "Model", "Variant", "Trim", "Fuel", "Transmission",
    "Drive", "Seater", "Ex-Showroom Price", "TCS", "Insurance",
    "Ext. Warranty (4th Yr)", "Ext. Warranty (4th & 5th Yr)", "Accessories Kit",
    "RSA (1 Yr)", "FASTag", "Registration\n(Individual / w/o Hyp.)",
    "On-Road Price\n(Individual / w/o Hyp.)", "Registration\n(Corporate / with Hyp.)",
    "On-Road Price\n(Corporate / with Hyp.)", "Source Sheet",
]


def _book(ex: float) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Price List"
    ws.append(_HEADER)
    tcs, ins, ew4, ew45, acc, rsa, fastag, reg = 8395.02, 40000, 0, 0, 5000, 0, 500, 3632
    onroad = ex + tcs + ins + ew4 + ew45 + acc + rsa + fastag + reg
    ws.append([1, "CV", "MAXX CITY", "MAXX CITY 1.5 LX BS6.2", "LX", "DIESEL", "MT", "2WD", "2",
               ex, tcs, ins, ew4, ew45, acc, rsa, fastag, reg, onroad, reg, onroad, "MAXX CITY"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def journey():
    """A tenant with its own OEM, the September and October books, and one
    Journey with a booking; all in one transaction, rolled back."""
    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        pytest.skip("DATABASE_URL is required for price-master date tests")
    engine = create_engine(database_url)
    conn = engine.connect()
    tx = conn.begin()
    s = uuid4().hex[:10]
    t = f"tenant-pmd-{s}"
    try:
        oem_id = conn.execute(text("INSERT INTO auditcore.oems (oem_code, oem_name) VALUES (:c,'O') RETURNING oem_id"),
                              {"c": f"PMD{s}"}).scalar_one()
        cat = conn.execute(text("INSERT INTO auditcore.product_categories (category_code, category_name) "
                                "VALUES (:c,'V') RETURNING product_category_id"), {"c": f"PMD-{s}"}).scalar_one()
        conn.execute(text("""INSERT INTO auditcore.projects (tenant_id, project_code, project_name, oem_id,
                             product_category_id, effective_start_date, timezone_name, project_status)
                             VALUES (:t, :pc, 'P', :o, :cat, CURRENT_DATE - 90, 'Asia/Kolkata', 'ACTIVE')"""),
                     {"t": t, "pc": f"PMD-{s}", "o": oem_id, "cat": cat})
        conn.execute(text("SELECT set_config('app.platform_super_admin','true',true)"))
        conn.execute(text("SELECT set_config('app.tenant_id',:t,true)"), {"t": t})
        dealer = conn.execute(text("INSERT INTO auditcore.dealers (tenant_id, dealer_code, dealer_name) "
                                   "VALUES (:t,:c,'D') RETURNING dealer_id"), {"t": t, "c": f"D{s}"}).scalar_one()
        outlet = conn.execute(text("INSERT INTO auditcore.dealer_outlets (tenant_id, dealer_id, outlet_code, outlet_name) "
                                   "VALUES (:t,:d,:c,'O') RETURNING outlet_id"),
                              {"t": t, "d": dealer, "c": f"O{s}"}).scalar_one()
        customer = conn.execute(text("""INSERT INTO auditcore.customers (tenant_id, dealer_id, outlet_id,
                                        customer_type_code, display_name) VALUES (:t,:d,:o,'INDIVIDUAL','C')
                                        RETURNING customer_id"""), {"t": t, "d": dealer, "o": outlet}).scalar_one()
        journey_id = conn.execute(text("""INSERT INTO auditcore.journeys (tenant_id, dealer_id, outlet_id, customer_id,
                                          journey_reference) VALUES (:t,:d,:o,:c,:r) RETURNING journey_id"""),
                                  {"t": t, "d": dealer, "o": outlet, "c": customer, "r": f"J{s}"}).scalar_one()
        conn.execute(text("INSERT INTO auditcore.bookings (tenant_id, journey_id, booking_date) VALUES (:t,:j,NULL)"),
                     {"t": t, "j": journey_id})
        skus = {}
        for on, ex in ((date(2026, 9, 3), 839_501.15), (date(2026, 10, 3), 852_000.62)):
            _, skus = ingest_price_list(conn, tenant_id=t, oem_id=oem_id, effective_from=on,
                                        parsed=parse_price_list(_book(ex)), actor_id="admin")
        conn.tenant_id = t  # type: ignore[attr-defined]
        conn.journey_id = journey_id  # type: ignore[attr-defined]
        conn.sku_id = next(iter(skus.values()))  # type: ignore[attr-defined]
        yield conn
    finally:
        tx.rollback()
        conn.close()
        engine.dispose()


def _set(c, **values) -> None:
    sets = ", ".join(f"{k} = :{k}" for k in values)
    c.execute(text(f"UPDATE auditcore.bookings SET {sets} WHERE tenant_id=:t AND journey_id=:j"),
              {**values, "t": c.tenant_id, "j": c.journey_id})


def _task(c, kind=PRICE_MASTER_NOT_IN_FORCE_TASK):
    return c.execute(text("SELECT task_status, title, assigned_role_code FROM auditcore.p2_tasks WHERE tenant_id=:t "
                          "AND journey_id=:j AND task_type=:k"),
                     {"t": c.tenant_id, "j": c.journey_id, "k": kind}).mappings().one_or_none()


def test_the_pricing_date_is_the_applied_date_else_the_booking_date_never_today(journey) -> None:
    c = journey
    assert journey_pricing_date(c, tenant_id=c.tenant_id, journey_id=c.journey_id) is None
    _set(c, booking_date=date(2026, 9, 20))
    assert journey_pricing_date(c, tenant_id=c.tenant_id, journey_id=c.journey_id) == date(2026, 9, 20)
    _set(c, pricing_effective_on=date(2026, 10, 5))
    assert journey_pricing_date(c, tenant_id=c.tenant_id, journey_id=c.journey_id) == date(2026, 10, 5)


@pytest.mark.parametrize(
    ("on", "version", "ex"),
    [(date(2026, 10, 1), 1, "839501.15"), (date(2026, 10, 3), 2, "852000.62"), (date(2026, 9, 3), 1, "839501.15")],
)
def test_the_book_in_force_on_the_pricing_date_prices_the_deal(journey, on, version, ex) -> None:
    c = journey
    _set(c, booking_date=on)
    found = journey_price_plan(c, tenant_id=c.tenant_id, journey_id=c.journey_id)
    assert found["onDate"] == on and found["plan"]["version_no"] == version and found["reason"] is None
    amount = c.execute(text("SELECT standard_amount FROM auditcore.price_list_items WHERE price_list_version_id=:v "
                            "AND product_sku_id=:s AND component_key='EX_SHOWROOM'"),
                       {"v": found["plan"]["price_list_version_id"], "s": c.sku_id}).scalar_one()
    assert amount == Decimal(ex)


def test_nothing_in_force_is_said_not_priced_on_the_newest_book(journey) -> None:
    """15 Aug 2026 is before the first book. It used to take the newest
    published book (October) "regardless of its effective window"."""
    c = journey
    _set(c, booking_date=date(2026, 8, 15))
    with pytest.raises(AuditCoreError) as raised:
        _price_plan_for_journey(c, tenant_id=c.tenant_id, journey_id=c.journey_id, effective_on=date(2026, 8, 15))
    assert raised.value.detail == "No price master in force on 2026-08-15"
    found = journey_price_plan(c, tenant_id=c.tenant_id, journey_id=c.journey_id)
    assert found["plan"] is None and found["reason"] == "No price master in force on 2026-08-15"


def test_a_task_says_no_price_master_is_in_force_and_closes_once_one_is(journey) -> None:
    c = journey
    _set(c, booking_date=date(2026, 8, 15))
    sync_price_master_in_force_task(c, tenant_id=c.tenant_id, journey_id=c.journey_id)
    task = _task(c)
    assert task["title"] == "No price master in force on 2026-08-15" and task["task_status"] != "VERIFIED_COMPLETE"
    _set(c, booking_date=date(2026, 9, 15))  # the PC corrects the date (or the August book is loaded)
    sync_price_master_in_force_task(c, tenant_id=c.tenant_id, journey_id=c.journey_id)
    assert _task(c)["task_status"] == "VERIFIED_COMPLETE"


def test_the_sku_pricing_panel_reads_the_book_of_the_pricing_date(journey) -> None:
    """It used to read whichever published book was newest: a September deal
    showed October figures once October was loaded."""
    from audit_core.uc03_journey_overview_projection import _sku_pricing_panel

    c = journey
    c.execute(text("INSERT INTO auditcore.journey_products (tenant_id, journey_id, product_sku_id, selection_status) "
                   "VALUES (:t, :j, :s, 'TENTATIVE')"), {"t": c.tenant_id, "j": c.journey_id, "s": c.sku_id})
    _set(c, booking_date=date(2026, 9, 20))
    panel = _sku_pricing_panel(c, tenant_id=c.tenant_id, journey_id=c.journey_id, reviewed_booking={})
    ex = next(m for m in panel["masterComponents"] if m["componentKey"] == "EX_SHOWROOM")
    assert ex["masterAmount"] == 839501.15
    assert panel["priceListVersionNo"] == 1 and panel["pricingDate"] == "2026-09-20"
    _set(c, booking_date=date(2026, 10, 3))
    panel = _sku_pricing_panel(c, tenant_id=c.tenant_id, journey_id=c.journey_id, reviewed_booking={})
    assert next(m for m in panel["masterComponents"] if m["componentKey"] == "EX_SHOWROOM")["masterAmount"] == 852000.62
    _set(c, booking_date=date(2026, 8, 15))
    assert _sku_pricing_panel(c, tenant_id=c.tenant_id, journey_id=c.journey_id, reviewed_booking={}) is None


def test_the_task_closes_when_the_pricing_date_is_cleared(journey) -> None:
    """Review of PR 468, S5: no pricing date any more, nothing to say about a
    book for it (PC_BOOKING_DATE_MISSING asks for the date)."""
    c = journey
    _set(c, booking_date=date(2026, 8, 15))
    sync_price_master_in_force_task(c, tenant_id=c.tenant_id, journey_id=c.journey_id)
    _set(c, booking_date=None)
    sync_price_master_in_force_task(c, tenant_id=c.tenant_id, journey_id=c.journey_id)
    assert _task(c)["task_status"] == "VERIFIED_COMPLETE"


def test_publishing_the_missing_book_closes_the_task_without_a_journey_event(journey) -> None:
    """Review of PR 468, S5: the task is raised for a Journey with no SKU,
    and loading a book that covers its date closes it at once."""
    c = journey
    _set(c, booking_date=date(2026, 8, 20))
    sync_price_master_in_force_task(c, tenant_id=c.tenant_id, journey_id=c.journey_id)
    assert _task(c)["task_status"] != "VERIFIED_COMPLETE"
    oem_id = c.execute(text("SELECT oem_id FROM auditcore.projects WHERE tenant_id=:t"), {"t": c.tenant_id}).scalar_one()
    stats: dict = {}
    ingest_price_list(c, tenant_id=c.tenant_id, oem_id=oem_id, effective_from=date(2026, 8, 1),
                      parsed=parse_price_list(_book(830_000.00)), actor_id="admin", stats=stats)
    assert stats["journeysRechecked"] == 1
    assert _task(c)["task_status"] == "VERIFIED_COMPLETE"


def test_a_pinned_sku_withdrawn_from_the_book_raises_a_pc_task_and_clears_its_standard(journey) -> None:
    """Review of PR 468, S4: a complete book that no longer lists the SKU
    (the old Scorpio N in October) must say so: a PC task, and no standard
    from an older book left standing."""
    from audit_core.uc03_deal_reconciliation import sync_deal_reconciliation

    c = journey
    oem_id = c.execute(text("SELECT oem_id FROM auditcore.projects WHERE tenant_id=:t"), {"t": c.tenant_id}).scalar_one()
    c.execute(text("INSERT INTO auditcore.journey_products (tenant_id, journey_id, product_sku_id, selection_status) "
                   "VALUES (:t, :j, :s, 'CONFIRMED')"), {"t": c.tenant_id, "j": c.journey_id, "s": c.sku_id})
    _set(c, booking_date=date(2026, 9, 20))
    sync_deal_reconciliation(c, tenant_id=c.tenant_id, journey_id=c.journey_id, correlation_id="")
    standard = c.execute(text("SELECT standard_amount FROM auditcore.commercial_lines WHERE tenant_id=:t AND journey_id=:j "
                              "AND component_key='ex_showroom_price'"), {"t": c.tenant_id, "j": c.journey_id}).scalar_one()
    assert standard == Decimal("839501.15")

    # A November book that no longer lists the vehicle.
    wb_rows = _book(900_000.00)
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(wb_rows))
    wb.active["C2"] = "MAXX HD"
    buf = BytesIO()
    wb.save(buf)
    ingest_price_list(c, tenant_id=c.tenant_id, oem_id=oem_id, effective_from=date(2026, 11, 1),
                      parsed=parse_price_list(buf.getvalue()), actor_id="admin")
    _set(c, pricing_effective_on=date(2026, 11, 5))  # the deal priced on its (later) invoice date
    result = sync_deal_reconciliation(c, tenant_id=c.tenant_id, journey_id=c.journey_id, correlation_id="")
    assert result["skuNotInPriceList"] is True and "error" not in result  # discounts still refreshed
    task = _task(c, SKU_NOT_IN_PRICE_MASTER_TASK)
    assert task["assigned_role_code"] == "PC" and task["task_status"] != "VERIFIED_COMPLETE"
    assert task["title"] == "SKU MAXX_CITY::MAXX CITY 1.5 LX BS6.2 is not in the price master for 2026-11-05 — set the SKU"
    standard = c.execute(text("SELECT standard_amount, price_list_item_id FROM auditcore.commercial_lines WHERE tenant_id=:t "
                              "AND journey_id=:j AND component_key='ex_showroom_price'"),
                         {"t": c.tenant_id, "j": c.journey_id}).one()
    assert tuple(standard) == (None, None)

    _set(c, pricing_effective_on=date(2026, 10, 20))  # back on a date whose book lists it
    sync_deal_reconciliation(c, tenant_id=c.tenant_id, journey_id=c.journey_id, correlation_id="")
    assert _task(c, SKU_NOT_IN_PRICE_MASTER_TASK)["task_status"] == "VERIFIED_COMPLETE"


def test_a_page_view_never_raises_the_withdrawn_sku_task(journey) -> None:
    """Re-review of PR 468, S10: the overview self-heal (raise_findings=False)
    neither raises tasks nor rewrites lines."""
    from audit_core.uc03_deal_reconciliation import sync_deal_reconciliation

    c = journey
    c.execute(text("INSERT INTO auditcore.journey_products (tenant_id, journey_id, product_sku_id, selection_status) "
                   "VALUES (:t, :j, :s, 'CONFIRMED')"), {"t": c.tenant_id, "j": c.journey_id, "s": c.sku_id})
    _set(c, pricing_effective_on=date(2026, 9, 20))
    sync_deal_reconciliation(c, tenant_id=c.tenant_id, journey_id=c.journey_id, correlation_id="")
    oem_id = c.execute(text("SELECT oem_id FROM auditcore.projects WHERE tenant_id=:t"), {"t": c.tenant_id}).scalar_one()
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(_book(900_000.00)))
    wb.active["C2"] = "MAXX HD"
    buf = BytesIO()
    wb.save(buf)
    ingest_price_list(c, tenant_id=c.tenant_id, oem_id=oem_id, effective_from=date(2026, 11, 1),
                      parsed=parse_price_list(buf.getvalue()), actor_id="admin")
    _set(c, pricing_effective_on=date(2026, 11, 5))
    sync_deal_reconciliation(c, tenant_id=c.tenant_id, journey_id=c.journey_id, correlation_id="", raise_findings=False)
    assert _task(c, SKU_NOT_IN_PRICE_MASTER_TASK) is None


def test_the_publish_sweep_re_prices_the_deal(journey) -> None:
    """Re-review of PR 468, S9: a withdrawn SKU's standards were cleared;
    publishing a book that lists it again re-prices the deal, not only
    closes the task."""
    from audit_core.uc03_deal_reconciliation import sync_deal_reconciliation

    c = journey
    oem_id = c.execute(text("SELECT oem_id FROM auditcore.projects WHERE tenant_id=:t"), {"t": c.tenant_id}).scalar_one()
    c.execute(text("INSERT INTO auditcore.journey_products (tenant_id, journey_id, product_sku_id, selection_status) "
                   "VALUES (:t, :j, :s, 'CONFIRMED')"), {"t": c.tenant_id, "j": c.journey_id, "s": c.sku_id})
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(_book(900_000.00)))
    wb.active["C2"] = "MAXX HD"
    buf = BytesIO()
    wb.save(buf)
    ingest_price_list(c, tenant_id=c.tenant_id, oem_id=oem_id, effective_from=date(2026, 11, 1),
                      parsed=parse_price_list(buf.getvalue()), actor_id="admin")
    _set(c, pricing_effective_on=date(2026, 11, 5))
    sync_deal_reconciliation(c, tenant_id=c.tenant_id, journey_id=c.journey_id, correlation_id="")
    assert _task(c, SKU_NOT_IN_PRICE_MASTER_TASK)["task_status"] != "VERIFIED_COMPLETE"

    # A corrected November book that lists MAXX CITY again.
    ingest_price_list(c, tenant_id=c.tenant_id, oem_id=oem_id, effective_from=date(2026, 11, 1),
                      parsed=parse_price_list(_book(860_000.00)), actor_id="admin")
    assert _task(c, SKU_NOT_IN_PRICE_MASTER_TASK)["task_status"] == "VERIFIED_COMPLETE"
    standard = c.execute(text("SELECT standard_amount FROM auditcore.commercial_lines WHERE tenant_id=:t AND journey_id=:j "
                              "AND component_key='ex_showroom_price'"), {"t": c.tenant_id, "j": c.journey_id}).scalar_one()
    assert standard == Decimal("860000.00")
