"""Unit tests for the OEM native master parsers.

The xlsx parsers are exercised with synthesised workbooks that mirror the OEM's
real layout. The two PDF parsers were validated against the actual Mahindra
Sept'26 source documents; a regression test over those runs only when
``MAHINDRA_FIXTURES_DIR`` points at a directory holding them.
"""
from __future__ import annotations

import os
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from audit_core.oem_master_parsers import (
    MasterParseError,
    _parse_consumer_tables,
    old_vehicle_requirement,
    parse_corporate_policy,
    parse_discount_grid,
    parse_master,
    parse_model_qualifiers,
    parse_price_list,
    split_variant_cell,
    variant_excluded,
)
from audit_core.oem_price_masters import _build_preview

_PRICE_HEADER = [
    "Sl. No.", "Category", "Model", "Variant", "Trim", "Fuel", "Transmission",
    "Drive", "Seater", "Ex-Showroom Price", "TCS", "Insurance",
    "Ext. Warranty (4th Yr)", "Ext. Warranty (4th & 5th Yr)", "Accessories Kit",
    "RSA (1 Yr)", "FASTag", "Registration\n(Individual / w/o Hyp.)",
    "On-Road Price\n(Individual / w/o Hyp.)", "Registration\n(Corporate / with Hyp.)",
    "On-Road Price\n(Corporate / with Hyp.)", "Source Sheet",
]


def _price_row(sl, model, variant, ex, reg_ind, reg_corp, source="Sheet A", category="PV"):
    tcs, ins, ew4, ew45, acc, rsa, fastag = 100, 200, 300, 400, 500, 60, 40
    onroad_ind = ex + tcs + ins + ew4 + ew45 + acc + rsa + fastag + reg_ind
    onroad_corp = ex + tcs + ins + ew4 + ew45 + acc + rsa + fastag + reg_corp
    return [
        sl, category, model, variant, variant.split()[0], "PETROL", "MT", "2WD", "5",
        ex, tcs, ins, ew4, ew45, acc, rsa, fastag, reg_ind, onroad_ind, reg_corp,
        onroad_corp, source,
    ]


def _price_workbook(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Price List"
    ws.append(_PRICE_HEADER)
    for row in rows:
        ws.append(row)
    ws.append([None] * 22)
    ws.append(["Effective date", None, "w.e.f. 03.09.2026"] + [None] * 19)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_price_list_parses_components_and_reconciles() -> None:
    content = _price_workbook(
        [
            _price_row(1, "THAR ROXX", "MX1 PMT 2WD", 1_000_000, 90_000, 91_500),
            _price_row(2, "THAR ROXX", "MX5 DAT 4WD", 1_800_000, 160_000, 161_500),
        ]
    )
    result = parse_price_list(content)
    assert not result.errors
    assert len(result.price_rows) == 2
    assert result.effective_from_hint is not None
    first = result.price_rows[0]
    assert first.components["EX_SHOWROOM"] == Decimal("1000000.00")
    assert first.components["REGISTRATION_INDIVIDUAL"] == Decimal("90000.00")
    assert first.components["REGISTRATION_CORPORATE"] == Decimal("91500.00")
    # ex + tcs 100 + ins 200 + ew4 300 + ew45 400 + acc 500 + rsa 60 + fastag 40 + reg_ind
    assert first.onroad_individual == Decimal("1091600.00")


def test_price_list_flags_unreconciled_row() -> None:
    good = _price_row(1, "THAR ROXX", "MX1 PMT 2WD", 1_000_000, 90_000, 91_500)
    bad = _price_row(2, "THAR ROXX", "MX5 DAT 4WD", 1_800_000, 160_000, 161_500)
    bad[18] = 999  # tamper the individual on-road total
    result = parse_price_list(_price_workbook([good, bad]))
    assert len(result.price_rows) == 1
    assert any("on-road (individual)" in e for e in result.errors)


def test_price_list_splits_commercial_registration_basis() -> None:
    rows = [
        _price_row(1, "BOLERO", "B4 BS6.2", 800_000, 60_000, 61_500, source="BOLERO NEO (PVT)"),
        _price_row(2, "BOLERO", "B4 BS6.2", 800_000, 64_000, 65_500, source="BOLERO NEO-COM"),
    ]
    result = parse_price_list(_price_workbook(rows))
    assert not result.errors
    bases = sorted(r.registration_basis for r in result.price_rows)
    assert bases == ["COMMERCIAL", "PRIVATE"]


def test_price_list_rejects_foreign_workbook() -> None:
    wb = Workbook()
    wb.active.append(["something", "else"])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(MasterParseError):
        parse_price_list(buf.getvalue())


def test_price_list_tolerates_a_currency_glyph_stuck_to_a_header_label() -> None:
    # Confirmed against a real uploaded master (Mahindra Consolidated List,
    # 03-Sep-2026): the Model column's header cell literally reads "Model₹"
    # -- the Rupee symbol stuck directly onto the label, every other header
    # cell unaffected. The strict header-layout check rejected the entire
    # file outright, silently discarding real Trim data (this exact
    # scenario is what left SCORPIO CLASSIC's trim missing from the Modify
    # Model picker for about a week -- see
    # test_price_list_ingest_persists_trim_across_sibling_variants_of_one_
    # model in test_oem_price_masters.py, which proved ingestion itself was
    # fine once a file actually got past this gate). A header cell is a
    # label, never itself a currency value, so tolerating a trailing
    # currency glyph here can't mask a genuine layout mismatch.
    header = list(_PRICE_HEADER)
    header[2] = "Model₹"
    wb = Workbook()
    ws = wb.active
    ws.title = "Price List"
    ws.append(header)
    ws.append(_price_row(1, "THAR ROXX", "MX1 PMT 2WD", 1_000_000, 90_000, 91_500))
    buf = BytesIO()
    wb.save(buf)

    result = parse_price_list(buf.getvalue())
    assert not result.errors
    assert len(result.price_rows) == 1
    assert result.price_rows[0].model_name == "THAR ROXX"


def test_price_list_preview_exposes_the_structured_attributes() -> None:
    # Confirmed live: an admin reviewing the Publish preview for a fresh
    # upload had no way to see Trim/Fuel/Transmission/Drive/Seater at all --
    # not because parsing dropped them (it doesn't), but because the preview
    # sample dict never included them in the first place, only model/
    # variant/category/registrationBasis/prices. The preview table itself
    # renders whatever keys the sample rows carry (AdminOemMastersPage.tsx's
    # SampleTable derives its columns from Object.keys(sample[0])), so
    # backfilling these keys here is the whole fix, no frontend change
    # needed.
    result = parse_price_list(
        _price_workbook([_price_row(1, "THAR ROXX", "MX1 PMT 2WD", 1_000_000, 90_000, 91_500)])
    )
    preview = _build_preview(result, "PRICE_LIST")
    row = preview["sample"][0]
    assert row["trim"] == "MX1"
    assert row["fuel"] == "PETROL"
    assert row["transmission"] == "MT"
    assert row["drive"] == "2WD"
    assert row["seater"] == "5"


def _corporate_workbook() -> bytes:
    wb = Workbook()
    policy = wb.active
    policy.title = "Corporate Policy"
    policy.append(["CORPORATE PRIVILEGE POLICY From 3rd Sept' 26 - 30th Sept' 26"])
    policy.append(["Personal Range of Vehicles"])
    header = [None, "Corporate Category", "THAR ROXX", None, None, "SCORPIO-N", None, None]
    policy.append(header)
    policy.append([None, None, "M&M ", "Dealer", "Total", "M&M ", "Dealer", "Total"])
    policy.append([None, "Cat-B", 0, 0, 0, 1800, 1800, 3600])
    policy.append([None, "Cat-Z (Signature)", 24000, 16000, 40000, 24000, 16000, 40000])

    companies = wb.create_sheet("Companies List")
    companies.append([None, '"Z" SIGNATURE', None, None, None, None, '"B"'])
    companies.append(
        [None, "S. No. ", "Corporate Type", "Corporate Description", "Corporate Code",
         None, "S. No. ", "Corporate Type", "Corporate Description", "Corporate Code"]
    )
    companies.append(
        [None, 1, "CF - M&M Group Companies", "Mahindra & Mahindra Ltd.", 10642,
         None, 1, "C1, C3", "Acme Corp Ltd.", 20001]
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_corporate_policy_matrix_and_company_list() -> None:
    result = parse_corporate_policy(_corporate_workbook())
    assert not result.errors
    # Cat-B THAR ROXX is 0 -> skipped; Cat-B SCORPIO-N and both Cat-Z rows kept
    cats = {(b.privilege_category, b.brand_alias) for b in result.corporate_benefits}
    assert ("B", "SCORPIO-N") in cats
    assert ("Z", "THAR ROXX") in cats
    assert all(b.m_and_m + b.dealer == b.total for b in result.corporate_benefits)

    codes = {c.corporate_code: c for c in result.corporate_companies}
    assert codes["10642"].privilege_category == "Z"
    assert codes["20001"].privilege_category == "B"


def test_corporate_policy_requires_both_sheets() -> None:
    wb = Workbook()
    wb.active.title = "Corporate Policy"
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(MasterParseError):
        parse_corporate_policy(buf.getvalue())


# ── PDF regression (opt-in) ─────────────────────────────────────────────────────
_FIXTURES = os.environ.get("MAHINDRA_FIXTURES_DIR")


@pytest.mark.skipif(not _FIXTURES, reason="MAHINDRA_FIXTURES_DIR not set")
@pytest.mark.parametrize(
    ("glob", "kind", "min_rows"),
    [
        ("*Consumer_Scheme*.pdf", "CONSUMER_SCHEME", 40),
        ("*Exchange_scheme*.pdf", "EXCHANGE_SCHEME", 30),
    ],
)
def test_pdf_schemes_against_source_documents(glob, kind, min_rows) -> None:
    matches = list(Path(_FIXTURES).glob(glob))
    assert matches, f"no fixture matching {glob}"
    result = parse_master(kind, matches[0].read_bytes())
    assert not result.errors
    assert len(result.discount_rows) >= min_rows
    for row in result.discount_rows:
        assert row.total_customer_offer >= 0


# ── the dealer's per-model price sheets (decision 2026-09-30) ──────────────────
def _dealer_pv_workbook(rows, *, model="VEERO", fuel="Diesel", wef="Price list w.e.f. Dt.03.09.2026") -> bytes:
    """Mirrors the PV/CV layout: one sheet, a MODEL NAME row, a Model & Variant
    header with two registration / on-road pairs (individual, corporate) and
    the sheet's own column-letter formulas."""
    wb = Workbook()
    ws = wb.active
    ws.title = model
    ws.append(["ADITYA MOTORS"])
    ws.append([wef])
    ws.append(["MODEL NAME", None, model, None, "FUEL TYPE", fuel, None, None, None, None, "STATE", "ODISHA"])
    ws.append(["Model & Variant", None, "Ex-showroom Price", "Tax Collection at Source (TCS)", "Insurance",
               "Extended Warranty (4th & 5th Year)", "Accessories Kit", "RSA (1 year)", "Fastag", None,
               "On Road Price - Individual", None, "On Road Price - Corporate", None])
    ws.append([None] * 10 + ["Registration\n without Hypoth", "On Road Price\nwithout\nHypoth",
                             "Registration\n without Hypoth", "On Road Price\nwithout\nHypoth"])
    ws.append([None, None, "(A)", "(B)", "(C)", "(D)", "(E)", "(F)", "(G)", None,
               "(H)", "(I)=(A+B+C+D+E+F+G+H)", "(J)", "(K)=(A+B+C+D+E+F+G+J)"])
    for row in rows:
        ws.append(row)
    ws.append(["NOTE-"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pv_row(variant, ex, reg_ind, reg_corp, *, ins=34000, tamper=None):
    tcs, ew, acc, rsa, fastag = 0, 0, 5000, 1520, 500
    onroad_ind = ex + tcs + ins + ew + acc + rsa + fastag + reg_ind
    onroad_corp = ex + tcs + ins + ew + acc + rsa + fastag + reg_corp
    row = [variant, None, ex, tcs, ins, ew, acc, rsa, fastag, None, reg_ind, onroad_ind, reg_corp, onroad_corp]
    if tamper is not None:
        row[11] = tamper
    return row


def _dealer_ev_workbook(sheets: dict[str, list[list]]) -> bytes:
    """Mirrors the EV layout: one sheet per model, Essential Accessories in
    place of the kit, no RSA / warranty, registration without and with
    hypothecation."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        ws.append([None, "ADITYA MOTORS"])
        ws.append([None, "MODEL NAME", title, "ODISHA", "STATE", "CUTTACK"])
        ws.append([None, "MODEL & VARIANT", "Ex-showroom Price with Charger", "Tax Collection at Source (TCS)",
                   "Insurance", "Essential Accessories", "Fastag", "On Road Price - Individual"])
        ws.append([None] * 7 + ["Registration without Hypothecation", "On Road Price without Hypothecation",
                                "Registration with Hypothecation", "On Road Price with Hypothecation"])
        ws.append([None, None, "(A)", "(B)", "(C)", "(D)", "(E)", "(F)", "(G)=(A+B+C+D+E+F)", "(H)", "(I)=(A+B+C+D+E+H)"])
        for row in rows:
            ws.append(row)
        ws.append([None, "N:B-"])
        ws.append([None, 1, "The buyer is free to insure elsewhere"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ev_row(variant, ex):
    tcs, ins, acc, fastag, reg_without, reg_with = round(ex * 0.01), 80000, 25000, 500, 140, 1640
    base = ex + tcs + ins + acc + fastag
    return [None, variant, ex, tcs, ins, acc, fastag, reg_without, base + reg_without, reg_with, base + reg_with]


def test_dealer_pv_sheet_parses_reconciles_and_dates_from_the_sheet() -> None:
    content = _dealer_pv_workbook([
        _pv_row("1.6XXL HD V2", 859_501, 3_632, 3_632),
        _pv_row("1.5XXL SD V2", 839_500, 51_110, 51_110),
        _pv_row("1.5XXL SD V2", 839_500, 51_110, 51_110),  # the sheet repeats a row verbatim
    ])
    result = parse_price_list(content, filename="veero_LNT_Price.xlsx")
    assert not result.errors, result.errors
    assert result.meta["layout"] == "DEALER_PER_MODEL" and result.meta["sheets"] == 1
    assert result.meta["duplicateRowsCollapsed"] == 1
    assert result.effective_from_hint == date(2026, 9, 3) and result.meta["effectiveFromSource"] == "SHEET"
    first, second = result.price_rows
    assert first.model_name == "VEERO" and first.variant_name == "1.6XXL HD V2" and first.fuel == "DIESEL"
    assert first.category == "ICE" and first.registration_basis == "STANDARD" and first.source_sheet == "VEERO"
    assert first.components["EX_SHOWROOM"] == Decimal("859501.00")
    assert first.components["REGISTRATION_INDIVIDUAL"] == Decimal("3632.00")
    assert "EXT_WARRANTY_4TH_YR" not in first.components  # the sheet has no such column
    assert first.onroad_individual == Decimal("904153.00") == first.onroad_corporate
    assert second.components["REGISTRATION_INDIVIDUAL"] == Decimal("51110.00")


def test_dealer_pv_sheet_rejects_a_row_off_its_own_formula() -> None:
    content = _dealer_pv_workbook([
        _pv_row("1.6XXL HD V2", 859_501, 3_632, 3_632),
        _pv_row("1.6XXL SD V2", 839_501, 3_632, 3_632, tamper=999),
    ])
    result = parse_price_list(content)
    assert len(result.price_rows) == 1
    assert any("1.6XXL SD V2" in e and "on-road is 999" in e for e in result.errors)


def test_dealer_ev_workbook_reads_every_model_sheet_and_the_file_name_date() -> None:
    content = _dealer_ev_workbook({
        "BE6": [_ev_row("BE 6 One B59 R18 NCH", 1_890_000), _ev_row("BE 6 Two B59 R19 C7", 2_240_000)],
        "xuv3XO EV": [_ev_row("XUV3XO EV - AX5 FH", 1_389_000)],
    })
    result = parse_price_list(content, filename="EV_PRICE_-01th_Sept26.xlsx")
    assert not result.errors, result.errors
    assert result.meta["sheets"] == 2 and result.meta["models"] == ["BE6", "xuv3XO EV"]
    assert result.effective_from_hint == date(2026, 9, 1) and result.meta["effectiveFromSource"] == "FILENAME"
    by_variant = {r.variant_name: r for r in result.price_rows}
    be6 = by_variant["BE 6 One B59 R18 NCH"]
    assert be6.components["REGISTRATION_INDIVIDUAL"] == Decimal("140.00")
    assert be6.components["REGISTRATION_CORPORATE"] == Decimal("1640.00")
    assert be6.onroad_corporate - be6.onroad_individual == Decimal("1500.00")
    assert be6.category == "UNSPECIFIED"  # nothing on the sheet says electric
    assert "RSA_1YR" not in be6.components and "ACCESSORIES_KIT" in be6.components
    ev = by_variant["XUV3XO EV - AX5 FH"]
    assert ev.fuel == "ELECTRIC" and ev.category == "BEV"


def test_dealer_sheet_and_file_name_dates_disagree_is_a_warning() -> None:
    content = _dealer_pv_workbook([_pv_row("1.6XXL HD V2", 859_501, 3_632, 3_632)])
    result = parse_price_list(content, filename="02_Nos_Added_PRICE_-07th__Sept26.xlsx")
    assert result.effective_from_hint == date(2026, 9, 3)
    assert any("file name says 2026-09-07" in w for w in result.warnings)


def test_dealer_sheet_preview_shows_absent_components_as_blank() -> None:
    result = parse_price_list(_dealer_pv_workbook([_pv_row("1.6XXL HD V2", 859_501, 3_632, 3_632)]))
    sample = _build_preview(result, "PRICE_LIST")["sample"][0]
    assert sample["exShowroom"] == "859501.00" and sample["extWarranty4thYr"] is None


# ── the dealer's discount grid (fifth master, decision 2026-09-30) ─────────────
def _grid_workbook(rows, parameters=(), *, title="Effective 1st Sep 2026") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Grid for Aug"
    ws.append([None, title])
    ws.append([None, "Model", "Booking Protection", "Agreed Buffer", "Insurance OD %", "Out of Territory"])
    for row in rows:
        ws.append([None, *row])
    ws.append([None, "Parameter", "Notes"])
    for parameter, note in parameters:
        ws.append([None, parameter, note])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_discount_grid_reads_models_values_and_parameters() -> None:
    from audit_core.oem_master_parsers import parse_discount_grid

    result = parse_discount_grid(_grid_workbook(
        [
            ["3XO", "30 days", 7000, 0.6, "Additional 3K"],
            ["SCORPIO N Old", "Out of scope", "Out of scope", "Out of scope", "Out of scope"],
            ["7XO", "60 days", "Nil", 0.5, "Nil"],
            ["PICKUP, MAXX, MAXX HD", "30 days", 7000, 0.5, "Additional 3K"],
        ],
        [("Penalty Amount", "₹30,000 per case for any policy breach"), ("Price List", "All dealerships follow it")],
    ), filename="Odisha_PV__CV_Discount_Grid-1_Aug_2026.xlsx")
    assert not result.errors, result.errors
    assert result.effective_from_hint == date(2026, 9, 1) and result.meta["effectiveFromSource"] == "SHEET"
    rows = {r.model_alias: r for r in result.grid_rows}
    xo = rows["3XO"]
    assert (xo.booking_protection_days, xo.agreed_buffer_amount, xo.insurance_od_percent, xo.out_of_territory_amount) == (
        30, Decimal("7000.00"), Decimal("60.00"), Decimal("3000.00"))
    old = rows["SCORPIO N Old"]
    assert old.in_scope is False and old.booking_protection_days is None and old.agreed_buffer_amount is None
    seven = rows["7XO"]
    assert seven.agreed_buffer_amount == Decimal("0.00") and seven.out_of_territory_amount == Decimal("0.00")
    assert rows["PICKUP, MAXX, MAXX HD"].model_aliases == ["PICKUP", "MAXX", "MAXX HD"]
    assert result.grid_parameters[0] == {"parameter": "Penalty Amount", "note": "₹30,000 per case for any policy breach"}
    assert _build_preview(result, "DISCOUNT_GRID")["sample"][0]["insuranceOdPercent"] == "60.00"


def test_discount_grid_rejects_an_unreadable_cell_and_a_foreign_workbook() -> None:
    from audit_core.oem_master_parsers import parse_discount_grid

    result = parse_discount_grid(_grid_workbook([["3XO", "thirty", 7000, 0.6, "Additional 3K"]]))
    assert result.grid_rows == [] and any("unreadable days" in e for e in result.errors)
    wb = Workbook()
    wb.active.append(["something", "else"])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(MasterParseError):
        parse_discount_grid(buf.getvalue())


# ── scheme wording (2026-10-20 product owner rules) ─────────────────────────────
@pytest.mark.parametrize(
    ("cell", "names"),
    [
        ("MX2 PRO PM AT & MT", ["MX2 PRO PM AT", "MX2 PRO PM MT"]),
        ("MX3 PRO PM AT/MT", ["MX3 PRO PM AT", "MX3 PRO PM MT"]),
        ("MX3 PM MT/AT", ["MX3 PM MT", "MX3 PM AT"]),
        ("MX3 DS AS/MT\nMX3 PRO DS MT", ["MX3 DS AS", "MX3 DS MT", "MX3 PRO DS MT"]),
        ("AX7 DS AS DT/MT DT\nAX7L DS MT DT", ["AX7 DS AS DT", "AX7 DS MT DT", "AX7L DS MT DT"]),
        ("LXT D AT/MT 4WD 4S HT BS6.2", ["LXT D AT 4WD 4S HT BS6.2", "LXT D MT 4WD 4S HT BS6.2"]),
        ("EL PRO FC - Battery 34.5 Kwh -\nCharger 7.2 Kw - E\nEL PRO FC - DT - Battery 34.5\nKwh - Charger 7.2 Kw - E",
         ["EL PRO FC - Battery 34.5 Kwh - Charger 7.2 Kw - E", "EL PRO FC - DT - Battery 34.5 Kwh - Charger 7.2 Kw - E"]),
        ("Z8 L D AT 2WD 6 & 7STR BS6.2 -\nN - ADAS",
         ["Z8 L D AT 2WD 6 STR BS6.2 - N - ADAS", "Z8 L D AT 2WD 7 STR BS6.2 - N - ADAS"]),
        ("V4/V6", ["V4", "V6"]),
        ("Z4 G / Z8 G / Z8L G 2WD", ["Z4 G", "Z8 G", "Z8L G 2WD"]),
        ("Supro Maxi Profit Truck VX CBC\nBS6.2 - Tipper / Hopper",
         ["Supro Maxi Profit Truck VX CBC BS6.2 - Tipper", "Supro Maxi Profit Truck VX CBC BS6.2 - Hopper"]),
        ("MX5 DMT 4WD & Mocha Interiors", ["MX5 DMT 4WD & Mocha Interiors"]),
        ("M2/M4+/M6+(Discontinued Model)", ["M2", "M4+", "M6+"]),
    ],
)
def test_variant_cells_are_joined_and_expanded(cell, names) -> None:
    assert split_variant_cell(cell) == names


def test_model_cell_qualifiers_are_parsed_not_dropped() -> None:
    base, q = parse_model_qualifiers("XUV 3XO (Excluding EDGE & REVX M(O) PM AT )")
    assert base == "XUV 3XO" and q == {"excludedVariantTexts": ["EDGE", "REVX M(O) PM AT"]}
    base, q = parse_model_qualifiers("XUV 3XO (Excluding EDGE & REVX M(O) PM AT) against Old Maruti Vehicle**")
    assert base == "XUV 3XO" and q["requiredOldVehicle"] == {"brand": "MARUTI"}
    assert q["excludedVariantTexts"] == ["EDGE", "REVX M(O) PM AT"]
    assert parse_model_qualifiers("XUV3XO for All Variants") == ("XUV3XO", {"allVariants": True})
    assert parse_model_qualifiers("Bolero Neo Plus (All Variants)") == ("Bolero Neo Plus", {"allVariants": True})
    assert parse_model_qualifiers("THAR ROXX (Not on Star Edition)") == (
        "THAR ROXX", {"excludedVariantTexts": ["Star Edition"]})
    assert parse_model_qualifiers("3XO REVX (Not on REVX Edge)") == ("3XO REVX", {"excludedVariantTexts": ["REVX Edge"]})
    assert parse_model_qualifiers("Scorpio-N including Refresh") == ("Scorpio-N", {"includes": ["Refresh"]})
    assert old_vehicle_requirement("Tata Ace Family") == {"brand": "TATA", "model": "ACE"}
    assert old_vehicle_requirement("Any Personal + Commercial Vehicle") is None
    assert old_vehicle_requirement("Not Applicable") is None
    assert variant_excluded("REVX EDGE(O) PM AT", ["EDGE"]) and variant_excluded("REVX M(O) PM AT", ["REVX M(O) PM AT"])
    assert not variant_excluded("REVX M(O) PM MT", ["EDGE", "REVX M(O) PM AT"])
    assert variant_excluded("AX7 DAT 2WD - Star Edition", ["Star Edition"])


def _bulletin_row(product, variants, factor, total, cash, other="₹ 0", description=""):
    contribution = f"{int(total.replace('₹', '').replace(',', '').strip()):,}"
    return [product, variants, "1", factor, contribution, "-", total, cash, other, description]


_HEADER_ROW = ["Product", "Variants", "M&M Contribution", "BNDP Factor", "M&M", "Dealer", "Total Consumer Offer",
               "Cash Discount", "Other Schemes", "Description"]


def test_a_product_cell_is_never_carried_over_a_page_break() -> None:
    """The September bulletin's XUV400 EV block opens page 2 with an empty
    product cell: carried, the Rs 2,50,000 EV offer read as XUV 3XO."""
    pages = [
        [[_HEADER_ROW,
          _bulletin_row("XUV 3XO", "AX5 PM MT", "1.18", "₹ 30,000", "₹ 30,000"),
          _bulletin_row("XUV 3XO", "AX7 DS AS DT/MT DT\nAX7L DS MT DT", "1.18", "₹ 20,000", "₹ 20,000")]],
        [[_bulletin_row(None, "EL PRO FC - Battery 34.5 Kwh -\nCharger 7.2 Kw - E", "1.05", "₹ 2,50,000", "₹ 2,50,000"),
          _bulletin_row("XUV400", "EC PRO - Battery 34.5 Kwh -\nCharger 3.2 Kw - E", "1.05", "₹ 50,000", "₹ 50,000")]],
    ]
    result = _parse_consumer_tables(pages, "w.e.f. 3rd Sep;26")
    assert not result.errors, result.errors
    bev = result.discount_rows[2]
    assert bev.model_alias == "XUV400" and bev.benefits == [("CASH_DISCOUNT", Decimal("250000.00"))]
    assert bev.variant_texts == ["EL PRO FC - Battery 34.5 Kwh - Charger 7.2 Kw - E"]
    assert bev.config["vehicleCategory"] == "EV"
    assert [r.model_alias for r in result.discount_rows] == ["XUV 3XO", "XUV 3XO", "XUV400", "XUV400"]


def test_a_row_split_by_the_page_break_is_one_row_and_an_unplaceable_one_is_refused() -> None:
    veero = [[_HEADER_ROW, _bulletin_row("VEERO", "V2\nV2 CNG", "1.18", "₹ 50,000", "₹ 45,000", "₹ 5,000")]]
    continued = [[_bulletin_row(None, "V4 CNG", "1.18", "₹ 50,000", "₹ 45,000", "₹ 5,000"),
                  _bulletin_row("MAXX CITY", "MAXX CITY 1.4 LX BS6.2", "1.18", "₹ 20,000", "₹ 15,000", "₹ 5,000")]]
    result = _parse_consumer_tables([veero, continued], "")
    assert not result.errors
    assert [(r.model_alias, r.variant_texts) for r in result.discount_rows] == [
        ("VEERO", ["V2", "V2 CNG", "V4 CNG"]), ("MAXX CITY", ["MAXX CITY 1.4 LX BS6.2"])]

    orphan = [[_bulletin_row(None, "V4 CNG", "1.18", "₹ 40,000", "₹ 40,000"),
               _bulletin_row("MAXX CITY", "MAXX CITY 1.4 LX BS6.2", "1.18", "₹ 20,000", "₹ 15,000", "₹ 5,000")]]
    result = _parse_consumer_tables([veero, orphan], "")
    assert any("cannot be placed" in e for e in result.errors)


def test_the_ev_bndp_factor_requires_an_ev_model() -> None:
    pages = [[[_HEADER_ROW, _bulletin_row("XUV 3XO", "EC PRO", "1.05", "₹ 2,50,000", "₹ 2,50,000"),
               _bulletin_row("XUV400", "EC PRO", "1.18", "₹ 50,000", "₹ 50,000")]]]
    result = _parse_consumer_tables(pages, "")
    assert result.discount_rows == []
    assert any("1.05 is the EV rate" in e for e in result.errors)
    assert any("not the EV rate 1.05" in e for e in result.errors)


def test_the_bev_carry_forward_email_is_refused_with_what_it_says() -> None:
    text = (
        "Consumer Schemes - BEVs - Sep'26\nDear All,\n"
        "We are continuing the same Consumer Schemes of August’26 in the month of September’26.\n"
        "There are no changes in schemes of –\nXEV 9e including Extended Warranty\nBE6 (old series)\n"
        "FE (old series)\nPls note there are no consumer schemes in –\nXEV9S\nBE6 Sporteq\nRegards\n"
    )
    with pytest.raises(MasterParseError) as caught:
        _parse_consumer_tables([[]], text)
    message = str(caught.value)
    assert "e-mail" in message and "Nothing was loaded" in message
    assert "XEV 9e including Extended Warranty, BE6 (old series), FE (old series)" in message
    assert "XEV9S, BE6 Sporteq" in message


def test_the_october_grid_reads_three_tiers_and_the_od_caps_from_its_notes() -> None:
    wb = Workbook()
    ws = wb.active
    ws.append([None, "Grid for ' October 2026"])
    ws.append([None, "Model", "Booking Protection", "MY - 2026\nWithout Inhouse Insurance or Extended Warranty",
               "MY - 2026\nWith Insurance  ", "With Extended Warranty + Insurance", "Out of Territory"])
    ws.append([None, "XUV 3XO", "30 days", 5000, 7000, 10000, "Additional 5K"])
    ws.append([None, "SCORPIO-N", "60 days", "Nil", "Nil", 5000, "Additional 5K"])
    ws.append([None, "PICKUP, MAXX City, MAXX HD", "30 days", 5000, 7000, 7000, "Additional 5K"])
    ws.append([])
    ws.append([None, "Notes"])
    ws.append([None, "Parameter", "Notes"])
    ws.append([None, "OD Discount  Insurance", (
        "Private Car: Maximum OD discount – 65%\nGoods Carrying Vehicle (GCV): Maximum OD discount – 60%\n"
        "Passenger Carrying Vehicle (PCV): Maximum OD discount – 50%\n3+3 Insurance: Maximum discount – 85%")])
    buf = BytesIO()
    wb.save(buf)
    result = parse_discount_grid(buf.getvalue(), filename="Odisha PV & CV Discount Grid-1Oct  2026 - Copy.xlsx")
    assert not result.errors, result.errors
    assert result.effective_from_hint == date(2026, 10, 1) and result.meta["effectiveFromSource"] == "FILENAME"
    rows = {r.model_alias: r for r in result.grid_rows}
    assert set(rows) == {"XUV 3XO", "SCORPIO-N", "PICKUP, MAXX City, MAXX HD"}  # "Notes" is not a model
    xo = rows["XUV 3XO"]
    assert xo.tiers == {"BASE": Decimal("5000.00"), "WITH_INSURANCE": Decimal("7000.00"),
                        "WITH_EW_INSURANCE": Decimal("10000.00")}
    assert xo.agreed_buffer_amount == Decimal("5000.00") and xo.out_of_territory_amount == Decimal("5000.00")
    assert xo.insurance_od_percent == Decimal("65.00")
    assert rows["PICKUP, MAXX City, MAXX HD"].insurance_od_percent == Decimal("60.00")  # goods carrier cap
    assert rows["SCORPIO-N"].tiers["BASE"] == Decimal("0.00")
    assert result.meta["odCapsPercent"] == {"PRIVATE_CAR": "65.00", "GCV": "60.00", "PCV": "50.00",
                                            "THREE_PLUS_THREE": "85.00"}


def _fixture(pattern: str) -> Path:
    matches = sorted(Path(_FIXTURES).rglob(pattern))
    if not matches:
        pytest.skip(f"no fixture matching {pattern}")
    return matches[0]


@pytest.mark.skipif(not _FIXTURES, reason="MAHINDRA_FIXTURES_DIR not set")
def test_real_september_bulletin_reads_the_bev_row_as_xuv400() -> None:
    result = parse_master("CONSUMER_SCHEME", _fixture("Consumer Scheme*Sep*26*.pdf").read_bytes())
    assert not result.errors
    bev = [r for r in result.discount_rows if r.total_customer_offer == Decimal("250000.00")]
    assert len(bev) == 1 and bev[0].model_alias == "XUV400" and bev[0].config["vehicleCategory"] == "EV"
    assert all(r.model_alias != "XUV 3XO" or r.config["vehicleCategory"] == "ICE" for r in result.discount_rows)


@pytest.mark.skipif(not _FIXTURES, reason="MAHINDRA_FIXTURES_DIR not set")
def test_real_october_exchange_xuv3xo_rows_parse_with_their_conditions() -> None:
    result = parse_master("EXCHANGE_SCHEME", _fixture("Exchange scheme Oct*26*.pdf").read_bytes())
    assert not result.errors
    xuv = [(r.model_alias, r.benefits) for r in result.discount_rows if parse_model_qualifiers(r.model_alias)[0] in
           {"XUV 3XO", "XUV3XO"}]
    amounts = sorted(str(b[0][1]) for _, b in xuv)
    assert amounts == ["20000.00", "40000.00", "40000.00", "60000.00"]
    maruti = next(a for a, b in xuv if b[0][1] == Decimal("60000.00"))
    assert parse_model_qualifiers(maruti)[1]["requiredOldVehicle"] == {"brand": "MARUTI"}


@pytest.mark.skipif(not _FIXTURES, reason="MAHINDRA_FIXTURES_DIR not set")
def test_real_october_grid_and_bev_email() -> None:
    grid_file = _fixture("*Discount Grid-1Oct*.xlsx")
    grid = parse_master("DISCOUNT_GRID", grid_file.read_bytes(), filename=grid_file.name)
    assert not grid.errors and grid.meta["layout"] == "DEALER_DISCOUNT_GRID_TIERED"
    assert {r.model_alias: r.tiers.get("WITH_EW_INSURANCE") for r in grid.grid_rows}["XUV 3XO"] == Decimal("10000.00")
    with pytest.raises(MasterParseError, match="e-mail"):
        parse_master("CONSUMER_SCHEME", _fixture("Consumer Schemes - BEVs*.pdf").read_bytes())
