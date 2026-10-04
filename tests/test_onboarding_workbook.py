"""Onboarding workbook: template, codes, parsing and the import plan (pure)."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook, load_workbook

from audit_core.onboarding_workbook import (
    DEALERSHIP_COLUMNS,
    PROJECT_COLUMNS,
    ExistingState,
    build_workbook,
    derive_dealer_code,
    export_rows,
    parse_workbook,
    plan_import,
)

OEMS = [{"oem_id": "m", "oem_code": "MAHINDRA", "oem_name": "Mahindra"},
        {"oem_id": "h", "oem_code": "HYUNDAI", "oem_name": "Hyundai"}]
SEGMENTS = [{"segment_id": "s1", "segment_code": "PV", "segment_name": "Passenger"}]


def _state(**kwargs):
    return ExistingState(oems=OEMS, segments=SEGMENTS, projects=kwargs.get("projects", []),
                         dealers=kwargs.get("dealers", {}), outlets=kwargs.get("outlets", {}))


def _book(projects, dealerships, *, project_header=None, dealer_header=None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    sheet = wb.create_sheet("Projects")
    sheet.append(project_header or ["Code", "Name", "OEM", "Active", "Start Date", "End Date"])
    for row in projects:
        sheet.append(row)
    sheet = wb.create_sheet("Dealerships")
    sheet.append(dealer_header or ["Code", "OEM", "Project Code", "Dealership", "Outlet Name", "Location",
                                   "State", "City", "Active", "PC Presence", "Monthly Car Sales Volume"])
    for row in dealerships:
        sheet.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


PROJECTS = [["JBR-01", "Mahindra Orissa", "Mahindra", "Yes", date(2026, 9, 4), None],
            ["JBR-02", "Hyundai Orissa", "Hyundai", "Yes", date(2026, 9, 14), None]]
OUTLETS = [
    ["AM-MAH-CUBE", "Mahindra", "JBR-01", "Aditya Motors", "Aditya Motors - Cube", "Plot 9", "Odisha", "Bhubaneswar", "Yes", "Onsite", 40],
    ["AM-MAH-PURI", "Mahindra", "JBR-01", "Aditya Motors", "Aditya Motors - Puri", "NH", "Odisha", "Puri", "Yes", "Satellite", 12],
    ["UH-HYU-PAHAL", "Hyundai", "JBR-02", "Utkal Hyundai", "Utkal Hyundai - Pahal", "Pahal", "Odisha", "Bhubaneswar", "Yes", None, None],
]


def test_dealer_code_is_initials_plus_oem():
    assert derive_dealer_code("Aditya Motors", "MAHINDRA") == "AM-MAH"
    assert derive_dealer_code("Utkal Hyundai", "HYUNDAI") == "UH-HYU"
    assert derive_dealer_code("Shivnath Motors Pvt. Ltd.", "MAHINDRA") == "SM-MAH"
    assert derive_dealer_code("Premier", "HYUNDAI") == "PR-HYU"


def test_a_new_workbook_plans_projects_dealers_and_outlets():
    plan = plan_import(parse_workbook(_book(PROJECTS, OUTLETS)), _state())
    assert plan["summary"] == {"projects": {"CREATE": 2}, "dealers": {"CREATE": 2}, "outlets": {"CREATE": 3},
                               "errors": 0}
    assert sorted(d["dealerCode"] for d in plan["dealers"]) == ["AM-MAH", "UH-HYU"]
    cube, puri, pahal = plan["outlets"]
    assert (cube["outletClassification"], cube["monthlyVehicleVolume"]) == ("ONSITE", 40)
    assert puri["outletClassification"] == "SATELLITE"
    assert pahal["outletClassification"] == "ONSITE"  # blank PC Presence = Onsite
    assert all(p["activate"] for p in plan["projects"])  # Active=Yes on a new project


def test_row_errors_are_reported_per_row():
    projects = [["JBR-01", "Mahindra Orissa", "Mahindra", "Yes", date(2026, 9, 4), date(2026, 9, 1)],
                ["JBR-01", "", "Tesla", "Maybe", None, None]]
    outlets = [
        ["AM-MAH-CUBE", "Hyundai", "JBR-01", "Aditya Motors", "Cube", "x", "Odisha", "Bbsr", "Yes", "Remote", -3],
        ["AM-MAH-CUBE", "Mahindra", "JBR-09", "Aditya Motors", "Cube 2", "x", "Odisha", "Bbsr", "Yes", "Onsite", 5],
    ]
    plan = plan_import(parse_workbook(_book(projects, outlets)), _state())
    first, second = plan["projects"]
    assert "End Date cannot be earlier than Start Date." in first["messages"]
    assert any("more than once" in m for m in second["messages"])
    assert any("not a known OEM" in m for m in second["messages"])
    assert any("Active must be Yes or No" in m for m in second["messages"])
    assert "Name is required." in second["messages"]
    bad_oem, unknown_project = plan["outlets"]
    assert any("does not match the project's OEM" in m for m in bad_oem["messages"])
    assert "PC Presence must be Onsite or Satellite." in bad_oem["messages"]
    assert any("Monthly Car Sales Volume" in m for m in bad_oem["messages"])
    assert any("JBR-09 is not in the Projects sheet" in m for m in unknown_project["messages"])
    assert plan["summary"]["errors"] == 4


def test_two_dealerships_with_the_same_derived_code_need_their_own_code():
    outlets = [
        ["AM-MAH-1", "Mahindra", "JBR-01", "Aditya Motors", "One", "x", "Odisha", "B", "Yes", "Onsite", 1],
        ["AM-MAH-2", "Mahindra", "JBR-01", "Anand Motors", "Two", "x", "Odisha", "B", "Yes", "Onsite", 1],
    ]
    plan = plan_import(parse_workbook(_book(PROJECTS[:1], outlets)), _state())
    assert any("give one of them its own Dealer Code" in m for m in plan["outlets"][1]["messages"])


def test_an_exported_workbook_round_trips_as_unchanged_then_updates():
    tenant = "tenant-1"
    state = _state(
        projects=[{"tenant_id": tenant, "business_code": "JBR-01", "project_code": "tenant-x",
                   "project_name": "Mahindra Orissa", "oem_code": "MAHINDRA", "effective_start_date": date(2026, 9, 4),
                   "effective_end_date": None, "timezone_name": "Asia/Kolkata", "project_status": "ACTIVE",
                   "segment_codes": ["PV"]}],
        dealers={tenant: [{"dealer_id": "d1", "dealer_code": "AM-MAH", "dealer_name": "Aditya Motors", "status": "ACTIVE"}]},
        outlets={tenant: [{"outlet_id": "o1", "dealer_id": "d1", "outlet_code": "AM-MAH-CUBE", "outlet_name": "Cube",
                           "outlet_classification": "ONSITE", "address_text": "Plot 9", "city": "Bhubaneswar",
                           "state_region": "Odisha", "postal_code": None, "latitude": None, "longitude": None,
                           "monthly_vehicle_volume": 40, "status": "ACTIVE"}]},
    )
    projects, dealerships = export_rows(state)
    content = build_workbook(oems=OEMS, projects=projects, dealerships=dealerships)
    exported = load_workbook(io.BytesIO(content))
    assert exported.sheetnames == ["Instructions", "Projects", "Dealerships"]
    assert [c.value for c in exported["Projects"][1]] == list(PROJECT_COLUMNS)
    assert [c.value for c in exported["Dealerships"][1]] == list(DEALERSHIP_COLUMNS)

    plan = plan_import(parse_workbook(content), state)
    assert plan["summary"] == {"projects": {"UNCHANGED": 1}, "dealers": {"UNCHANGED": 1},
                               "outlets": {"UNCHANGED": 1}, "errors": 0}

    sheet = exported["Dealerships"]
    sheet.cell(row=2, column=DEALERSHIP_COLUMNS.index("PC Presence") + 1, value="Satellite")
    sheet.cell(row=2, column=DEALERSHIP_COLUMNS.index("Monthly Car Sales Volume") + 1, value=55)
    exported["Projects"].cell(row=2, column=PROJECT_COLUMNS.index("Name") + 1, value="Mahindra Odisha")
    buffer = io.BytesIO()
    exported.save(buffer)
    plan = plan_import(parse_workbook(buffer.getvalue()), state)
    assert plan["projects"][0]["action"] == "UPDATE" and plan["projects"][0]["changes"] == {
        "name": ["Mahindra Orissa", "Mahindra Odisha"]}
    [outlet] = plan["outlets"]
    assert outlet["action"] == "UPDATE"
    assert outlet["changes"] == {"outletClassification": ["ONSITE", "SATELLITE"], "monthlyVehicleVolume": [40, 55]}


def test_existing_records_without_codes_are_matched_by_name_and_get_codes():
    tenant = "tenant-1"
    state = _state(
        projects=[{"tenant_id": tenant, "business_code": "JBR-01", "project_code": "t", "project_name": "Mahindra Orissa",
                   "oem_code": "MAHINDRA", "effective_start_date": date(2026, 9, 4), "effective_end_date": None,
                   "timezone_name": "Asia/Kolkata", "project_status": "CONFIGURING"}],
        dealers={tenant: [{"dealer_id": "d1", "dealer_code": "a" * 32, "dealer_name": "aditya motors", "status": "ACTIVE"}]},
        outlets={tenant: [{"outlet_id": "o1", "dealer_id": "d1", "outlet_code": "b" * 32, "outlet_name": "Aditya Motors - Cube",
                           "outlet_classification": "ONSITE", "address_text": "Plot 9", "city": "Bhubaneswar",
                           "state_region": "Odisha", "postal_code": None, "latitude": None, "longitude": None,
                           "monthly_vehicle_volume": 40, "status": "ACTIVE"}]},
    )
    plan = plan_import(parse_workbook(_book(PROJECTS[:1], OUTLETS[:1])), state)
    [dealer] = plan["dealers"]
    assert dealer["dealerId"] == "d1" and dealer["changes"]["dealerCode"] == ["a" * 32, "AM-MAH"]
    [outlet] = plan["outlets"]
    assert outlet["outletId"] == "o1" and outlet["changes"]["outletCode"] == ["b" * 32, "AM-MAH-CUBE"]


def test_oem_and_start_date_cannot_change_on_an_existing_project():
    state = _state(projects=[{"tenant_id": "t1", "business_code": "JBR-01", "project_code": "t", "project_name": "M",
                              "oem_code": "MAHINDRA", "effective_start_date": date(2026, 9, 1), "effective_end_date": None,
                              "timezone_name": "Asia/Kolkata", "project_status": "ACTIVE"}])
    plan = plan_import(parse_workbook(_book([["JBR-01", "M", "Hyundai", "Yes", date(2026, 9, 4), None]], [])), state)
    messages = plan["projects"][0]["messages"]
    assert "OEM cannot be changed by import once the project exists." in messages
    assert "Start Date cannot be changed by import once the project exists." in messages


def test_missing_required_columns_and_unreadable_files_are_reported():
    parsed = parse_workbook(_book([], [], project_header=["Code", "Name"]))
    assert any("missing column(s): OEM, Start Date" in e for e in parsed.errors)
    try:
        parse_workbook(b"not an excel file")
    except ValueError as exc:
        assert "readable Excel workbook" in str(exc)
    else:
        raise AssertionError("expected ValueError")


def test_a_preview_only_needs_the_projects_the_workbook_mentions():
    from audit_core.onboarding_imports import workbook_tenants

    projects = [{"tenant_id": "t1", "business_code": "JBR-01"}, {"tenant_id": "t2", "business_code": "JBR-02"},
                {"tenant_id": "t3", "business_code": None}]
    parsed = parse_workbook(_book([["JBR-01", "M", "Mahindra", "Yes", date(2026, 9, 4), None]],
                                  [["UH-HYU-1", "Hyundai", "JBR-02", "Utkal Hyundai", "One", "x", "Odisha", "B", "Yes", "Onsite", 1]]))
    assert workbook_tenants(parsed, projects) == {"t1", "t2"}


# ------------------------------------------------ the OEM geofencing master (Hyundai_Master)

MASTER_HEADER = ["Dealer Group", "Code", "OEM", "Outlet Name", "Complete Address (Recommended)", "State",
                 "City (Source)", "Google Maps URL", "Google Place ID", "Latitude", "Longitude",
                 "Dealership Code", "Outlet Code", "PC Presence"]
MASTER_ROWS = [
    ["Aditya", "E7205", "Hyundai", "Aditya Hyundai, Tamando", "Plot No. 11, NH-5, Tamando, Odisha 752054", "Odisha",
     "Bhubaneshwar", "https://maps.example/1", "ChIJ1", 20.230755, 85.737483, "AH-HYU", "AH-HYU-TAMANDO", "Yes"],
    ["Premier", "E7215", "Hyundai", "Premier Hyundai, Baripada", "Baripada, Odisha", "Odisha", "Baripada",
     "https://maps.example/2", "ChIJ2", 21.899926, 86.758471, "PH-HYU", "PH-HYU-BARIPADA", "Yes"],
]


def _master(rows=None, header=None, title="Hyundai_Master") -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.title = title
    sheet.append(header or MASTER_HEADER)
    for row in (MASTER_ROWS if rows is None else rows):
        sheet.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _hyundai_project(tenant="t-hyu", code="JBR-02"):
    return {"tenant_id": tenant, "business_code": code, "project_code": "t", "project_name": "Hyundai Orissa",
            "oem_code": "HYUNDAI", "effective_start_date": date(2026, 9, 14), "effective_end_date": None,
            "timezone_name": "Asia/Kolkata", "project_status": "ACTIVE"}


def test_the_oem_geofencing_master_is_read_as_it_is():
    parsed = parse_workbook(_master())
    assert not parsed.errors and len(parsed.dealerships.rows) == 2
    plan = plan_import(parsed, _state(projects=[_hyundai_project()]))
    assert plan["summary"]["errors"] == 0 and plan["summary"]["outlets"] == {"CREATE": 2}
    assert sorted((d["dealerCode"], d["dealerName"]) for d in plan["dealers"]) == [
        ("AH-HYU", "Aditya Hyundai"), ("PH-HYU", "Premier Hyundai")]
    tamando = plan["outlets"][0]
    assert tamando["outletCode"] == "AH-HYU-TAMANDO" and tamando["projectCode"] == "JBR-02"
    assert (tamando["latitude"], tamando["longitude"]) == ("20.2307550", "85.7374830")
    assert tamando["outletClassification"] == "ONSITE" and tamando["city"] == "Bhubaneshwar"
    assert tamando["addressText"].startswith("Plot No. 11")


def test_the_master_needs_its_oems_only_project_or_a_project_code():
    two = _state(projects=[_hyundai_project(), _hyundai_project("t-2", "JBR-03")])
    plan = plan_import(parse_workbook(_master()), two)
    assert plan["summary"]["errors"] == 2
    assert any("2 Hyundai projects" in m for m in plan["outlets"][0]["messages"])
    none = plan_import(parse_workbook(_master()), _state())
    assert any("no Hyundai project" in m for m in none["outlets"][0]["messages"])
    with_code = _master(header=MASTER_HEADER + ["Project Code"], rows=[r + ["JBR-03"] for r in MASTER_ROWS])
    plan = plan_import(parse_workbook(with_code), two)
    assert plan["summary"]["errors"] == 0 and plan["outlets"][0]["tenantId"] == "t-2"


def test_an_outlet_onboarded_under_the_oems_code_is_updated_not_duplicated():
    tenant = "t-hyu"
    state = _state(
        projects=[_hyundai_project(tenant)],
        dealers={tenant: [{"dealer_id": "d1", "dealer_code": "AH-HYU", "dealer_name": "Aditya Hyundai", "status": "ACTIVE"}]},
        outlets={tenant: [{"outlet_id": "o1", "dealer_id": "d1", "outlet_code": "E7205",
                           "outlet_name": "Aditya Hyundai, Tamando", "outlet_classification": "ONSITE",
                           "address_text": "Plot No. 11, NH-5, Tamando, Odisha 752054", "city": "Bhubaneshwar",
                           "state_region": "Odisha", "postal_code": None, "latitude": None, "longitude": None,
                           "monthly_vehicle_volume": None, "status": "ACTIVE"}]},
    )
    plan = plan_import(parse_workbook(_master()), state)
    tamando, baripada = plan["outlets"]
    assert tamando["action"] == "UPDATE" and tamando["outletId"] == "o1"
    assert tamando["changes"]["outletCode"] == ["E7205", "AH-HYU-TAMANDO"]
    assert tamando["changes"]["latitude"][1] == "20.2307550"
    assert baripada["action"] == "CREATE"
    assert plan["summary"]["dealers"] == {"UNCHANGED": 1, "CREATE": 1}


def test_the_master_reports_bad_coordinates_and_presence_per_row():
    rows = [MASTER_ROWS[0][:9] + [120.0, 85.7] + MASTER_ROWS[0][11:13] + ["Maybe"], MASTER_ROWS[1]]
    plan = plan_import(parse_workbook(_master(rows=rows)), _state(projects=[_hyundai_project()]))
    bad = plan["outlets"][0]["messages"]
    assert any("Latitude must be between" in m for m in bad) and any("PC Presence" in m for m in bad)
    assert plan["outlets"][1]["action"] == "CREATE"


def test_a_sheet_that_is_not_an_outlet_master_is_still_not_found():
    parsed = parse_workbook(_master(header=["Name", "Notes"], rows=[["a", "b"]]))
    assert not parsed.dealerships.rows and parsed.projects.errors
