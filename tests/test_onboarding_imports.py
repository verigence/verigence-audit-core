"""Excel onboarding API on Postgres: template, upload -> preview -> apply,
round-trip export, validation failures and SuperAdmin-only access.
Provisioning (Security + DI) and activation are stubbed at the import's
boundary; they have their own tests."""
from __future__ import annotations

import io
import os
from datetime import date
from uuid import uuid4

import pytest
from conftest import delete_tenant_data
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, text

from audit_core import onboarding_imports
from audit_core.dependencies import (
    HumanAdminRequest,
    get_engine,
    require_super_admin_request,
)
from audit_core.main import app
from audit_core.security_integration import SecurityAdminContext

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _admin(super_admin: bool = True) -> HumanAdminRequest:
    return HumanAdminRequest(
        user_id="superadmin-onboarding", bearer_token="token",
        admin_context=SecurityAdminContext(user_id="superadmin-onboarding", is_super_admin=super_admin, admin_scopes=()),
    )


@pytest.fixture
def onboarding(monkeypatch):
    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        pytest.skip("DATABASE_URL is required for onboarding integration tests")
    engine = create_engine(database_url)
    suffix = uuid4().hex[:8].upper()
    created: list[str] = []
    activations: list[str] = []

    def fake_create(entry, state, *, import_id, admin_request, engine):
        tenant_id = f"tenant-onb-{uuid4().hex}"
        oem_id = next(o["oem_id"] for o in state.oems if o["oem_code"] == entry["oemCode"])
        with engine.begin() as connection:
            connection.execute(
                text(
                    """
                    INSERT INTO auditcore.projects (tenant_id, project_code, project_name, oem_id,
                        effective_start_date, timezone_name, project_status)
                    VALUES (:t, :t, :name, :oem, :start, 'Asia/Kolkata', 'CONFIGURING')
                    """
                ),
                {"t": tenant_id, "name": entry["name"], "oem": oem_id, "start": entry["startDate"]},
            )
        created.append(tenant_id)
        return tenant_id

    def fake_activate(tenant_id, *, import_id, admin_request, engine):
        activations.append(tenant_id)
        return "Resolve blocking Project Readiness checks before activation: PROJECT_MASTERS_READY."

    monkeypatch.setattr(onboarding_imports, "_create_project", fake_create)
    monkeypatch.setattr(onboarding_imports, "_activate", fake_activate)
    app.dependency_overrides[get_engine] = lambda: engine
    app.dependency_overrides[require_super_admin_request] = lambda: _admin()
    try:
        yield {"engine": engine, "suffix": suffix, "created": created, "activations": activations}
    finally:
        app.dependency_overrides.clear()
        for tenant_id in created:
            delete_tenant_data(engine, tenant_id)
        with engine.begin() as connection:
            connection.execute(text("DELETE FROM auditcore.onboarding_imports WHERE created_by_actor_id='superadmin-onboarding'"))
        engine.dispose()


def _workbook(suffix: str, *, bad: bool = False) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    projects = wb.create_sheet("Projects")
    projects.append(["Code", "Name", "OEM", "Active", "Start Date", "End Date"])
    projects.append([f"T{suffix}-01", "Mahindra Orissa", "Mahindra", "Yes", date(2026, 9, 4), None])
    projects.append([f"T{suffix}-02", "Hyundai Orissa", "Hyundai" if not bad else "Tesla", "Yes", date(2026, 9, 14), None])
    dealerships = wb.create_sheet("Dealerships")
    dealerships.append(["Code", "OEM", "Project Code", "Dealership", "Outlet Name", "Location", "State", "City",
                        "Active", "PC Presence", "Monthly Car Sales Volume"])
    dealerships.append(["AM-MAH-CUBE", "Mahindra", f"T{suffix}-01", "Aditya Motors", "Aditya Motors - Cube",
                        "Plot 9", "Odisha", "Bhubaneswar", "Yes", "Onsite", 40])
    dealerships.append(["AM-MAH-PURI", "Mahindra", f"T{suffix}-01", "Aditya Motors", "Aditya Motors - Puri",
                        "NH 316", "Odisha", "Puri", "No", "Satellite", 12])
    dealerships.append(["UH-HYU-PAHAL", "Hyundai", f"T{suffix}-02", "Utkal Hyundai", "Utkal Hyundai - Pahal",
                        "Pahal", "Odisha", "Bhubaneswar", "Yes", "Satellite", 25])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload(client, content: bytes):
    return client.post("/v1/onboarding/imports", files={"file": ("onboarding.xlsx", content, XLSX)})


def test_template_download(onboarding):
    response = TestClient(app).get("/v1/onboarding/workbook")
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Instructions", "Projects", "Dealerships"]
    assert book["Dealerships"].max_row == 1  # headers only


def test_upload_preview_apply_and_round_trip(onboarding):
    client = TestClient(app, raise_server_exceptions=False)
    suffix = onboarding["suffix"]
    preview = _upload(client, _workbook(suffix))
    assert preview.status_code == 201, preview.text
    body = preview.json()
    assert body["status"] == "PREVIEW_READY"
    assert body["plan"]["summary"] == {"projects": {"CREATE": 2}, "dealers": {"CREATE": 2},
                                       "outlets": {"CREATE": 3}, "errors": 0}
    assert "input" not in body["plan"]

    applied = client.post(f"/v1/onboarding/imports/{body['importId']}:apply")
    assert applied.status_code == 200, applied.text
    result = applied.json()
    assert result["status"] == "APPLIED"
    assert {p["status"] for p in result["result"]["projects"]} == {"DONE"}
    # Active=Yes: activation was attempted; readiness keeps them in setup and says why
    assert len(onboarding["activations"]) == 2
    assert all("Stays in setup" in p["message"] for p in result["result"]["projects"])

    engine = onboarding["engine"]
    with engine.begin() as connection:
        rows = connection.execute(
            text(
                """
                SELECT p.business_code, d.dealer_code, o.outlet_code, o.outlet_classification,
                       o.monthly_vehicle_volume, o.status
                FROM auditcore.dealer_outlets o
                JOIN auditcore.dealers d ON d.tenant_id=o.tenant_id AND d.dealer_id=o.dealer_id
                JOIN auditcore.projects p ON p.tenant_id=o.tenant_id
                WHERE o.tenant_id = ANY(:t) ORDER BY o.outlet_code
                """
            ),
            {"t": onboarding["created"]},
        ).all()
    assert [tuple(r) for r in rows] == [
        (f"T{suffix}-01", "AM-MAH", "AM-MAH-CUBE", "ONSITE", 40, "ACTIVE"),
        (f"T{suffix}-01", "AM-MAH", "AM-MAH-PURI", "SATELLITE", 12, "INACTIVE"),
        (f"T{suffix}-02", "UH-HYU", "UH-HYU-PAHAL", "SATELLITE", 25, "ACTIVE"),
    ]

    # the same file again: nothing to create, it cannot be applied twice
    again = _upload(client, _workbook(suffix)).json()
    assert again["plan"]["summary"]["projects"] == {"UPDATE": 2}  # still to activate
    assert again["plan"]["summary"]["outlets"] == {"UNCHANGED": 3}
    assert client.post(f"/v1/onboarding/imports/{body['importId']}:apply").status_code == 409

    # the export carries the new records with their IDs
    export = client.get("/v1/onboarding/workbook", params={"data": "true"})
    book = load_workbook(io.BytesIO(export.content))
    codes = {row[0].value: row for row in book["Dealerships"].iter_rows(min_row=2) if row[0].value}
    assert {"AM-MAH-CUBE", "AM-MAH-PURI", "UH-HYU-PAHAL"} <= set(codes)
    header = [c.value for c in book["Dealerships"][1]]
    assert codes["AM-MAH-PURI"][header.index("PC Presence")].value == "Satellite"
    assert codes["AM-MAH-PURI"][header.index("Outlet ID")].value


def test_errors_block_the_apply(onboarding):
    client = TestClient(app, raise_server_exceptions=False)
    body = _upload(client, _workbook(onboarding["suffix"], bad=True)).json()
    assert body["status"] == "VALIDATION_FAILED"
    bad = [p for p in body["plan"]["projects"] if p["action"] == "ERROR"]
    assert bad and any("not a known OEM" in m for m in bad[0]["messages"])
    response = client.post(f"/v1/onboarding/imports/{body['importId']}:apply")
    assert response.status_code == 409
    assert onboarding["created"] == []


def test_only_a_super_admin_can_onboard(onboarding):
    app.dependency_overrides[require_super_admin_request] = lambda: _admin(super_admin=False)
    client = TestClient(app, raise_server_exceptions=False)
    assert client.get("/v1/onboarding/workbook").status_code == 403
    assert _upload(client, _workbook(onboarding["suffix"])).status_code == 403
